"""
WLCA calculation workbook - Dashboard and Compare sheets  v0.1  (draft for review)
Called by wlca_calc_template.build(); not run on its own.

Drafted by an AI agent (Claude) - review before trusting.

Dashboard  one option at a time (drop-down at the top): headline tiles, targets on their own boundaries,
           carbon by module and by element group, how firm the number is (modelled / assumed / benchmarked),
           top 10 materials, elements against the GLA typical development, structure against Price & Myers,
           carbon budget for the elements not yet designed.
Compare    every option side by side: table plus stacked element chart and the budget-target chart.

All charts are native Excel charts on cells that hold live formulas, so they move when an input,
factor or quantity changes. Fonts, colours and the logo come from the config "branding" block
(kept private); without it the sheets use neutral defaults.
"""
from openpyxl.chart import BarChart, DoughnutChart, LineChart, Reference
from openpyxl.chart.axis import ChartLines
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.marker import DataPoint
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.text import RichText, Text
from openpyxl.chart.title import Title
from openpyxl.drawing.line import LineProperties
from openpyxl.drawing.text import CharacterProperties, Font as DFont, Paragraph, ParagraphProperties, RegularTextRun
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

from wlca_calc_template import GROUPS, SUMMARY, SUM_ROW, name, sf

DEFAULT = dict(
    font="Arial", logo=None,
    project="1F4E79", target="7F2A3C", ink="222222", muted="7F7F7F", tile="F2F2F2", rule="D9D9D9",
    groups=["8C6D4F", "2F4F4F", "5B8DB8", "E07B1A", "7A3B5A", "79A88B", "D9534F", "9FB3BC"],
    contingency="D9D2C5", contingency_bar="8FB9BE", warn="FCE8C3",
    modules=["1F5F66", "5B8DB8", "2C4A68", "B8C4C8", "B5A82A"],      # A1-A3, A4, A5, B4, C1-C4
    modules_b23=["C98A3E", "E6C190"],                                # B2 maintenance, B3 repair (PS default allowances)
    basis=["1F5F66", "E07B1A", "9FB3BC", "B8C4C8", "D9D2C5"],
    good="DCEFE1", bad="F8D7D3",
)
BOUNDARY_TEXT = {"UP_X8": "A1-A5 excl. ext. works", "UP_ALL": "A1-A5 all elements",
                 "UP_SC": "A1-A5 shell & core, excl. ext. works", "UP_WB": "A1-A5 incl. Cat B + loose FF&E",
                 "AC_X8": "A-C excl. ext. works", "AC_ALL": "A-C (excl. B6/B7)"}
GROUP_LABELS = ["Substructure", "Superstructure", "Facade", "Internal finishes", "FF&E", "Services (MEP)", "External works", "Site (A5.1/A5.2)"]


# building- / asset-level results that are reported to SIGFIG significant figures (PS 6.1.4)
BUILDING_LEVEL = {"UP_ALL", "AC_ALL", "UP_T", "UP_X8", "UP_SC", "UP_WB", "SCN_UP_ALL", "BIO_STORED",
                  "B6_ND", "B6_D", "B7_ND", "B7_D", "WLC_ND", "WLC_D", "EUI", "AC_DEC"}


# --------------------------------------------------------------------------------------- logo
def image_size(data):
    """(format, width, height) of a PNG, JPEG or GIF from its bytes - enough to place it without Pillow."""
    import struct
    if data[:8] == b"\x89PNG\r\n\x1a\n":
        w, h = struct.unpack(">II", data[16:24])
        return "png", w, h
    if data[:6] in (b"GIF87a", b"GIF89a"):
        w, h = struct.unpack("<HH", data[6:10])
        return "gif", w, h
    if data[:2] == b"\xff\xd8":
        i = 2
        while i + 9 < len(data):
            if data[i] != 0xFF:
                i += 1
                continue
            marker, seg = data[i + 1], struct.unpack(">H", data[i + 2:i + 4])[0]
            if 0xC0 <= marker <= 0xCF and marker not in (0xC4, 0xC8, 0xCC):       # start of frame
                h, w = struct.unpack(">HH", data[i + 5:i + 9])
                return "jpeg", w, h
            i += 2 + seg
    raise ValueError("not a PNG, JPEG or GIF image")


def logo_image(path):
    """openpyxl image for the logo. openpyxl reads images through Pillow; where Pillow cannot be imported (Rhino 8's
    Python may only see another program's copy) the file's bytes are embedded as they are - PNG, JPEG or GIF."""
    from openpyxl.drawing.image import Image
    try:
        return Image(path)
    except ImportError:
        pass

    class RawImage(Image):
        def __init__(self, data):              # no Pillow: size from the file header, bytes written unchanged
            self.ref = path
            self._raw = data
            self.format, self.width, self.height = image_size(data)

        def _data(self):
            return self._raw
    with open(path, "rb") as f:
        return RawImage(f.read())


# --------------------------------------------------------------------------------------- styling helpers
def _brand(b):
    out = dict(DEFAULT)
    out.update({k: v for k, v in (b or {}).items() if v})
    return out


def _cp(B, size=900, bold=False, color=None):
    return CharacterProperties(latin=DFont(typeface=B["font"]), sz=size, b=bold, solidFill=color or B["ink"])


def _rich(B, size=900, bold=False, color=None):
    cp = _cp(B, size, bold, color)
    return RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=cp), endParaRPr=cp)])


def _title(B, s):
    cp = _cp(B, 1100, True)
    rt = RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=cp), r=[RegularTextRun(rPr=cp, t=s)])])
    return Title(tx=Text(rich=rt), overlay=False)


def _style_chart(ch, B, title, y_title=None, legend=True, w=17.0, h=8.5):
    ch.title = _title(B, title)
    ch.width, ch.height = w, h
    ch.style = 2
    for ax in (ch.x_axis, ch.y_axis):
        ax.delete = False
        ax.txPr = _rich(B, 900, color=B["muted"])
        ax.spPr = GraphicalProperties(ln=LineProperties(solidFill=B["rule"]))
    if isinstance(ch, BarChart):
        ch.y_axis.scaling.min = 0       # bars always start at zero
    gl = ChartLines()
    gl.spPr = GraphicalProperties(ln=LineProperties(solidFill="EDEDED"))
    ch.y_axis.majorGridlines = gl
    ch.x_axis.majorGridlines = None
    if y_title:
        ch.y_axis.title = _title(B, y_title)
        ch.y_axis.title.tx.rich.p[0].r[0].rPr = _cp(B, 900, False, B["muted"])
    if legend:
        ch.legend.position = "b"
        ch.legend.txPr = _rich(B, 900)
    else:
        ch.legend = None


def _fill(series, hexcol):
    series.graphicalProperties.solidFill = hexcol
    series.graphicalProperties.line.solidFill = "FFFFFF"     # 2px surface gap between adjacent fills
    series.graphicalProperties.line.width = 19050


def _labels(series, B, fmt="#,##0", pos=None):
    dl = DataLabelList()
    dl.showVal = True
    dl.showSerName = dl.showCatName = dl.showLegendKey = dl.showPercent = False
    dl.numFmt = fmt
    dl.txPr = _rich(B, 800, color=B["ink"])
    if pos:
        dl.position = pos
    series.dLbls = dl


def _points(series, colours):
    for i, c in enumerate(colours):
        pt = DataPoint(idx=i)
        pt.graphicalProperties.solidFill = c
        pt.graphicalProperties.line.solidFill = "FFFFFF"
        series.dPt.append(pt)


def _f(B, size=10, bold=False, color=None):
    return Font(name=B["font"], size=size, bold=bold, color=color or B["ink"])


def _note(ws, B, r, text, c2=9, h=22):
    """Muted note wrapped inside A:I so it never runs under the charts on the right."""
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=c2)
    c = ws.cell(r, 1, text)
    c.font = _f(B, 8, color=B["muted"])
    c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[r].height = h


def _section(ws, B, r, text, c1=1, c2=18):
    ws.cell(r, c1, text).font = _f(B, 12, True)
    for c in range(c1, c2 + 1):
        ws.cell(r, c).border = Border(bottom=Side(style="thin", color=B["rule"]))


# --------------------------------------------------------------------------------------- Dashboard
def build_dashboard(wb, ctx, opts, branding):
    B = _brand(branding)
    ws = wb["Dashboard"]
    ws.sheet_view.showGridLines = False
    for c in range(1, 19):
        ws.column_dimensions[chr(64 + c)].width = 10.5
    ws.column_dimensions["J"].width = 24      # result column ("Meets before contingency only")
    for r in range(1, 110):
        ws.row_dimensions[r].height = 15      # fixed rows so chart anchors land the same in Excel and LibreOffice
    TGc, TGb = ctx["TGc"], ctx["TGb"]

    # ---- header
    if B.get("logo"):
        try:
            img = logo_image(B["logo"])
            ratio = img.width / float(img.height)
            img.height, img.width = 30, int(30 * ratio)          # 30 px high ...
            if img.width > 220:                                    # ... and at most 220 px wide (wide wordmarks)
                img.width, img.height = 220, int(220 / ratio)
            ws.add_image(img, "A1")
        except Exception as e:  # missing Pillow or file - keep going without the logo
            ws["A1"] = "(logo not loaded: %s)" % e
    ws.row_dimensions[1].height = 28
    ws.row_dimensions[3].height = 26         # 18 pt title
    ws["A3"] = "Whole life carbon assessment"
    ws["A3"].font = _f(B, 18, True)
    ws["A4"] = '=PROJ&"  |  "&USE&"  |  RIBA stage "&STAGE&"  |  RICS WLCA PS 2nd ed., 60-year RSP"'
    ws["A4"].font = _f(B, 10, color=B["muted"])
    ws["M3"] = "Option"
    ws["M3"].font = _f(B, 10, color=B["muted"])
    ws["M3"].alignment = Alignment(horizontal="right")
    ws.merge_cells("N3:R3")
    sel = ws["N3"]
    sel.value = opts[0]["name"]
    sel.font = _f(B, 14, True, B["project"])
    sel.fill = PatternFill("solid", fgColor="FFF2CC")
    dv = DataValidation(type="list", formula1="=OPT_NAMES", allow_blank=False)
    dv.add("N3")
    ws.add_data_validation(dv)
    name(wb, "DASH_OPT", "Dashboard!$N$3")
    ws["M4"] = "=IFERROR(MATCH(DASH_OPT,OPT_NAMES,0),1)"
    ws["M4"].font = _f(B, 8, color="FFFFFF")
    name(wb, "OPT_IDX", "Dashboard!$M$4")
    ws["N4"] = "Pick an option; every chart below follows it. Compare shows all options."
    ws["N4"].font = _f(B, 8, color=B["muted"])

    # ---- chart data block (T:AA) - live formulas, CHOOSE over the options' Summary blocks
    dc = 20      # column T
    ws.cell(4, dc, "Chart data for the selected option (formulas - do not edit)").font = _f(B, 8, True, B["muted"])
    for key, lbl, unit in SUMMARY:
        r = SUM_ROW[key]
        ws.cell(r, dc, key).font = _f(B, 8, color=B["muted"])
        ws.cell(r, dc + 1, "=CHOOSE(OPT_IDX,%s)" % ",".join("%s!$Y$%d" % (o["RS"], r) for o in opts)).font = _f(B, 8, color=B["muted"])
        ws.cell(r, dc + 1).number_format = "0%" if key.startswith("SH_") else "#,##0.0"
        ws.cell(r, dc + 2, lbl).font = _f(B, 8, color=B["muted"])
    last_key = max(SUM_ROW.values())
    KEYS = "Dashboard!$T$5:$T$%d" % last_key
    VALS = "Dashboard!$U$5:$U$%d" % last_key
    v = lambda key: "Dashboard!$U$%d" % SUM_ROW[key]           # noqa: E731
    pick = lambda key_expr: "INDEX(%s,MATCH(%s,%s,0))" % (VALS, key_expr, KEYS)   # noqa: E731
    tsum = lambda scheme, target, metric: ('SUMIFS(%s,%s,"%s",%s,%s,%s,USE,%s,"%s")'                       # noqa: E731
                                          % (TGc("F"), TGc("A"), scheme, TGc("B"), target, TGc("C"), TGc("E"), metric))
    for c, wdt in zip(range(dc, dc + 9), [12, 11, 40, 2, 26, 11, 11, 11, 11]):
        ws.column_dimensions[ws.cell(1, c).column_letter].width = wdt

    # small chart tables in X:AB (short labels for chart categories)
    xc = 24      # column X
    tables = {}

    def table(r0, title, rows, header=None, x0=None):
        x0 = x0 or xc
        ws.cell(r0, x0, title).font = _f(B, 8, True, B["muted"])
        if header:
            for j, h in enumerate(header):
                ws.cell(r0, x0 + 1 + j, h).font = _f(B, 8, True, B["muted"])
        for i, row in enumerate(rows, start=1):
            for j, val in enumerate(row):
                cc = ws.cell(r0 + i, x0 + j, val)
                cc.font = _f(B, 8, color=B["muted"])
                if j:
                    cc.number_format = "#,##0"     # chart labels and axes follow the source format
        return r0 + 1, r0 + len(rows)

    mc = xc + 6       # column AD: the modules table has 7 rows and would run into the element-group table in X
    tables["mod"] = table(5, "Modules (kgCO2e/m2, incl. contingency)",
                          [("A1-A3", "=%s" % v("M_A13")), ("A4", "=%s" % v("M_A4")), ("A5", "=%s" % v("M_A5")),
                           ("B2", "=%s" % v("M_B2")), ("B3", "=%s" % v("M_B3")),
                           ("B4", "=%s" % v("M_B4")), ("C1-C4", "=%s" % v("M_C"))], x0=mc)
    tables["grp"] = table(12, "Element groups A1-A5 (kgCO2e/m2)",
                          [(GROUP_LABELS[i], "=%s" % v("E_%d" % i)) for i in range(len(GROUPS))] + [("Contingency", "=%s" % v("E_CONT"))])
    tables["basis"] = table(23, "Basis of A1-A5 (kgCO2e/m2)",
                            [("Modelled - project information", "=%s" % v("B_MOD")), ("Modelled - [ASSUMED] quantities", "=%s" % v("B_ASM")),
                             ("Benchmarked (not yet designed)", "=%s" % v("B_BM")), ("Site lumps (A5.1, A5.2)", "=%s" % v("B_SITE")),
                             ("Contingency", "=%s" % v("B_CONT"))])
    top = []
    for j in range(10):
        rr = ctx["m_first"] + j
        top.append(("=CHOOSE(OPT_IDX,%s)" % ",".join("%s!$AV$%d" % (o["MT"], rr) for o in opts),
                    "=CHOOSE(OPT_IDX,%s)" % ",".join("%s!$AU$%d" % (o["MT"], rr) for o in opts)))
    tables["top"] = table(30, "Top 10 materials A1-A3 (tCO2e)", top)
    gla_rows = []
    for i in range(7):
        gcol = chr(ord("N") + i)       # GLA_Breakdown group columns N..T on Targets
        pct = 'SUMIFS(%s,%s,USE,%s,"UP_ALL")' % (TGb(gcol), TGb("L"), TGb("M"))
        gla_rows.append((GROUP_LABELS[i], "=%s*(1+CONT)" % v("E_%d" % i),
                         "=%s*%s" % (pct, tsum("GLA", '"GLA aspirational"', "UP_ALL")),
                         "=%s*%s" % (pct, tsum("GLA", '"GLA benchmark"', "UP_ALL"))))
    tables["gla"] = table(42, "Elements vs GLA typical development (A1-A5, kgCO2e/m2, incl. contingency)", gla_rows,
                          ["Project", "GLA aspirational", "GLA benchmark"])
    bud = [("Target (no contingency)", "=BUD_VAL"),
           ("Modelled + site, incl. contingency", "=" + sf(pick('"MOD_"&BUD_KEY'))),
           ("Left for elements not yet designed", "=" + sf("BUD_VAL-%s" % pick('"MOD_"&BUD_KEY'))),
           ("Not yet designed - typical benchmarks", "=" + sf(pick('"BM_"&BUD_KEY'))),
           ("Not yet designed - scenario tier", "=" + sf(pick('"BMS_"&BUD_KEY')))]
    tables["bud"] = table(52, "Carbon budget (kgCO2e/m2, incl. contingency)", bud)

    # ---- headline tiles (row 6-9)
    _section(ws, B, 6, "Headline - incl. %s contingency" % "PS Table 10")
    ws["J6"] = '="Contingency "&TEXT(CONT,"0%")&" (PS Table 10, RIBA stage "&STAGE&")"'
    ws["J6"].font = _f(B, 9, color=B["muted"])
    # UKNZCBS headline: offices use the shell & core limit (5.1.6.2); other sectors the single limit
    uk_lim = "IF(%s>0,%s,%s)" % (v("UK_LIM_SC"), v("UK_LIM_SC"), v("UK_LIM"))
    uk_val = "IF(%s>0,%s,%s)" % (v("UK_LIM_SC"), v("UP_SC"), v("UP_X8"))
    tiles = [
        ("Upfront carbon A1-A5", "=" + sf(v("UP_ALL")), "#,##0",
         '="kgCO2e/m2 GIA | "&TEXT(%s,"#,##0")&" before contingency"' % sf(v("UP_ALL") + "/(1+CONT)")),
        ("Embodied carbon A-C", "=" + sf(v("AC_ALL")), "#,##0",
         '="incl. biogenic, excl. B6/B7 | "&TEXT(%s,"#,##0")&" before contingency"' % sf("(%s-%s-%s)/(1+CONT)+%s+%s" % (
             v("AC_ALL"), v("BIO_STORED"), v("BIO_RELEASED"), v("BIO_STORED"), v("BIO_RELEASED")))),
        ("Upfront total", "=" + sf(v("UP_T")), "#,##0", "tCO2e"),
        ("LETI upfront band", "=%s" % v("LETI_UP"), "@", '="life cycle: "&%s' % v("LETI_LC")),
        ("UKNZCBS upfront", '=IF({lim}=0,"n/a",IF({val}<={lim},"Meets","Above"))'.format(lim=uk_lim, val=uk_val), "@",
         '=IF({lim}=0,"no limit for this use",IF({sc}>0,"shell & core ","")&TEXT({val},"#,##0")&" vs limit "&TEXT({lim},"#,##0")&" ("&UK_YEAR&")")'
         .format(lim=uk_lim, val=sf(uk_val), sc=v("UK_LIM_SC"))),
        ("Resting on assumptions", "=%s" % v("SH_ASM"), "0%", "of A1-A5 - replace as information arrives"),
    ]
    for i, (lbl, val, fmt, sub) in enumerate(tiles):
        c0 = 1 + 3 * i
        for rr in (7, 8, 9):
            ws.merge_cells(start_row=rr, start_column=c0, end_row=rr, end_column=c0 + 2)
            for cc in range(c0, c0 + 3):
                ws.cell(rr, cc).fill = PatternFill("solid", fgColor=B["tile"])
        ws.cell(7, c0, lbl).font = _f(B, 9, color=B["muted"])
        x = ws.cell(8, c0, val)
        x.font = _f(B, 20, True, B["project"])
        x.number_format = fmt
        ws.cell(9, c0, sub).font = _f(B, 8, color=B["muted"])
        for rr in (7, 8, 9):
            ws.cell(rr, c0).alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[8].height = 30
    # biogenic carbon reported separately alongside the upfront figures (PS: not in A1-A5; included in A-C)
    ws["A10"] = ('=IF(%s=0,"No biogenic carbon stored in the modelled products.","Biogenic carbon stored in products (A1-A3, reported separately, '
                 'not in A1-A5): "&TEXT(%s,"#,##0")&" kgCO2e/m2 GIA. It is released at end of life, so it nets to zero within A-C.")'
                 # PS 5.6.4: state the R1 assumption behind the incineration route
                 '&" Incineration: "&IF(LEFT(EFW,6)="Energy","energy recovery assumed (R1; PS 5.6.4 UK Table 27 data) - C3, '
                 'electricity benefit in D.","treated as disposal (not R1) - C4, no D1.")'
                 % (v("BIO_STORED"), sf(v("BIO_STORED"))))
    ws["A10"].font = _f(B, 9, color=B["muted"])
    ws.conditional_formatting.add("M8", CellIsRule(operator="equal", formula=['"Above"'], font=Font(name=B["font"], size=20, bold=True, color="C0392B")))

    # ---- targets table + chart (rows 11-22)
    _section(ws, B, 11, "Against targets - each compared on its own boundary")
    hdr = ["Target", "", "", "Boundary", "", "Before contingency", "Incl. contingency", "Target (no cont.)",
           "Margin", "Result (incl. contingency)"]
    for j, h in enumerate(hdr, start=1):
        ws.cell(12, j, h).font = _f(B, 9, True, B["muted"])
        ws.cell(12, j).alignment = Alignment(wrap_text=True, vertical="bottom")
    ws.row_dimensions[12].height = 26
    trows = [
        ("SH_UK", '"UKNZCBS "&UK_YEAR&" shell & core limit"', "UP_SC", "UK_LIM_SC"),
        ("SH_UK", '"UKNZCBS "&UK_YEAR&IF(UK_LIM_SC>0," whole-building limit"," limit")', "UP_WB", "UK_LIM"),
        ("SH_LETI", '"LETI 2020 design target (C)"', "UP_X8", tsum("LETI upfront", '"LETI band C"', "UP_X8")),
        ("SH_LETI", '"LETI 2030 design target (A)"', "UP_X8", tsum("LETI upfront", '"LETI band A"', "UP_X8")),
        ("SH_RIBA", '"RIBA 2030 Challenge - 2025"', "AC_X8", tsum("RIBA 2030", '"RIBA 2025"', "AC_X8")),
        ("SH_RIBA", '"RIBA 2030 Challenge - 2030"', "AC_X8", tsum("RIBA 2030", '"RIBA 2030"', "AC_X8")),
        ("SH_GLA", '"GLA benchmark A1-A5"', "UP_ALL", tsum("GLA", '"GLA benchmark"', "UP_ALL")),
        ("SH_GLA", '"GLA aspirational A1-A5"', "UP_ALL", tsum("GLA", '"GLA aspirational"', "UP_ALL")),
        ("SH_GLA", '"GLA benchmark A-C"', "AC_ALL", tsum("GLA", '"GLA benchmark"', "AC_ALL")),
        ("SH_GLA", '"GLA aspirational A-C"', "AC_ALL", tsum("GLA", '"GLA aspirational"', "AC_ALL")),
    ]
    t0 = 13
    for i, (sh, lbl, key, tgt) in enumerate(trows):
        r = t0 + i
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
        ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=5)
        ws.cell(r, 1, '=IF(AND(%s="Y",N(%s)>0),%s,"")' % (sh, tgt, lbl))
        ws.cell(r, 4, '=IF(A%d="","",%s)' % (r, '"%s"' % BOUNDARY_TEXT[key]))
        # summary values include contingency on every target boundary, so "before" = value / (1 + CONT)
        # shown to SIGFIG significant figures (PS 6.1.4); the result below is judged on the unrounded values
        ws.cell(r, 6, '=IF(A%d="",NA(),IF(ISNUMBER(%s),%s,NA()))' % (r, v(key), sf(v(key) + "/(1+CONT)")))
        ws.cell(r, 7, '=IF(A%d="",NA(),IF(ISNUMBER(%s),%s,NA()))' % (r, v(key), sf(v(key))))
        ws.cell(r, 8, '=IF(A%d="",NA(),%s)' % (r, tgt))
        ws.cell(r, 9, '=IF(A{r}="","",IF(ISNA(G{r}),"",H{r}-G{r}))'.format(r=r))
        # judged on the reportable figure (incl. contingency, PS 4.10); middle state shows what the contingency alone tips over
        ws.cell(r, 10, '=IF(A{r}="","",IF(NOT(ISNUMBER({v})),"Not assessed",IF({v}<=H{r},"Meets",IF({v}/(1+CONT)<=H{r},"Meets before contingency only","Above"))))'
                .format(r=r, v=v(key)))
        for c in range(1, 11):
            ws.cell(r, c).font = _f(B, 9)
            ws.cell(r, c).border = Border(bottom=Side(style="hair", color=B["rule"]))
        for c in (6, 7, 8, 9):
            ws.cell(r, c).number_format = '#,##0;-#,##0;0;""'
    t1 = t0 + len(trows) - 1
    ws.conditional_formatting.add("J%d:J%d" % (t0, t1), CellIsRule(operator="equal", formula=['"Meets"'], fill=PatternFill("solid", fgColor=B["good"])))
    ws.conditional_formatting.add("J%d:J%d" % (t0, t1), CellIsRule(operator="equal", formula=['"Meets before contingency only"'],
                                                                    fill=PatternFill("solid", fgColor=B["warn"])))
    ws.conditional_formatting.add("J%d:J%d" % (t0, t1), CellIsRule(operator="equal", formula=['"Above"'], fill=PatternFill("solid", fgColor=B["bad"])))
    _note(ws, B, t1 + 1, "Targets and limits are bare thresholds (no contingency); the result uses the project figure incl. the PS Table 10 contingency, "
                         "the figure the PS requires for reporting (4.10). Building-level results shown to 2 s.f. at RIBA 0-2 and 3 s.f. from RIBA 3 (PS 6.1.4); "
                         "results are judged on the unrounded values. UKNZCBS offices: shell & core = this result less the Cat A share of finishes, FF&E and services "
                       "(Project inputs, 100% shell & core by default); whole building needs a Cat B + loose FF&E allowance. A-C includes B2 / B3 (PS default allowances); B1, B5 not assessed.", h=36)
    na = FormulaRule(formula=["ISNA(F%d)" % t0], font=Font(color="FFFFFF"))      # hide #N/A kept for the chart
    ws.conditional_formatting.add("F%d:H%d" % (t0, t1), na)
    # chart rows (X:AA from row 60): per target one stacked project bar (before + contingency) and one target bar
    c0r = 60
    for j, h in enumerate(["Targets chart", "Project before contingency", "Contingency (PS Table 10)", "Target (no contingency)"]):
        ws.cell(c0r - 1, xc + j, h).font = _f(B, 8, True, B["muted"])
    for i in range(len(trows)):
        r, p_, q_ = t0 + i, c0r + 2 * i, c0r + 2 * i + 1
        vals = [("=A%d" % r, "=F%d" % r, "=G%d-F%d" % (r, r), "=NA()"), ("", "=NA()", "=NA()", "=H%d" % r)]
        for rr_, row_ in zip((p_, q_), vals):
            for j, val in enumerate(row_):
                cc = ws.cell(rr_, xc + j, val)
                cc.font = _f(B, 8, color=B["muted"])
                if j:
                    cc.number_format = "#,##0"
    c1r = c0r + 2 * len(trows) - 1
    ch = BarChart()
    ch.type = "bar"
    ch.grouping = "stacked"
    ch.overlap = 100
    ch.gapWidth = 30
    ch.add_data(Reference(ws, min_col=xc + 1, max_col=xc + 3, min_row=c0r - 1, max_row=c1r), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=xc, min_row=c0r, max_row=c1r))
    _fill(ch.series[0], B["project"])
    _fill(ch.series[1], B["contingency_bar"])
    _fill(ch.series[2], B["target"])
    ch.gapWidth = 15                          # values are in the table beside the chart
    ch.x_axis.scaling.orientation = "maxMin"
    _style_chart(ch, B, "Project (before + contingency) against targets (kgCO2e/m2 GIA)", w=15.5, h=7.6)
    ws.add_chart(ch, "L12")

    # ---- where the carbon is (rows 26-42)
    _section(ws, B, 27, "Where the carbon is")
    a, b = tables["mod"]
    ch = BarChart()
    ch.type = "col"
    ch.add_data(Reference(ws, min_col=mc + 1, min_row=a, max_row=b), titles_from_data=False)
    ch.set_categories(Reference(ws, min_col=mc, min_row=a, max_row=b))
    m5 = (list(B["modules"]) + DEFAULT["modules"])[:5]
    _points(ch.series[0], m5[:3] + list(B.get("modules_b23") or DEFAULT["modules_b23"])[:2] + m5[3:5])
    _labels(ch.series[0], B)
    ch.gapWidth = 50
    _style_chart(ch, B, "By life-cycle module (kgCO2e/m2 GIA)", legend=False, w=15.5, h=8)
    ws.add_chart(ch, "A28")
    a, b = tables["grp"]
    ch = BarChart()
    ch.type = "bar"
    ch.add_data(Reference(ws, min_col=xc + 1, min_row=a, max_row=b), titles_from_data=False)
    ch.set_categories(Reference(ws, min_col=xc, min_row=a, max_row=b))
    _points(ch.series[0], B["groups"] + [B["contingency"]])
    _labels(ch.series[0], B)
    ch.x_axis.scaling.orientation = "maxMin"
    ch.gapWidth = 40
    _style_chart(ch, B, "Upfront A1-A5 by element group (kgCO2e/m2 GIA)", legend=False, w=17.5, h=8)
    ws.add_chart(ch, "J28")

    # ---- how firm + hotspots (rows 45-61)
    _section(ws, B, 45, "How firm the number is, and the carbon hotspots")
    a, b = tables["basis"]
    ch = DoughnutChart(holeSize=55)
    ch.add_data(Reference(ws, min_col=xc + 1, min_row=a, max_row=b), titles_from_data=False)
    ch.set_categories(Reference(ws, min_col=xc, min_row=a, max_row=b))
    _points(ch.series[0], B["basis"])
    dl = DataLabelList()
    dl.showPercent = True
    dl.showVal = dl.showCatName = dl.showSerName = dl.showLegendKey = False
    dl.txPr = _rich(B, 800)
    ch.series[0].dLbls = dl
    ch.title = _title(B, "Basis of upfront A1-A5")
    ch.legend.position = "r"
    ch.legend.txPr = _rich(B, 900)
    ch.width, ch.height = 15.5, 8
    ws.add_chart(ch, "A46")
    a, b = tables["top"]
    ch = BarChart()
    ch.type = "bar"
    ch.add_data(Reference(ws, min_col=xc + 1, min_row=a, max_row=b), titles_from_data=False)
    ch.set_categories(Reference(ws, min_col=xc, min_row=a, max_row=b))
    _fill(ch.series[0], B["project"])
    _labels(ch.series[0], B)
    ch.x_axis.scaling.orientation = "maxMin"
    ch.gapWidth = 40
    _style_chart(ch, B, "Top 10 materials, A1-A3 (tCO2e)", legend=False, w=17.5, h=8)
    ws.add_chart(ch, "J46")

    # ---- elements vs GLA + structure vs P&M (rows 63-80)
    _section(ws, B, 63, "Elements against the GLA typical development, and structure against Price & Myers")
    for j, h in enumerate(["Element group", "", "Project", "GLA aspirational", "GLA benchmark", "Position", "", "", ""], start=1):
        ws.cell(64, j, h).font = _f(B, 9, True, B["muted"])
        ws.cell(64, j).alignment = Alignment(wrap_text=True, vertical="bottom")
    ws.row_dimensions[64].height = 26
    a, b = tables["gla"]
    for i in range(7):
        r, src = 65 + i, a + i
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
        ws.merge_cells(start_row=r, start_column=6, end_row=r, end_column=9)
        ws.cell(r, 1, "=X%d" % src)
        for j, c in enumerate((3, 4, 5)):
            ws.cell(r, c, "=%s%d" % (("Y", "Z", "AA")[j], src)).number_format = "#,##0"
        if i < 3:
            ws.cell(r, 6, '=IF(E{r}=0,"",IF(C{r}<=D{r},"Better than aspirational",IF(C{r}<=E{r},"Within typical","Above typical")))'.format(r=r))
        else:
            ws.cell(r, 6, '="From benchmark - not yet designed"')
        for c in range(1, 10):
            ws.cell(r, c).font = _f(B, 9)
            ws.cell(r, c).border = Border(bottom=Side(style="hair", color=B["rule"]))
    ws.conditional_formatting.add("F65:F71", CellIsRule(operator="equal", formula=['"Better than aspirational"'], fill=PatternFill("solid", fgColor=B["good"])))
    ws.conditional_formatting.add("F65:F71", CellIsRule(operator="equal", formula=['"Above typical"'], fill=PatternFill("solid", fgColor=B["bad"])))
    _note(ws, B, 72, "GLA values = typical-development share x GLA total (London Plan Guidance). Project incl. contingency; site A5.1/A5.2 sits inside the GLA groups so is not shown.")
    ws.cell(74, 1, '="Structure against Price & Myers - "&PM_GROUP&" (n="&PM_N&"), kgCO2e/m2 GIA, before contingency"').font = _f(B, 9, True)
    for j, h in zip((1, 3, 4, 5, 6, 7), ["Stages", "Project", "P&M P25", "P&M median", "P&M P75", "Position"]):
        ws.cell(75, j, h).font = _f(B, 9, True, B["muted"])
    for r, lbl, key, q in ((76, "A1-A5", "STRUCT", ("PM_P25", "PM_P50", "PM_P75")), (77, "A1-C4", "STRUCT_AC", ("PM_C25", "PM_C50", "PM_C75"))):
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
        ws.merge_cells(start_row=r, start_column=7, end_row=r, end_column=9)
        ws.cell(r, 1, lbl)
        ws.cell(r, 3, "=%s" % v(key)).font = _f(B, 9, True, B["project"])
        for c, nm_ in zip((4, 5, 6), q):
            ws.cell(r, c, "=%s" % nm_)
        ws.cell(r, 7, '=IF(N(D{r})=0,"no P&M range",IF(C{r}<D{r},"Below P25 - check for omissions",IF(C{r}<=F{r},"Within P25-P75","Above P75 - look for savings")))'.format(r=r))
        for c in range(1, 10):
            if c != 3:
                ws.cell(r, c).font = _f(B, 9)
            ws.cell(r, c).border = Border(bottom=Side(style="hair", color=B["rule"]))
        for c in (3, 4, 5, 6):
            ws.cell(r, c).number_format = "#,##0"
    ws.conditional_formatting.add("G76:G77", CellIsRule(operator="equal", formula=['"Within P25-P75"'], fill=PatternFill("solid", fgColor=B["good"])))
    ws.conditional_formatting.add("G76:G77", FormulaRule(formula=['OR(LEFT(G76,5)="Below",LEFT(G76,5)="Above")'], fill=PatternFill("solid", fgColor=B["bad"])))
    for r, t in ((78, "P&M scope: structural engineer's elements only - foundations, piles, basement walls, slabs, frame, roof slabs, structural stairs; "
                      "includes site activities (A5a). Excludes facade, finishes, non-loadbearing / secondary elements, services and external works."),
                 (79, "Project: modelled structure (RICS 1, 2.1, 2.2, 2.4) plus the whole-building A5.2 site lump (and C1 for A1-C4). "
                      "The lumps are not split by element, so the project figure leans high."),
                 (80, "P&M method: RICS 2017 / IStructE v2 with ICE v3 factors; A1-C4 includes B1 carbonation. A range to sense-check, not a target.")):
        _note(ws, B, r, t)
    ch = BarChart()
    ch.type = "bar"
    ch.grouping = "clustered"
    ch.add_data(Reference(ws, min_col=xc + 1, max_col=xc + 3, min_row=a - 1, max_row=b), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=xc, min_row=a, max_row=b))
    for s_, c in zip(ch.series, [B["project"], B["groups"][5], B["target"]]):
        _fill(s_, c)
    _labels(ch.series[0], B)
    ch.x_axis.scaling.orientation = "maxMin"
    ch.gapWidth = 50
    _style_chart(ch, B, "Element groups vs GLA typical development (kgCO2e/m2)", w=17.5, h=8.5)
    ws.add_chart(ch, "J64")

    # ---- budget (rows 82-98)
    _section(ws, B, 82, "Carbon budget for the elements not yet designed")
    ws.cell(83, 1, '="Budget target: "&BUD&" ("&BUD_KEY&"), "&TEXT(BUD_VAL,"#,##0")&" kgCO2e/m2"').font = _f(B, 9, True)
    a, b = tables["bud"]
    for i in range(5):
        r, src = 84 + i, a + i
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
        ws.cell(r, 1, "=X%d" % src).font = _f(B, 9)
        x = ws.cell(r, 5, "=Y%d" % src)
        x.number_format = "#,##0"
        x.font = _f(B, 9, i == 2)
    ws.cell(89, 1, '=IF(Y%d<Y%d,"Typical benchmarks for the remaining elements would exceed the budget by "&TEXT(Y%d-Y%d,"#,##0")&" kgCO2e/m2",'
                   '"Typical benchmarks for the remaining elements fit within the budget")' % (a + 2, a + 3, a + 3, a + 2)).font = _f(B, 9, True)
    ws.cell(90, 1, '="Scenario: non-modelled at "&SCN&" - upfront "&TEXT(%s,"#,##0")&" and A-C "&TEXT(%s,"#,##0")&" kgCO2e/m2 (not a compliance result)"'
            % (v("SCN_UP_ALL"), v("SCN_AC_ALL"))).font = _f(B, 9)
    ch = BarChart()
    ch.type = "col"
    ch.add_data(Reference(ws, min_col=xc + 1, min_row=a, max_row=b), titles_from_data=False)
    ch.set_categories(Reference(ws, min_col=xc, min_row=a, max_row=b))
    _points(ch.series[0], [B["target"], B["project"], B["basis"][2], B["groups"][5], B["basis"][1]])
    _labels(ch.series[0], B)
    ch.gapWidth = 60
    _style_chart(ch, B, "Carbon budget (kgCO2e/m2 GIA)", legend=False, w=17.5, h=8)
    ws.add_chart(ch, "J83")

    # ---- operational energy and water, whole life (rows 92-99, left of the budget chart)
    _section(ws, B, 92, "Operational carbon and whole life (B6 / B7)", c2=9)
    ws.cell(93, 1, '=%s&"  |  no contingency on B6 / B7"' % v("OP_STATUS")).font = _f(B, 9, True)
    ws.merge_cells("E94:F94")
    ws.merge_cells("G94:I94")
    for j, h in zip((1, 5, 7), ["kgCO2e/m2 GIA over the RSP", "Non-decarbonised", "Decarbonised (Table 15 / FES)"]):
        ws.cell(94, j, h).font = _f(B, 9, True, B["muted"])
        ws.cell(94, j).alignment = Alignment(horizontal="left" if j == 1 else "right")
    for r, lbl, a_, b_ in ((95, "Embodied A-C (incl. contingency)", "AC_ALL", "AC_DEC"), (96, "B6 operational energy", "B6_ND", "B6_D"),
                           (97, "B7 operational water", "B7_ND", "B7_D"), (98, "Whole life carbon A-C + B6 + B7", "WLC_ND", "WLC_D")):
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
        ws.merge_cells(start_row=r, start_column=5, end_row=r, end_column=6)
        ws.merge_cells(start_row=r, start_column=7, end_row=r, end_column=9)
        ws.cell(r, 1, lbl).font = _f(B, 9, r == 98)
        for c, key in ((5, a_), (7, b_)):
            x = ws.cell(r, c, "=" + sf(v(key)))
            x.number_format = "#,##0"
            x.font = _f(B, 9, r == 98, B["project"] if r == 98 else None)
        for c in range(1, 10):
            ws.cell(r, c).border = Border(bottom=Side(style="hair", color=B["rule"]))
    ws.cell(99, 1, ('=IF(ISNUMBER(%s),"Energy use intensity "&TEXT(%s,"#,##0")&" kWh/m2 GIA/yr - UKNZCBS limit "&TEXT(%s,"#,##0")&" ("&UK_YEAR&"): "'
                    '&IF(%s<=%s,"meets","above")&IF(RIBA_EUI25>0,"; RIBA 2030 Climate Challenge 2025 <"&RIBA_EUI25&", 2030 <"&RIBA_EUI30,""),'
                    '"Energy use intensity: not assessed")&IF(%s<0,"  |  D2 exported PV "&TEXT(%s,"#,##0")&" (memo)","")')
            % (v("EUI"), sf(v("EUI")), v("EUI_LIM"), v("EUI"), v("EUI_LIM"), v("D2_PV"), v("D2_PV"))).font = _f(B, 9)
    ws.cell(100, 1, ('=IF(%s>0,"Water "&TEXT(%s,"#,##0")&" m3/yr"&IF(N(OCC)>0," for "&TEXT(OCC,"#,##0")&" occupants ("&TEXT(%s*1000/OCC/OCC_DAYS,"0.0")&" l/person/day; '
                     'RIBA 2025 <"&RIBA_W_LPD&", budget "&TEXT(W_BUDGET,"#,##0")&" m3/yr)",""),"Water: not assessed - add occupancy (step 1) or the engineer\'s figure")')
             % (v("W_M3"), v("W_M3"), v("W_M3"))).font = _f(B, 9)

    ws.cell(101, 1, "Notes").font = _f(B, 9, True)
    notes = ["Headline and targets include contingency (PS Table 10). The scenario and the budget re-use typical or tier benchmarks for elements not yet designed - "
             "they are optioneering aids, never compliance results.",
             "Rows tagged [ASSUMED] in the take-off rest on office defaults (no SI, no detail yet); the doughnut shows their share.",
             "Sources: UKNZCBS Pilot (2024) Annex A; LETI Embodied Carbon Target Alignment (2022); RIBA 2030 Climate Challenge v2 (2021); "
             "GLA WLCA London Plan Guidance (2022); Price & Myers Embodied Carbon Building Database (May 2025). Values on the Targets sheet.",
             "B6 / B7: DESNZ GHG conversion factors and NESO Future Energy Scenarios (most conservative pathway, excl. BECCS) - see REF_B6B7."]
    for i, t in enumerate(notes):
        ws.cell(102 + i, 1, t).font = _f(B, 8, color=B["muted"])
    ws.freeze_panes = "A6"


# --------------------------------------------------------------------------------------- Compare
def build_compare(wb, ctx, opts, branding):
    B = _brand(branding)
    ws = wb["Compare"]
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 56
    n = len(opts)
    ws["A1"] = "Options compared"
    ws["A1"].font = _f(B, 18, True)
    ws["A2"] = '=PROJ&"  |  "&USE&"  |  same library, benchmarks and inputs for every option; values incl. contingency unless stated; building-level results to "&SIGFIG&" s.f. (PS 6.1.4)"'
    ws["A2"].font = _f(B, 10, color=B["muted"])
    rows = [("Option", "NAME", "@"), ("GIA (m2)", "GIA", "#,##0"),
            ("Upfront A1-A5 (kgCO2e/m2, incl. contingency)", "UP_ALL", "#,##0"), ("Embodied A-C (kgCO2e/m2, incl. contingency)", "AC_ALL", "#,##0"),
            ("Upfront A1-A5 before contingency", "UP_ALL/", "#,##0"), ("Embodied A-C before contingency", "AC_ALL/", "#,##0"),
            ("Upfront total (tCO2e)", "UP_T", "#,##0"),
            ("Biogenic carbon stored, A1-A3 (reported separately; not in A1-A5)", "BIO_STORED", "#,##0"),
            ("A1-A5 excl. ext. works - UKNZCBS / LETI boundary", "UP_X8", "#,##0"),
            ("A1-A5 shell & core excl. ext. works (UKNZCBS offices)", "UP_SC", "#,##0"),
            ("UKNZCBS shell & core limit (offices, no contingency)", "UK_LIM_SC", "#,##0"),
            ("A1-A5 incl. Cat B + loose FF&E (UKNZCBS whole building)", "UP_WB", "#,##0"),
            ("UKNZCBS whole-building limit (no contingency)", "UK_LIM", "#,##0"),
            ("LETI upfront band", "LETI_UP", "@"), ("LETI life-cycle band", "LETI_LC", "@")] + \
           [("%s A1-A5 (before contingency)" % GROUP_LABELS[i], "E_%d" % i, "#,##0") for i in range(len(GROUPS))] + \
           [("Contingency A1-A5", "E_CONT", "#,##0"), ("Share resting on [ASSUMED] quantities", "SH_ASM", "0%"),
            ("Structure A1-A5 incl. A5.2 (P&M scope, before contingency)", "STRUCT", "#,##0"),
            ("Structure A1-C4 incl. A5.2, C1 (P&M scope, before contingency)", "STRUCT_AC", "#,##0"),
            ("Scenario: upfront with non-modelled at scenario tier", "SCN_UP_ALL", "#,##0"),
            ("Operational basis (B6 / B7)", "OP_STATUS", "@"),
            ("Energy use intensity (kWh/m2 GIA/yr)", "EUI", "#,##0"), ("UKNZCBS EUI limit (kWh/m2 GIA/yr)", "EUI_LIM", "#,##0"),
            ("B6 operational energy, without grid decarbonisation (kgCO2e/m2, RSP)", "B6_ND", "#,##0"),
            ("B7 operational water, without grid decarbonisation (kgCO2e/m2, RSP)", "B7_ND", "#,##0"),
            ("Embodied A-C, decarbonised (PS Table 15)", "AC_DEC", "#,##0"),
            ("Whole life A-C + B6 + B7, non-decarbonised", "WLC_ND", "#,##0"),
            ("Whole life A-C + B6 + B7, decarbonised", "WLC_D", "#,##0"),
            ("QA flags open", "QA_FLAGS", "0")]
    for j, o in enumerate(opts):
        ws.column_dimensions[chr(66 + j)].width = 18
    r0 = 4
    for i, (lbl, key, fmt) in enumerate(rows):
        r = r0 + i
        ws.cell(r, 1, lbl).font = _f(B, 10, i == 0)
        for j, o in enumerate(opts):
            ref = "%s!$Y$%d" % (o["RS"], SUM_ROW[key.rstrip("/")])
            if key.endswith("/"):
                ref = ref + "/(1+CONT)"
            c = ws.cell(r, 2 + j, "=" + (sf(ref) if key.rstrip("/") in BUILDING_LEVEL else ref))
            c.number_format = fmt
            c.font = _f(B, 10, i == 0, B["project"] if i == 0 else None)
            c.alignment = Alignment(horizontal="right", wrap_text=(i == 0), vertical="bottom")
            c.border = Border(bottom=Side(style="hair", color=B["rule"]))
        ws.cell(r, 1).border = Border(bottom=Side(style="hair", color=B["rule"]))
    rend = r0 + len(rows) - 1
    # chart data: element groups x options (rows) and budget comparison
    g0 = rend + 3
    ws.cell(g0 - 1, 1, "Chart data (formulas)").font = _f(B, 8, True, B["muted"])
    ws.cell(g0, 1, "Element group").font = _f(B, 8, True, B["muted"])
    for j, o in enumerate(opts):
        ws.cell(g0, 2 + j, "=%s!$Y$%d" % (o["RS"], SUM_ROW["NAME"])).font = _f(B, 8, color=B["muted"])
    labels = GROUP_LABELS + ["Contingency"]
    keys = ["E_%d" % i for i in range(len(GROUPS))] + ["E_CONT"]
    for i, (lbl, key) in enumerate(zip(labels, keys), start=1):
        ws.cell(g0 + i, 1, lbl).font = _f(B, 8, color=B["muted"])
        for j, o in enumerate(opts):
            x = ws.cell(g0 + i, 2 + j, "=%s!$Y$%d" % (o["RS"], SUM_ROW[key]))
            x.font = _f(B, 8, color=B["muted"])
            x.number_format = "#,##0"
    g1 = g0 + len(labels)
    b0 = g1 + 2
    ws.cell(b0, 1, "Budget comparison").font = _f(B, 8, True, B["muted"])
    ws.cell(b0 + 1, 1, '="Project ("&BUD_KEY&", incl. contingency)"').font = _f(B, 8, color=B["muted"])
    ws.cell(b0 + 2, 1, '="Target: "&BUD&" (no contingency)"').font = _f(B, 8, color=B["muted"])
    for j, o in enumerate(opts):
        ws.cell(b0, 2 + j, "=%s!$Y$%d" % (o["RS"], SUM_ROW["NAME"])).font = _f(B, 8, color=B["muted"])
        ws.cell(b0 + 1, 2 + j).number_format = "#,##0"
        ws.cell(b0 + 1, 2 + j, '=INDEX(%s!$Y$5:$Y$%d,MATCH(BUD_KEY,%s!$X$5:$X$%d,0))'
                % (o["RS"], max(SUM_ROW.values()), o["RS"], max(SUM_ROW.values()))).font = _f(B, 8, color=B["muted"])
        ws.cell(b0 + 2, 2 + j, "=BUD_VAL").font = _f(B, 8, color=B["muted"])

    ch = BarChart()
    ch.type = "col"
    ch.grouping = "stacked"
    ch.overlap = 100
    ch.gapWidth = 80
    ch.add_data(Reference(ws, min_col=1, max_col=1 + n, min_row=g0 + 1, max_row=g1), from_rows=True, titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=2, max_col=1 + n, min_row=g0))
    for s_, c in zip(ch.series, B["groups"] + [B["contingency"]]):
        _fill(s_, c)
    _style_chart(ch, B, "Upfront A1-A5 by element group (kgCO2e/m2 GIA)", w=17.5, h=10)
    ws.add_chart(ch, "%s4" % chr(66 + n + 1))

    bc = BarChart()
    bc.type = "col"
    bc.add_data(Reference(ws, min_col=1, max_col=1 + n, min_row=b0 + 1), from_rows=True, titles_from_data=True)
    bc.set_categories(Reference(ws, min_col=2, max_col=1 + n, min_row=b0))
    _fill(bc.series[0], B["project"])
    _labels(bc.series[0], B)
    bc.gapWidth = 120
    ln = LineChart()
    ln.add_data(Reference(ws, min_col=1, max_col=1 + n, min_row=b0 + 2), from_rows=True, titles_from_data=True)
    s_ = ln.series[0]
    s_.graphicalProperties.line.solidFill = B["target"]
    s_.graphicalProperties.line.width = 28575
    s_.graphicalProperties.line.dashStyle = "dash"
    s_.marker.symbol = "none"
    s_.smooth = False
    _style_chart(bc, B, "Against the budget target (kgCO2e/m2 GIA)", w=17.5, h=8)
    bc += ln
    ws.add_chart(bc, "%s26" % chr(66 + n + 1))
    ws.freeze_panes = "B4"
