"""
WLCA calculation template generator  v0.2  (draft for review)
Office WLCA template - build step 5: take-off export + build-up library + benchmarks -> Excel workbook

Drafted by an AI agent (Claude) - review before trusting.

Writes ONE self-contained .xlsx with live formulas, so an assessor can change an input, a factor or a
build-up in Excel and see the result move. Nothing in this script is project data or library data:
every number comes from the three inputs named in the config.

  inputs                              -> sheets
  take-off export(s) (schema v0.1 CSV) -> O<k>_Takeoff   one sheet set per design option (v0.2)
  build-up library (.xlsx)            -> Materials, Layers, Buildups, Rebar_Rates, Offsite_Params, REF_*
  benchmark table (.xlsx)             -> Benchmarks     kgCO2e/m2 GIA by use class and element group
  targets table (.xlsx, optional)     -> Targets        UKNZCBS / LETI / RIBA / GLA values, GLA breakdown, P&M ranges
  branding (optional)                 -> Dashboard / Compare fonts, colours, logo (see wlca_dashboard.py)

Calculation (RICS WLCA PS 2nd ed., module logic ported from the office dwelling workbook):
  Lines     take-off row x build-up layer (16 slots per row, looked up live from Layers)
            A1-A3 = qty x factor; A3 offsite (CWCT) = mass x (assembly + A2 to fabricator)
            + fabrication waste WR/(1-WR) x (A1-A3 + A2); A4 = mass x A4/kg (Materials, lorry loads);
            A5.3 = site WR x (A1-A3 + A3 offsite + A4 + C2-C4 of the waste)  [PS 5.1.4: WR x factor];
            B4 = replacements x (A1-A3 + A3 offsite + A4 + A5.3), replacements = ROUNDUP(RSP/life) - 1;
            C2-C4 and D = mass x per-kg factor x (1 + replacements)
  Materials per material: lorry loads for A4 (T17 distances, REF_Vehicles) and C2 (Table 26
            distances, 50% laden); end of life per Table 23 route (PS 5.6.4 / 5.6.5):
              reuse + recycling -> C3 = share x Table 28 stand-in (no EPD processing data)
              incineration      -> share x fossil combustion CO2, in C3 (energy recovery, UK Table 27
                                   simplification) or C4 (disposal); D1 = -share x Table 27 kWh/t x grid
              landfill          -> C4 = share x Table 28
            library C3 / C4 / D values (EPD, adjusted to the project route) replace the defaults
  Results   RICS element categories x modules; site lumps A5.1 / A5.2 / C1; contingency; per m2 GIA
  QA        checks with OK / FLAG / INFO, and a scope coverage table

Options (v0.2): every option gets O<k>_Takeoff, O<k>_Lines, O<k>_Mat (project-mass-dependent A4/C2),
O<k>_Results (with a Summary block the Dashboard and Compare sheets read) and O<k>_QA. Library, benchmarks,
targets and project inputs are shared, so options are compared like for like.

Run:  python wlca_calc_template.py config.json        (see config.example.json)
"""
import csv
import io
from copy import copy
import json
import os
import sys

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

VERSION = "0.2"
SLOTS = 16                      # layers per build-up the Lines sheet can expand (library max is 15)
SPARE_ROWS = 30                 # blank Takeoff rows for manual additions

# end of life (PS 5.6.4 / 5.6.5). Category IDs are the library's REF_T28_C4 keys.
T28_BIO, T28_INERT = "C4-WOOD-BIOBASED", "C4-AGG-CONC-BRICK-INSUL"
EFW_OPTS = ("Energy recovery - C3 + D1", "Disposal - C4, no D1")
STANDIN_OPTS = ("Table 28 wood and biobased", "Table 28 inert")

INPUT = PatternFill("solid", fgColor="FFF2CC")
HEAD = PatternFill("solid", fgColor="D9E1F2")
SECTION = PatternFill("solid", fgColor="EDEDED")
TOTAL = PatternFill("solid", fgColor="E2EFDA")
GREEN = PatternFill("solid", fgColor="C6EFCE")
AMBER = PatternFill("solid", fgColor="FFEB9C")
RED = PatternFill("solid", fgColor="FFC7CE")
BOLD = Font(bold=True)
TITLE = Font(bold=True, size=14)
WRAP = Alignment(wrap_text=True, vertical="top")
THIN = Border(bottom=Side(style="thin", color="BFBFBF"))

# Reporting rows: (code or criteria, label, kind). kind "x" = exact code, "p" = code prefix.
# Codes follow the office scope matrix (RICS building element categories).
MODELLED_ROWS = [
    ("0", "0 Facilitating works", "p"),
    ("1.1", "1.1 Foundations and piling", "x"),
    ("1.2.1", "1.2.1 Lowest slab", "x"),
    ("1.2.2", "1.2.2 Basement suspended slabs", "x"),
    ("1.2.3", "1.2.3 Basement retaining walls", "x"),
    ("2.1.1", "2.1.1 Frame - vertical (columns, core walls)", "x"),
    ("2.1.2", "2.1.2 Frame - horizontal (beams)", "x"),
    ("2.2.1", "2.2.1 Upper floors and roof slabs", "x"),
    ("2.3", "2.3 Roof - other elements", "x"),
    ("2.4", "2.4 Stairs, ramps, guarding", "p"),
    ("2.5.1", "2.5.1 External - opaque envelope", "x"),
    ("2.5.2", "2.5.2 External - glazed envelope / curtain walling", "x"),
    ("2.5.3", "2.5.3 External - roof coverings", "x"),
    ("2.6", "2.6 Windows and external doors", "x"),
    ("2.7", "2.7 Internal walls and partitions", "x"),
    ("2.8", "2.8 Internal doors", "x"),
    ("3", "3 Finishes", "p"),
    ("4", "4 FF&E", "p"),
    ("5", "5 Services (MEP)", "p"),
    ("6", "6 Prefabricated buildings", "p"),
    ("7", "7 Works to existing buildings", "p"),
    ("8", "8 External works", "p"),
]
# GLA "typical development" element groups used by the Dashboard; and Price & Myers structural scope
GROUPS = ["Substructure", "Superstructure", "Facade", "Internal finishes", "FFE", "Services/MEP", "External works", "Site"]
GROUP_OF = {"0": "Site", "1.1": "Substructure", "1.2.1": "Substructure", "1.2.2": "Substructure", "1.2.3": "Substructure",
            "2.1.1": "Superstructure", "2.1.2": "Superstructure", "2.2.1": "Superstructure", "2.3": "Superstructure",
            "2.4": "Superstructure", "2.5.1": "Facade", "2.5.2": "Facade", "2.5.3": "Superstructure", "2.6": "Facade",
            "2.7": "Superstructure", "2.8": "Superstructure", "3": "Internal finishes", "4": "FFE", "5": "Services/MEP",
            "6": "Superstructure", "7": "Superstructure", "8": "External works",
            "2.7/2.8": "Superstructure", "A5.1": "Site", "A5.2": "Site", "C1": "Site"}
STRUCTURE = {"1.1", "1.2.1", "1.2.2", "1.2.3", "2.1.1", "2.1.2", "2.2.1", "2.4"}
# Summary block on O<k>_Results (key, label, unit) - Dashboard and Compare read these by row
SUMMARY = [
    ("GIA", "GIA", "m2"), ("UP_T", "A1-A5 total incl. contingency", "tCO2e"), ("AC_T", "A-C total incl. contingency", "tCO2e"),
    ("UP_ALL", "A1-A5, all categories (GLA)", "kgCO2e/m2"), ("UP_X8", "A1-A5 excl. external works (UKNZCBS, LETI)", "kgCO2e/m2"),
    ("BC_ALL", "B-C excl. B6/B7 (GLA)", "kgCO2e/m2"), ("AC_ALL", "A-C incl. biogenic stored and released (GLA)", "kgCO2e/m2"),
    ("AC_X8", "A-C incl. biogenic, excl. external works (LETI life cycle, RIBA)", "kgCO2e/m2"),
    ("M_A13", "A1-A3 incl. contingency", "kgCO2e/m2"), ("M_A4", "A4", "kgCO2e/m2"), ("M_A5", "A5", "kgCO2e/m2"),
    ("M_B4", "B4", "kgCO2e/m2"), ("M_C", "C1-C4", "kgCO2e/m2"), ("M_D", "D (memo)", "kgCO2e/m2"),
    ("M_B2", "B2 maintenance incl. contingency (PS 5.2.2)", "kgCO2e/m2"), ("M_B3", "B3 repair incl. contingency (PS 5.2.3)", "kgCO2e/m2"),
    ("B2_BASIS", "B2 basis (the greater of GIA x rate and % of A1-A5)", ""),
] + [("E_%d" % i, "A1-A5 %s (before contingency)" % g, "kgCO2e/m2") for i, g in enumerate(GROUPS)] + [
    ("E_CONT", "A1-A5 contingency", "kgCO2e/m2"),
    ("B_MOD", "A1-A5 modelled from project information", "kgCO2e/m2"), ("B_ASM", "A1-A5 modelled on [ASSUMED] quantities", "kgCO2e/m2"),
    ("B_BM", "A1-A5 benchmarked", "kgCO2e/m2"), ("B_SITE", "A1-A5 site lumps", "kgCO2e/m2"), ("B_CONT", "A1-A5 contingency", "kgCO2e/m2"),
    ("SH_ASM", "Share of A1-A5 on [ASSUMED] quantities", ""),
    ("STRUCT", "Structure A1-A5 + A5.2 site lump, before contingency (P&M scope)", "kgCO2e/m2"),
    ("STRUCT_AC", "Structure A-C + A5.2 and C1 site lumps, before contingency (P&M scope)", "kgCO2e/m2"),
    ("MOD_UP_ALL", "Modelled + site, A1-A5", "kgCO2e/m2"), ("MOD_UP_X8", "Modelled + site, A1-A5 excl. cat 8", "kgCO2e/m2"),
    ("MOD_AC_ALL", "Modelled + site, A-C", "kgCO2e/m2"), ("MOD_AC_X8", "Modelled + site, A-C excl. cat 8", "kgCO2e/m2"),
    ("BM_UP_ALL", "Benchmarked, A1-A5", "kgCO2e/m2"), ("BM_UP_X8", "Benchmarked, A1-A5 excl. cat 8", "kgCO2e/m2"),
    ("BM_AC_ALL", "Benchmarked, A-C", "kgCO2e/m2"), ("BM_AC_X8", "Benchmarked, A-C excl. cat 8", "kgCO2e/m2"),
    ("SCN_UP_ALL", "Scenario: A1-A5 with non-modelled at the scenario tier", "kgCO2e/m2"),
    ("SCN_AC_ALL", "Scenario: A-C with non-modelled at the scenario tier", "kgCO2e/m2"),
    ("BMS_UP_ALL", "Benchmarked at scenario tier, A1-A5", "kgCO2e/m2"), ("BMS_UP_X8", "Benchmarked at scenario tier, A1-A5 excl. cat 8", "kgCO2e/m2"),
    ("BMS_AC_ALL", "Benchmarked at scenario tier, A-C", "kgCO2e/m2"), ("BMS_AC_X8", "Benchmarked at scenario tier, A-C excl. cat 8", "kgCO2e/m2"),
    ("BIO_STORED", "Biogenic carbon stored in products, A1-A3 (reported separately; not in A1-A5)", "kgCO2e/m2"),
    ("BIO_RELEASED", "Biogenic carbon released at end of life (in A-C)", "kgCO2e/m2"),
    ("UP_SC", "A1-A5 shell & core excl. external works (UKNZCBS offices)", "kgCO2e/m2"),
    ("MOD_UP_SC", "Modelled + site, A1-A5 shell & core", "kgCO2e/m2"), ("BM_UP_SC", "Benchmarked, A1-A5 shell & core", "kgCO2e/m2"),
    ("BMS_UP_SC", "Benchmarked at scenario tier, A1-A5 shell & core", "kgCO2e/m2"),
    ("UP_WB", "A1-A5 whole building incl. Cat B allowance (UKNZCBS offices)", "kgCO2e/m2"),
    ("UK_LIM", "UKNZCBS upfront limit, whole building (construction start year)", "kgCO2e/m2"),
    ("UK_LIM_SC", "UKNZCBS upfront limit, offices shell & core", "kgCO2e/m2"), ("LETI_UP", "LETI upfront band", ""),
    ("LETI_LC", "LETI life-cycle band", ""), ("QA_FLAGS", "QA flags open", ""), ("NAME", "Option", ""),
    # operational energy and water (PS 5.3 / 5.4): no contingency; both grid scenarios always reported (PS 4.11.4)
    ("EUI", "Energy use intensity incl. on-site PV used", "kWh/m2 GIA/yr"), ("EUI_LIM", "UKNZCBS EUI limit (construction start year)", "kWh/m2 GIA/yr"),
    ("B6_ND", "B6 operational energy over the RSP, without grid decarbonisation", "kgCO2e/m2"),
    ("B6_D", "B6 operational energy over the RSP, with grid decarbonisation (FES)", "kgCO2e/m2"),
    ("B7_ND", "B7 operational water over the RSP, without grid decarbonisation", "kgCO2e/m2"),
    ("B7_D", "B7 operational water over the RSP, with grid decarbonisation (FES)", "kgCO2e/m2"),
    ("D2_PV", "D2 exported PV over the RSP (memo, decarbonised grid)", "kgCO2e/m2"),
    ("WLC_ND", "Whole life carbon A-C + B6 + B7, without grid decarbonisation", "kgCO2e/m2"),
    ("WLC_D", "Whole life carbon A-C + B6 + B7, with grid decarbonisation (FES)", "kgCO2e/m2"),
    ("OP_STATUS", "Operational basis", ""),
    # material decarbonisation (PS 4.11.5 Table 15): B2-B4 and C1-C2 x 0.5; A1-A5, C3-C4 and biogenic unchanged; D1 x 0.5 (memo)
    ("AC_DEC", "Embodied A-C with material decarbonisation (PS Table 15), incl. contingency", "kgCO2e/m2"),
    ("D_DEC", "D with material decarbonisation (memo)", "kgCO2e/m2"),
    ("W_M3", "Water used (engineer, else occupancy x RIBA 2025 target)", "m3/yr"),
]
OP_ROUTES = ("Engineer prediction (TM54)", "Indicative - EUI estimate or UKNZCBS limit", "Not assessed")
OP_KEYS = (("elec", "Grid electricity (kWh/yr)", "E_ELEC"), ("gas", "Natural gas (kWh/yr)", "E_GAS"), ("dh", "District heat (kWh/yr)", "E_DH"),
           ("pv_self", "PV used on site (kWh/yr)", "E_PVS"), ("pv_export", "PV exported (kWh/yr)", "E_PVX"), ("water", "Water (m3/yr)", "W"),
           ("eui", "Indicative EUI (kWh/m2 GIA/yr)", "EUI_IN"))
OCC_DAYS = {"Office": 250, "Residential": 365, "Schools & Universities": 190, "Retail": 310}
SUM_ROW = dict((k, 5 + i) for i, (k, _, _) in enumerate(SUMMARY))     # row of each key on O<k>_Results (cols X:Z)


def sf(expr):
    """Excel formula rounding `expr` to SIGFIG significant figures (RICS PS 6.1.4: 2 at RIBA 0-2, 3 from RIBA 3).
    For building- / asset-level results only; text passes through, so 'not assessed' still reads."""
    return ('IF(ISNUMBER({x}),IF({x}=0,0,ROUND({x},SIGFIG-1-INT(LOG10(ABS({x}))))),{x})'.format(x=expr))

# Element groups the benchmark table can cover -> key used to drop overlapping modelled layers
BENCH_KEYS = [("2.7/2.8", "2.7 + 2.8 Internal walls and doors"), ("3", "3 Finishes"),
              ("4", "4 FF&E (building-related, S&C + Cat A)"), ("5", "5 Services (MEP)"), ("8", "8 External works")]
# Scope coverage table: code, label, expected default ("Y", "N" or "DEMO" = only if demolition)
COVERAGE = [
    ("0.1", "Demolition / site clearance", "DEMO"), ("1.1", "Foundations and piling", "Y"),
    ("1.2.1", "Lowest slab", "Y"), ("1.2.2", "Basement suspended slabs", "Y"), ("1.2.3", "Basement retaining walls", "Y"),
    ("2.1.1", "Frame - vertical", "Y"), ("2.1.2", "Frame - horizontal (beams, transfers)", "Y"),
    ("2.2.1", "Upper floors and roof slabs", "Y"), ("2.3", "Roof - other elements", "Y"), ("2.4", "Stairs and ramps", "Y"),
    ("2.5.1", "External - opaque envelope", "Y"), ("2.5.2", "External - glazed envelope", "Y"),
    ("2.5.3", "External - roof coverings", "Y"), ("2.6", "Windows and external doors", "Y"),
    ("2.7", "Internal walls and partitions", "Y"), ("2.8", "Internal doors", "Y"), ("3", "Finishes", "Y"),
    ("4", "FF&E", "Y"), ("5", "Services (MEP)", "Y"), ("6", "Prefabricated buildings", "N"),
    ("7", "Works to existing buildings", "N"), ("8", "External works", "Y"),
]
def coverage_expected(code, default, header):
    """Whether the scope check expects an element, from what the take-off export says the design has
    (basement levels, facade mix). Returns (expected, reason); exports without that information keep the default."""
    h = header or {}

    def num(k):
        try:
            return float(h[k])
        except (KeyError, TypeError, ValueError):
            return None
    n_b, depth = num("basement_levels"), num("basement_depth_m")
    cw = num("facade_curtain_wall_share")
    if code == "1.2.2" and n_b is not None:
        if n_b >= 2:
            return "Y", "%d basement levels: suspended slabs between the lowest slab and ground" % n_b
        return "N", ("one basement level: its floor is the lowest slab (1.2.1) and the ground-floor slab over it is a suspended floor "
                     "at the ground line (2.2.1, BR-03) - no basement suspended slabs") if n_b == 1 else "no basement"
    if code == "1.2.3" and n_b is not None:
        return ("Y", "basement %.1f m deep" % (depth or 0)) if (n_b >= 1 or (depth or 0) > 0) else ("N", "no basement")
    if code == "2.5.2" and cw is not None:
        return ("Y", "curtain wall %.0f%% of the net facade" % (100 * cw)) if cw > 0 else \
            ("N", "no curtain wall in the facade mix: the glazing is in punched windows (2.6)")
    return default, ""


MODULES = ["A1-A3", "A4", "A5.1", "A5.2", "A5.3", "A1-A5", "B4", "C1", "C2", "C3", "C4", "A-C", "D", "Bio"]


# --------------------------------------------------------------------------------------- inputs
def read_export(path):
    """Interchange schema v0.1 CSV: '# key: value' header lines, then a CSV table."""
    header, lines = {}, []
    with open(path, encoding="utf-8-sig") as f:
        for ln in f:
            if ln.startswith("#"):
                k, _, v = ln[1:].partition(":")
                header[k.strip()] = v.strip()
            elif ln.strip():
                lines.append(ln)
    rows = list(csv.DictReader(io.StringIO("".join(lines))))
    for r in rows:
        r["quantity"] = float(r["quantity"])
    return header, rows


def read_sheet(wb, name, header_row=4, first=5):
    ws = wb[name]
    hdr = [c.value for c in ws[header_row]]
    out = []
    for r in range(first, ws.max_row + 1):
        vals = [ws.cell(r, c).value for c in range(1, len(hdr) + 1)]
        if vals[0] in (None, ""):
            continue
        out.append(vals)
    return hdr, out


def read_benchmarks(spec):
    """Benchmark table with a header row naming 'Building Type', 'Building Element Category' and module
    columns (A1-A3, A4, A5, B4, C1..C4, optionally D). Returns dict rows."""
    wb = load_workbook(spec["path"], data_only=True, read_only=True)
    ws = wb[spec.get("sheet") or wb.sheetnames[0]]
    hdr, out = None, []
    for row in ws.iter_rows(values_only=True):
        cells = [str(v).strip() if v is not None else "" for v in row]
        if hdr is None:
            if "Building Type" in cells and "Building Element Category" in cells:
                hdr = cells
            continue
        d = dict(zip(hdr, row))
        if not d.get("Building Type") or not d.get("Building Element Category"):
            continue
        out.append(d)
    if hdr is None:
        raise ValueError("Benchmark sheet: header with 'Building Type' / 'Building Element Category' not found")
    return out


def bench_key(label):
    s = str(label).strip()
    if s.startswith("2.7"):
        return "2.7/2.8"
    for k in ("3", "4", "5", "8"):
        if s.startswith(k + " ") or s == k:
            return k
    return None


# --------------------------------------------------------------------------------------- helpers
def col(n):
    return get_column_letter(n)


def head(ws, row, labels, widths=None, height=45):
    for i, h in enumerate(labels, start=1):
        c = ws.cell(row, i, h)
        c.font = BOLD
        c.fill = HEAD
        c.alignment = WRAP
    ws.row_dimensions[row].height = height
    if widths:
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[col(i)].width = w


def title(ws, text, note):
    ws["A1"] = text
    ws["A1"].font = TITLE
    ws["A2"] = note
    ws["A2"].alignment = Alignment(wrap_text=False)


def name(wb, nm, ref):
    dn = DefinedName(nm, attr_text=ref)
    try:
        wb.defined_names[nm] = dn          # openpyxl >= 3.1
    except TypeError:
        wb.defined_names.append(dn)


# --------------------------------------------------------------------------------------- build
def read_targets(path):
    """Targets workbook: sheets Targets, GLA_Breakdown, PM_Structure (header row 1). Returns dict of row lists."""
    if not path:
        return None
    wb = load_workbook(path, data_only=True, read_only=True)
    out = {}
    for nm in ("Targets", "GLA_Breakdown", "PM_Structure"):
        out[nm] = [list(r) for r in wb[nm].iter_rows(values_only=True)] if nm in wb.sheetnames else []
    return out


def _num(v):
    try:
        return float(v) if v not in (None, "") else None
    except (TypeError, ValueError):
        return None


def grid_from_library(lib):
    """(factor, source) from the library's REF_T27_D1 'Grid carbon factor' row; (None, None) if absent."""
    if "REF_T27_D1" not in lib.sheetnames:
        return None, None
    ws = lib["REF_T27_D1"]
    for r in range(1, ws.max_row + 1):
        a = ws.cell(r, 1).value
        if isinstance(a, str) and a.lower().startswith("grid carbon factor"):
            try:
                return float(ws.cell(r, 3).value), ws.cell(r, 4).value
            except (TypeError, ValueError):
                return None, None
    return None, None


def options_from_cfg(cfg):
    """v0.2: cfg["options"] = [{"name", "takeoff_csv", "facade_route" (optional)}]; v0.1 configs with a single takeoff_csv still work."""
    opts = cfg.get("options") or [{"name": None, "takeoff_csv": cfg["takeoff_csv"]}]
    out = []
    for i, o in enumerate(opts, start=1):
        header, rows = read_export(o["takeoff_csv"])
        out.append(dict(k=i, name=o.get("name") or header.get("model_version") or "Option %d" % i,
                        csv=o["takeoff_csv"], header=header, rows=rows, facade_route=o.get("facade_route"),
                        operational=o.get("operational") or {}))
    return out


def copy_sheet(src, dst):
    for row in src.iter_rows():
        for c in row:
            if c.value is not None:
                d = dst.cell(c.row, c.column, c.value)
                if c.has_style:
                    d.font = copy(c.font)
                    d.number_format = c.number_format
                    d.alignment = copy(c.alignment)
                    d.fill = copy(c.fill)
    for k, dim in src.column_dimensions.items():
        dst.column_dimensions[k].width = dim.width


def build(cfg):
    opts = options_from_cfg(cfg)
    header = opts[0]["header"]
    lib = load_workbook(cfg["library_xlsx"])
    bench = read_benchmarks(cfg["benchmarks"])
    tg = read_targets(cfg.get("targets_xlsx"))
    site = dict(A51=35.0, A52=40.0, C1=12.0)
    site.update(cfg.get("site", {}))
    params = dict(contingency=0.15, rsp=60, facade_route="Unitised curtain wall", cw_build_up="CW-UNIT-01",
                  facade_opaque_build_up="EW-SFS-BRK-01",
                  road_vehicle="RV-17PLUS", eol_vehicle="RV-17PLUS", sea_factor=0.01592,
                  eol_distances_km=[50, 50, 50], benchmark_c1="Exclude", cw_reference_a1a5=None,
                  construction_start_year=2027, show_targets=dict(UKNZCBS="Y", LETI="Y", RIBA="Y", GLA="Y"),
                  scenario="GLA aspirational", budget_target="UKNZCBS", pm_group=None,
                  sc_share=dict(finishes=1.0, ffe=1.0, services=1.0), catb_allowance=None,
                  efw_route=EFW_OPTS[0], bio_standin=STANDIN_OPTS[1], grid_d1=None, grid_d1_source=None,
                  op_route=OP_ROUTES[2], op_year=None, dh_factor=0.0, occupancy=None, occ_days=None,
                  b2_rate=10.0, b2_share=0.01, b3_share=0.25, b3_mep_share=0.10)
    params.update(cfg.get("params", {}))
    if params["op_year"] is None:
        params["op_year"] = int(params["construction_start_year"]) + 3
    if params["grid_d1"] is None:
        # the library's REF_T27_D1 carries the grid factor next to Table 27 ("Grid carbon factor" row)
        params["grid_d1"], params["grid_d1_source"] = grid_from_library(lib)

    wb = Workbook()
    wb.remove(wb.active)
    for nm in ["README", "Project", "Dashboard", "Compare"]:
        wb.create_sheet(nm)
    for o in opts:
        for s in ("Results", "QA", "Takeoff", "Lines", "Mat"):
            wb.create_sheet("O%d_%s" % (o["k"], s))
    for nm in ["Materials", "Layers", "Buildups", "Benchmarks", "Targets"]:
        wb.create_sheet(nm)
    for nm in ["Rebar_Rates", "Offsite_Params", "REF_T17_Transport", "REF_T18_Waste", "REF_T20_ServiceLife",
               "REF_T23_EoL", "REF_T28_C4", "REF_T27_D1", "REF_B6B7", "REF_Vehicles"]:
        if nm not in lib.sheetnames:
            if nm == "Offsite_Params":
                raise ValueError("Library has no Offsite_Params sheet (needs library v0.8+)")
            continue
        dst = wb.create_sheet(nm)
        copy_sheet(lib[nm], dst)
        dst.sheet_properties.tabColor = "BFBFBF"

    lib_src = os.path.basename(cfg["library_xlsx"])
    bm_src = cfg["benchmarks"].get("source") or os.path.basename(cfg["benchmarks"]["path"])

    # =================================================================================== Materials (shared)
    mh, mats = read_sheet(lib, "Materials")
    M = wb["Materials"]
    title(M, "Materials - factors (from the library) and per-kg module factors",
          "Columns A-W copied from %s. X onwards: formulas. Project-mass-dependent A4 / C2 per kg are on each O<k>_Mat sheet. "
          "Yellow = may be overridden for this project." % lib_src)
    m_first, m_last = 5, 5 + len(mats) - 1
    rng = lambda c: "Materials!$%s$%d:$%s$%d" % (c, m_first, c, m_last)   # noqa: E731
    # col W (23) = fossil CO2 released on combustion, per kg (library v0.12+); blank on older libraries -> 0
    mcols = (mh + [None])[:22] + [mh[22] if len(mh) > 22 and mh[22] else "Combustion CO2, fossil (kgCO2e/kg)"]
    calc_hdr = ["Life (yrs, T20)", "Site waste rate (T18)", "Road km (T17)", "Sea km (T17)", "Reuse", "Recycling",
                "Incineration", "Disposal", "T28 (kgCO2e/t)",
                "Combustion CO2 used (kgCO2e/kg)", "C3 reuse + recycling: T28 stand-in (kgCO2e/kg)",
                "Incineration: share x combustion CO2 (kgCO2e/kg)", "EfW electricity (kWh/t, T27)",
                "D1 energy recovery (kgCO2e/kg)", "",
                "C3 (kgCO2e/kg)", "C4 (kgCO2e/kg)", "D (kgCO2e/kg)"]
    head(M, 4, [h if h is not None else "" for h in mcols] + calc_hdr,
         [20, 34, 8, 9, 9, 10, 10, 11, 30, 7, 11, 13, 13, 13, 16, 12, 18, 10, 10, 10, 10, 24, 11] + [9] * len(calc_hdr))
    for i, m in enumerate(mats):
        r = m_first + i
        for c in range(1, 24):
            v = m[c - 1] if c - 1 < len(m) else None
            if c == 18:
                v = "=COUNTIF(Layers!$C:$C,A%d)" % r
            if isinstance(v, str) and v.startswith("=") and c != 18:
                v = None
            M.cell(r, c, v)
        for c in (6, 7, 19, 20, 21, 23):
            M.cell(r, c).fill = INPUT
        f = {
            24: '=IFERROR(INDEX(REF_T20_ServiceLife!$D:$D,MATCH(M{r},REF_T20_ServiceLife!$A:$A,0))*1,RSP)',
            25: '=IFERROR(INDEX(REF_T18_Waste!$C:$C,MATCH(L{r},REF_T18_Waste!$A:$A,0))*1,0)',
            26: '=IFERROR(INDEX(REF_T17_Transport!$D:$D,MATCH(K{r},REF_T17_Transport!$A:$A,0))*1,0)',
            27: '=IFERROR(N(INDEX(REF_T17_Transport!$E:$E,MATCH(K{r},REF_T17_Transport!$A:$A,0))),0)',
            28: '=IFERROR(INDEX(REF_T23_EoL!$D:$D,MATCH(N{r},REF_T23_EoL!$A:$A,0))/100,0)',
            29: '=IFERROR(INDEX(REF_T23_EoL!$E:$E,MATCH(N{r},REF_T23_EoL!$A:$A,0))/100,0)',
            30: '=IFERROR(INDEX(REF_T23_EoL!$F:$F,MATCH(N{r},REF_T23_EoL!$A:$A,0))/100,0)',
            31: '=IFERROR(INDEX(REF_T23_EoL!$G:$G,MATCH(N{r},REF_T23_EoL!$A:$A,0))/100,1)',
            32: '=IFERROR(INDEX(REF_T28_C4!$C:$C,MATCH(O{r},REF_T28_C4!$A:$A,0))*1,0)',
            # ---- PS 5.6.4 / 5.6.5 end of life per route (XW 16/10/2026)
            33: '=N(W{r})',
            # reuse + recycling share: processing emissions; with no EPD C3 the PS fallback is the landfill default (Table 28).
            # Bio-based: the inert factor by default - Table 28 wood is landfill methane, charged on the landfilled share only
            34: ('=(AB{r}+AC{r})*IF(AND(O{r}="%s",BIO_STANDIN="%s"),'
                 'IFERROR(INDEX(REF_T28_C4!$C:$C,MATCH("%s",REF_T28_C4!$A:$A,0))*1,0),AF{r})/1000'
                 % (T28_BIO, STANDIN_OPTS[1], T28_INERT)),
            35: '=AD{r}*AG{r}',                                     # fossil CO2 from burning the incinerated share
            36: ('=IFERROR(INDEX(REF_T27_D1!$C:$C,MATCH(IF(O{r}="%s","EFW-BIOMASS","EFW-NONBIOMASS"),'
                 'REF_T27_D1!$A:$A,0))*1,0)' % T28_BIO),
            37: '=IF(EFW="%s",-AD{r}*AJ{r}/1000*GRID_D1,0)' % EFW_OPTS[0],   # Annex K ex. 2c: electricity only
            39: '=IF(S{r}<>"",S{r},AH{r}+IF(EFW="%s",AI{r},0))' % EFW_OPTS[0],
            40: '=IF(T{r}<>"",T{r},AE{r}*AF{r}/1000+IF(EFW="%s",0,AI{r}))' % EFW_OPTS[0],
            41: '=IF(U{r}<>"",N(U{r}),AK{r})',     # a library / EPD D value replaces the EfW estimate (no double counting)
        }
        for c, ftxt in f.items():
            M.cell(r, c, ftxt.format(r=r))
    M.freeze_panes = "C5"

    # =================================================================================== Layers (shared)
    lh, lays = read_sheet(lib, "Layers")
    lays.sort(key=lambda v: (str(v[0]), float(v[1] or 0)))       # contiguous per build-up, by Seq
    L = wb["Layers"]
    title(L, "Layers - Appendix M panels (from the library), quantity per ONE declared unit of the build-up",
          "Copied from %s; mass and unit conversion recomputed against this workbook's Materials." % lib_src)
    l_first, l_last = 5, 5 + len(lays) - 1
    head(L, 4, ["build_up_code", "Seq", "material_id", "Quantity", "Qty unit", "Basis / derivation",
                "Material declared unit", "Qty in material's declared unit", "Mass (kg)", "A1-A3 (kgCO2e)",
                "Biogenic (kgCO2e)", "Layer check", "Reports to RICS code (blank = build-up's)", "Assumed?",
                "Service life override (yrs)", "Site waste rate override", "Offsite fabrication waste rate (A3)"],
         [18, 5, 22, 10, 7, 50, 9, 11, 10, 10, 10, 14, 12, 8, 10, 10, 11])
    MA = rng("A")
    for i, v in enumerate(lays):
        r = l_first + i
        for c in (1, 2, 3, 4, 5, 6, 13, 14, 15, 16, 17):
            x = v[c - 1] if c - 1 < len(v) else None
            L.cell(r, c, x)
        L.cell(r, 4).fill = INPUT
        mm = 'MATCH(C{r},%s,0)' % MA
        L.cell(r, 7, ('=IFERROR(INDEX(%s,%s),"?")' % (rng("C"), mm)).format(r=r))
        L.cell(r, 8, ('=IFERROR(IF(E{r}=G{r},D{r},IF(AND(E{r}="m3",G{r}="kg"),D{r}*INDEX(%s,%s),'
                      'IF(AND(E{r}="kg",G{r}="m3"),D{r}/INDEX(%s,%s),NA()))),"")' % (rng("E"), mm, rng("E"), mm)).format(r=r))
        L.cell(r, 9, ('=IFERROR(H{r}*INDEX(%s,%s),"")' % (rng("D"), mm)).format(r=r))
        L.cell(r, 10, ('=IFERROR(IF(INDEX(%s,%s)="","",H{r}*INDEX(%s,%s)),"")' % (rng("F"), mm, rng("F"), mm)).format(r=r))
        L.cell(r, 11, ('=IFERROR(H{r}*INDEX(%s,%s),"")' % (rng("G"), mm)).format(r=r))
        L.cell(r, 12, ('=IF(COUNTIF(%s,C{r})=0,"Unknown material",IF(H{r}="","Unit mismatch",'
                       'IF(J{r}="","Factor needed",IF(COUNTIF(Buildups!$A:$A,A{r})=0,"Unknown build-up","OK"))))' % MA).format(r=r))
    L.freeze_panes = "D5"
    LA = "Layers!$A$1:$A$%d" % (l_last + SLOTS)

    # =================================================================================== Buildups (shared)
    bh, bus = read_sheet(lib, "Buildups")
    B = wb["Buildups"]
    title(B, "Build-ups - codes the take-off points to (from the library)", "Results per declared unit, recomputed from Layers.")
    b_first, b_last = 5, 5 + len(bus) - 1
    head(B, 4, ["build_up_code", "RICS code", "Description", "Declared unit", "Basis", "Version", "Status", "Provenance",
                "Layers", "Layers not OK", "Mass (kg/unit)", "A1-A3 (kgCO2e/unit)", "Library check", "Offsite (CWCT A3)"],
         [18, 8, 50, 8, 50, 7, 10, 30, 7, 8, 10, 11, 26, 9])
    off_idx = bh.index("Offsite manufactured (CWCT A3)") if "Offsite manufactured (CWCT A3)" in bh else None
    for i, v in enumerate(bus):
        r = b_first + i
        for c in range(1, 9):
            B.cell(r, c, v[c - 1])
        B.cell(r, 9, '=COUNTIF(Layers!$A:$A,A%d)' % r)
        B.cell(r, 10, '=COUNTIFS(Layers!$A:$A,A%d,Layers!$L:$L,"<>OK")' % r)
        B.cell(r, 11, '=IF(I{r}=0,"",SUMIF(Layers!$A:$A,A{r},Layers!$I:$I))'.format(r=r))
        B.cell(r, 12, '=IF(OR(I{r}=0,J{r}>0),"",SUMIF(Layers!$A:$A,A{r},Layers!$J:$J))'.format(r=r))
        B.cell(r, 13, '=IF(I{r}=0,"No layers - basis needed",IF(J{r}>0,"Incomplete - "&J{r}&" layer(s) not OK","OK"))'.format(r=r))
        B.cell(r, 14, (v[off_idx] if off_idx is not None else None) or "N")
    B.freeze_panes = "B5"
    BA, BD, BM, BN, BG = ("Buildups!$%s$%d:$%s$%d" % (c, b_first, c, b_last) for c in "ADMNG")

    # =================================================================================== Benchmarks (shared)
    K = wb["Benchmarks"]
    title(K, "Benchmarks - kgCO2e per m2 GIA for elements not taken off (RICS PS 4.5.2 method 2)",
          "Source: %s. Applied when the use class matches Project and Apply = Y." % bm_src)
    head(K, 4, ["Use class", "Element group", "Key", "A1-A3", "A4", "A5", "B4", "C1", "C2", "C3", "C4", "D",
                "Contingency already inside? (Y/N)", "Apply (Y/N)", "Applied", "Source / note"],
         [20, 34, 8, 8, 8, 8, 8, 8, 8, 8, 8, 8, 12, 8, 8, 60])
    k_first = 5
    kr = 0
    for d in bench:
        key = bench_key(d["Building Element Category"])
        if key is None:
            continue
        r = k_first + kr
        kr += 1
        K.cell(r, 1, str(d["Building Type"]).strip())
        K.cell(r, 2, str(d["Building Element Category"]).strip())
        K.cell(r, 3, key)
        for c, mod in zip(range(4, 13), ["A1-A3", "A4", "A5", "B4", "C1", "C2", "C3", "C4", "D"]):
            v = d.get(mod)
            K.cell(r, c, float(v) if v not in (None, "") else 0.0)
        K.cell(r, 13, "N").fill = INPUT
        K.cell(r, 14, "Y").fill = INPUT
        K.cell(r, 15, '=IF(AND(A{r}=USE,N{r}="Y"),1,0)'.format(r=r))
        K.cell(r, 16, bm_src)
    k_last = k_first + kr - 1
    KC, KO = "Benchmarks!$C$%d:$C$%d" % (k_first, k_last), "Benchmarks!$O$%d:$O$%d" % (k_first, k_last)
    KR = lambda c: "Benchmarks!$%s$%d:$%s$%d" % (c, k_first, c, k_last)    # noqa: E731
    dv = DataValidation(type="list", formula1='"Y,N"', allow_blank=False)
    dv.add("M%d:N%d" % (k_first, k_last))
    K.add_data_validation(dv)
    K.freeze_panes = "D5"

    # =================================================================================== Targets (shared, private data)
    TG = wb["Targets"]
    title(TG, "Targets and reference ranges", "Source: %s. Each scheme is compared on its own boundary (see README)."
          % (os.path.basename(cfg["targets_xlsx"]) if cfg.get("targets_xlsx") else "none configured - target sections read n/a"))
    tgt = (tg or {}).get("Targets") or [["scheme", "target", "use_class", "year_or_band", "metric", "value", "unit", "source", "note"]]
    for i, row in enumerate(tgt):
        for c, v in enumerate(row[:9], start=1):
            TG.cell(4 + i, c, v)
    t_rows = 4 + len(tgt) - 1
    for c in TG[4]:
        c.font, c.fill = BOLD, HEAD
    gb = (tg or {}).get("GLA_Breakdown") or [["use_class", "metric"] + GROUPS[:7]]
    gb_col = 12
    for i, row in enumerate(gb):
        for c, v in enumerate(row[:9], start=gb_col):
            TG.cell(4 + i, c, v)
    gb_last = 4 + len(gb) - 1
    pm = (tg or {}).get("PM_Structure") or [["group", "use_class", "n", "p10", "p25", "p50", "p75", "p90", "ac_p25", "ac_p50", "ac_p75"]]
    pm_col = 23
    for i, row in enumerate(pm):
        for c, v in enumerate(row[:11], start=pm_col):
            TG.cell(4 + i, c, v)
    pm_last = 4 + len(pm) - 1
    for c in list(TG[4])[gb_col - 1:]:
        if c.value is not None:
            c.font, c.fill = BOLD, HEAD
    TGc = lambda c: "Targets!$%s$5:$%s$%d" % (c, c, max(5, t_rows))          # noqa: E731
    TGb = lambda c: "Targets!$%s$5:$%s$%d" % (c, c, max(5, gb_last))        # noqa: E731
    TGp = lambda c: "Targets!$%s$5:$%s$%d" % (c, c, max(5, pm_last))        # noqa: E731
    for c, w in zip("ABCDEFGHI", [16, 22, 20, 12, 8, 8, 14, 50, 30]):
        TG.column_dimensions[c].width = w

    # =================================================================================== Project (shared)
    P = wb["Project"]
    title(P, "Project inputs", "Yellow cells are inputs. Header values come from the take-off export(s); change them here if needed.")
    use_classes = sorted({K.cell(r, 1).value for r in range(k_first, k_last + 1)})
    show = params["show_targets"]
    pm_groups = [r[0] for r in pm[1:] if r and r[0]]
    pm_default = params.get("pm_group") or next((r[0] for r in pm[1:] if r and r[1] == header.get("use_class", "Office")), pm_groups[0] if pm_groups else "")
    items = [
        ("Project ID", header.get("project_id", ""), "", "From export header", "PROJ"),
        ("Use class", header.get("use_class", "Office"), "", "Selects benchmark rows and targets", "USE"),
        ("RIBA stage", int(float(header.get("riba_stage", 2))), "", "", "STAGE"),
        ("Fit-out scope", header.get("fitout_scope", "Shell & core + Cat A"), "",
         "Benchmarks cover building-related items (shell & core + Cat A, London Plan Guidance). Cat B loose items excluded.", None),
        ("Estimated occupancy", _num(params["occupancy"] if params["occupancy"] is not None else header.get("occupancy")), "people",
         "Peak occupants (offices: workplaces; homes: residents; schools: pupils). Sets the water budget (B7) when there is no "
         "public health engineer's figure: occupancy x RIBA 2025 potable water target x occupied days", "OCC"),
        ("Occupied days per year", params["occ_days"] or OCC_DAYS.get(header.get("use_class", "Office"), 250), "days",
         "Office 250, homes 365, schools 190 (RIBA's school target is already per pupil per year)", "OCC_DAYS"),
        ("Reference study period", params["rsp"], "yrs", "PS 4.2", "RSP"),
        ("Contingency", params["contingency"], "", "PS Table 10 - 15% at RIBA 0-2. Applied to A-C (not D, not biogenic).", "CONT"),
        ("Opaque facade build-up to route", params["facade_opaque_build_up"], "",
         "The take-off's opaque facade code (CONFIG build_ups.facade_opaque). Soffits and other 2.5.1 rows are not routed.", "OPQ_BU"),
        ("Curtain wall build-up", params["cw_build_up"], "", "", "CW_BU"),
        ("A5.1 rate (pre-construction demolition)", site["A51"], "kgCO2e/m2 existing GIA", "PS default", "R_A51"),
        ("A5.2 rate (construction activities)", site["A52"], "kgCO2e/m2 GIA", "PS default", "R_A52"),
        ("C1 rate (deconstruction / demolition)", site["C1"], "kgCO2e/m2 GIA", site.get("C1_note", "Lump sum"), "R_C1"),
        ("B2 maintenance - rate", params["b2_rate"], "kgCO2e/m2 GIA",
         "PS 5.2.2 / London Plan Guidance 2.5.12: 10 kgCO2e/m2 GIA or a share of A1-A5, whichever is greater - a one-off total over the RSP "
         "(XW 26/10/2026), covering all element categories", "B2_RATE"),
        ("B2 maintenance - share of A1-A5", params["b2_share"], "",
         "1% of A1-A5: all elements incl. external works and the site lumps, before contingency (contingency is added at the end)", "B2_PCT"),
        ("B3 repair - share of B2 (non-MEP)", params["b3_share"], "",
         "PS 5.2.3: 25% of B2 for the relevant items. B2 split by the non-MEP share of A1-A5 (services = element group 5)", "B3_PCT"),
        ("B3 repair - MEP, share of services A1-A3", params["b3_mep_share"], "",
         "PS 5.2.3 / CIBSE TM65: 10% of A1-A3 of element group 5, initial installation (its replacements are in B4)", "B3_MEP"),
        ("Benchmark C1 shares", params["benchmark_c1"], "",
         "Exclude = C1 only from the whole-building lump above (avoids counting C1 twice)", "BM_C1"),
        ("A4 road vehicle (REF_Vehicles mode ID)", params["road_vehicle"], "", "PS 5.1.3: average rigid HGV", "VEH_A4"),
        ("A4 sea freight factor", params["sea_factor"], "kgCO2e/t.km", "Per tonne-km; ships are not loaded as lorries", "SEA_F"),
        ("C2 vehicle (REF_Vehicles mode ID)", params["eol_vehicle"], "", "PS 5.6.3: average rigid HGV, 50% laden", "VEH_C2"),
        ("C2 distance - waste processing / recycling", params["eol_distances_km"][0], "km",
         "PS Table 26: average of the two closest facilities - PLACEHOLDER, replace per site", "D_PROC"),
        ("C2 distance - energy from waste", params["eol_distances_km"][1], "km", "PLACEHOLDER, replace per site", "D_EFW"),
        ("C2 distance - landfill", params["eol_distances_km"][2], "km", "PLACEHOLDER, replace per site", "D_LAND"),
        ("Incineration (Table 23 route) reported as", params["efw_route"], "",
         "PS 5.6.4: energy recovery only in R1 plants. Default uses the PS UK simplification - Table 27 generic EfW data, "
         "reported in C3 with the energy benefit in D1 - although the PS notes typical UK plants run slightly below R1. "
         "'Disposal' moves the combustion emissions to C4 and drops D1. A-C is the same either way.", "EFW"),
        ("Grid electricity factor for D1 energy recovery", params["grid_d1"], "kgCO2e/kWh",
         params.get("grid_d1_source") or "DESNZ GHG conversion factors, UK electricity generation (excl. T&D) - latest year", "GRID_D1"),
        ("C3 stand-in for recovered bio-based material", params["bio_standin"], "",
         "Wood and bio-based products follow the Table 23 split (e.g. engineered timber 5/35/59/1%). The Table 28 wood factor is fugitive "
         "methane from landfill, so it applies to the landfilled share only (C4). Reused / recycled share: inert factor as the processing "
         "stand-in (default, XW 16/10/2026); 'wood' charges methane on it too.", "BIO_STANDIN"),
        ("TARGETS", None, None, None, None),
        ("Construction start year (UKNZCBS limit year)", params["construction_start_year"], "", "UKNZCBS 4.2.5.3: limits for the year construction commences", "UK_YEAR"),
        ("Show UKNZCBS", show.get("UKNZCBS", "Y"), "", "Upfront A1-A5 excl. external works. Offices: shell & core limit (5.1.6.2) and whole-building limit (incl. Cat B + loose FF&E)", "SH_UK"),
        ("Shell & core share - internal finishes", params["sc_share"]["finishes"], "",
         "UKNZCBS offices: share of element group 3 counted as shell & core (rest = Cat A). 100% = strict default until a split is known", "SC_FIN"),
        ("Shell & core share - FF&E", params["sc_share"]["ffe"], "", "As above, element group 4 (building-related FF&E)", "SC_FFE"),
        ("Shell & core share - services", params["sc_share"]["services"], "", "As above, element group 5 (MEP)", "SC_MEP"),
        ("Cat B fit-out + loose FF&E allowance", params["catb_allowance"], "kgCO2e/m2 GIA",
         "UKNZCBS offices whole-building limit only. Before contingency. Leave blank = whole-building limit 'not assessed'", "CATB"),
        ("Show LETI", show.get("LETI", "Y"), "", "Upfront and life-cycle bands; excl. external works and non-fixed FF&E", "SH_LETI"),
        ("Show RIBA 2030 Climate Challenge", show.get("RIBA", "Y"), "", "A1-A5, B1-B5, C1-C4 incl. sequestration; categories 0-7", "SH_RIBA"),
        ("Show GLA (London Plan Guidance)", show.get("GLA", "Y"), "", "A1-A5 and A-C (excl. B6, B7), all categories incl. external works", "SH_GLA"),
        ("Non-modelled scenario tier", params["scenario"], "",
         "Scenario only - never the headline. Benchmarked elements re-run at GLA aspirational / GLA benchmark per module group", "SCN"),
        ("Budget target", params["budget_target"], "", "Carbon budget chart: target minus modelled + site = what is left for elements not yet designed", "BUD"),
        ("Price & Myers structure group", pm_default, "", "Structural A1-A5 range (structure only; RICS 2017 / IStructE v2 method, ICE v3)", "PM_GROUP"),
        ("OPERATIONAL", None, None, None, None),
        ("Operational energy (B6) basis", params["op_route"], "",
         "PS 5.3: whole life carbon needs a predictive energy model (TM54 or equivalent, never Part L). 'Indicative' uses each option's EUI estimate "
         "(options table), else the UKNZCBS limit, x GIA as all-electric use. Energy and water per option are in the options table.", "OP_ROUTE"),
        ("First full year of operation", params["op_year"], "", "Start of the 60-year FES grid average (PS G1-G2); default construction start + 3", "OP_YEAR"),
        ("District heating carbon factor", params["dh_factor"] or 0.0, "kgCO2e/kWh",
         "Network's published factor incl. losses (PS 5.3.5); used in both scenarios unless the network publishes a projection", "F_DH"),
    ]
    head(P, 4, ["", "Input", "Value", "Unit", "Note / source"], [3, 44, 24, 22, 110])
    r = 5
    lists = {"USE": ",".join(use_classes), "BM_C1": "Exclude,Include",
             "SH_UK": "Y,N", "SH_LETI": "Y,N", "SH_RIBA": "Y,N", "SH_GLA": "Y,N", "SCN": "GLA aspirational,GLA benchmark,None",
             "BUD": "UKNZCBS,LETI 2030 (A),LETI 2020 (C),RIBA 2030,GLA aspirational A1-A5,GLA benchmark A1-A5",
             "PM_GROUP": ",".join(pm_groups), "EFW": ",".join(EFW_OPTS), "BIO_STANDIN": ",".join(STANDIN_OPTS),
             "OP_ROUTE": ",".join(OP_ROUTES)}
    for lbl, val, unit, note, nm in items:
        if lbl in ("TARGETS", "OPERATIONAL"):
            P.cell(r, 2, {"TARGETS": "Targets and benchmarking", "OPERATIONAL": "Operational energy and water (B6 / B7)"}[lbl]).font = BOLD
            r += 1
            continue
        P.cell(r, 2, lbl)
        c = P.cell(r, 3, val)
        c.fill = INPUT
        P.cell(r, 4, unit)
        P.cell(r, 5, note)
        if nm:
            name(wb, nm, "Project!$C$%d" % r)
        if nm in ("CONT", "SC_FIN", "SC_FFE", "SC_MEP", "B2_PCT", "B3_PCT", "B3_MEP"):
            c.number_format = "0%"
        if nm in lists and lists[nm]:
            v = DataValidation(type="list", formula1='"%s"' % lists[nm])
            v.add(c.coordinate)
            P.add_data_validation(v)
        r += 1

    # ---- options table
    r += 1
    P.cell(r, 2, "Design options (one take-off export each)").font = BOLD
    r += 1
    P.cell(r - 1, 5, "Facade route: 'Unitised curtain wall' sends 2.5.1 rows using the opaque facade build-up (OPQ_BU) to the curtain wall "
                     "build-up (CW_BU), reported in 2.5.2 (spandrels are part of the system); 'As exported' keeps the take-off split.")
    fv = DataValidation(type="list", formula1='"Unitised curtain wall,As exported"')
    P.add_data_validation(fv)
    for c, h in enumerate(["#", "Option name", "GIA (m2)", "Existing GIA demolished (m2)", "Facade route", "Take-off export / GIA source"]
                          + [lbl_ for _, lbl_, _ in OP_KEYS], start=1):
        x = P.cell(r, c, h)
        x.font, x.fill = BOLD, HEAD
    r += 1
    opt_first = r
    for o in opts:
        h = o["header"]
        P.cell(r, 1, o["k"])
        P.cell(r, 2, o["name"]).fill = INPUT
        P.cell(r, 3, float(h.get("gia_m2") or 0)).fill = INPUT
        P.cell(r, 4, float(h.get("existing_gia_demolished_m2") or 0)).fill = INPUT
        P.cell(r, 5, o.get("facade_route") or params["facade_route"]).fill = INPUT
        fv.add(P.cell(r, 5).coordinate)
        P.cell(r, 6, "%s | %s" % (os.path.basename(o["csv"]), h.get("gia_source", "")))
        name(wb, "O%d_NAME" % o["k"], "Project!$B$%d" % r)
        name(wb, "O%d_GIA" % o["k"], "Project!$C$%d" % r)
        name(wb, "O%d_GIA_DEMO" % o["k"], "Project!$D$%d" % r)
        name(wb, "O%d_FACADE" % o["k"], "Project!$E$%d" % r)
        for j, (key_, _, nm_) in enumerate(OP_KEYS):          # B6 / B7 inputs per option (engineer's prediction route)
            v_ = o["operational"].get(key_)
            c_ = P.cell(r, 7 + j, float(v_) if v_ not in (None, "") else None)
            c_.fill = INPUT
            c_.number_format = "#,##0"
            name(wb, "O%d_%s" % (o["k"], nm_), "Project!$%s$%d" % (get_column_letter(7 + j), r))
        r += 1
    opt_last = r - 1
    name(wb, "OPT_NAMES", "Project!$B$%d:$B$%d" % (opt_first, opt_last))

    r += 1
    P.cell(r, 2, "Derived parameters (formulas)").font = BOLD
    r += 1
    veh = lambda nm_, colx: '=INDEX(REF_Vehicles!$%s:$%s,MATCH(%s,REF_Vehicles!$A:$A,0))' % (colx, colx, nm_)  # noqa: E731
    tsum = lambda scheme, target, metric: ('SUMIFS(%s,%s,"%s",%s,%s,%s,USE,%s,"%s")'                          # noqa: E731
                                          % (TGc("F"), TGc("A"), scheme, TGc("B"), target, TGc("C"), TGc("E"), metric))
    gla = lambda tier, metric: tsum("GLA", '"GLA %s"' % tier, metric)                                         # noqa: E731
    rb = lambda id_: 'IFERROR(INDEX(REF_B6B7!$C:$C,MATCH("%s",REF_B6B7!$A:$A,0))*1,0)' % id_                    # noqa: E731
    # FES average over the RSP (PS G1-G2): years before the series take its first value, after 2050 its last
    fy, fv = "REF_B6B7!$G$5:$G$60", "REF_B6B7!$H$5:$H$60"
    FES_AVG_F = ("=IFERROR((SUMPRODUCT((%s>=OP_YEAR)*(%s<=OP_YEAR+RSP-1)*%s)"
                 "+MAX(0,OP_YEAR+RSP-1-MAX(%s))*INDEX(%s,MATCH(MAX(%s),%s,0))"
                 "+MAX(0,MIN(%s)-OP_YEAR)*INDEX(%s,MATCH(MIN(%s),%s,0)))/RSP,0)"
                 % (fy, fy, fv, fy, fv, fy, fy, fy, fv, fy, fy))
    derived = [
        ("A4 vehicle payload", veh("VEH_A4", "C"), "t", "V1_PAY"),
        ("A4 vehicle volume", veh("VEH_A4", "D"), "m3", "V1_VOL"),
        ("A4 c1 outward", veh("VEH_A4", "E"), "kgCO2e/km", "V1_C1"),
        ("A4 empty running factor", veh("VEH_A4", "F"), "", "V1_ERF"),
        ("A4 c2 return", veh("VEH_A4", "G"), "kgCO2e/km", "V1_C2"),
        ("C2 vehicle payload", veh("VEH_C2", "C"), "t", "V2_PAY"),
        ("C2 vehicle volume", veh("VEH_C2", "D"), "m3", "V2_VOL"),
        ("C2 50% laden factor", veh("VEH_C2", "H"), "kgCO2e/km", "V2_F50"),
        ("CWCT factory assembly (A3)", '=INDEX(Offsite_Params!$C:$C,MATCH("CWCT-A3-ASSEMBLY",Offsite_Params!$A:$A,0))', "kgCO2e/kg", "CWCT_ASM"),
        ("CWCT transport to fabricator (A2)", '=INDEX(Offsite_Params!$C:$C,MATCH("CWCT-A2-FAB",Offsite_Params!$A:$A,0))', "kgCO2e/kg", "CWCT_A2F"),
        ("Reporting precision, building-level results (significant figures)", "=IF(STAGE<=2,2,3)", "s.f.", "SIGFIG"),
        ("UKNZCBS upfront limit", "=" + tsum("UKNZCBS", '"UKNZCBS "&UK_YEAR', "UP_X8"), "kgCO2e/m2", "UK_LIM"),
        ("Scenario factor A1-A5 (tier / GLA benchmark)",
         '=IF(SCN="None",1,IFERROR(IF(SCN="GLA aspirational",%s,%s)/%s,1))' % (gla("aspirational", "UP_ALL"), gla("benchmark", "UP_ALL"), gla("benchmark", "UP_ALL")), "", "SCN_FUP"),
        ("Scenario factor B-C (tier / GLA benchmark)",
         '=IF(SCN="None",1,IFERROR(IF(SCN="GLA aspirational",%s,%s)/%s,1))' % (gla("aspirational", "BC_ALL"), gla("benchmark", "BC_ALL"), gla("benchmark", "BC_ALL")), "", "SCN_FBC"),
        ("UKNZCBS offices shell & core limit", "=" + tsum("UKNZCBS", '"UKNZCBS S&C "&UK_YEAR', "UP_X8"), "kgCO2e/m2", "UK_LIM_SC"),
        ("Budget target value",
         '=IF(BUD="UKNZCBS",IF(UK_LIM_SC>0,UK_LIM_SC,UK_LIM),IF(BUD="LETI 2030 (A)",%s,IF(BUD="LETI 2020 (C)",%s,IF(BUD="RIBA 2030",%s,IF(BUD="GLA aspirational A1-A5",%s,%s)))))'
         % (tsum("LETI upfront", '"LETI band A"', "UP_X8"), tsum("LETI upfront", '"LETI band C"', "UP_X8"),
            tsum("RIBA 2030", '"RIBA 2030"', "AC_X8"), gla("aspirational", "UP_ALL"), gla("benchmark", "UP_ALL")), "kgCO2e/m2", "BUD_VAL"),
        ("Budget target boundary", '=IF(BUD="RIBA 2030","AC_X8",IF(LEFT(BUD,3)="GLA","UP_ALL",IF(AND(BUD="UKNZCBS",UK_LIM_SC>0),"UP_SC","UP_X8")))', "", "BUD_KEY"),
        ("P&M structure P25", '=SUMIFS(%s,%s,PM_GROUP)' % (TGp("AA"), TGp("W")), "kgCO2e/m2", "PM_P25"),
        ("P&M structure median", '=SUMIFS(%s,%s,PM_GROUP)' % (TGp("AB"), TGp("W")), "kgCO2e/m2", "PM_P50"),
        ("P&M structure P75", '=SUMIFS(%s,%s,PM_GROUP)' % (TGp("AC"), TGp("W")), "kgCO2e/m2", "PM_P75"),
        ("P&M sample size", '=SUMIFS(%s,%s,PM_GROUP)' % (TGp("Y"), TGp("W")), "projects", "PM_N"),
        ("P&M structure A1-C4 P25", '=SUMIFS(%s,%s,PM_GROUP)' % (TGp("AE"), TGp("W")), "kgCO2e/m2", "PM_C25"),
        ("P&M structure A1-C4 median", '=SUMIFS(%s,%s,PM_GROUP)' % (TGp("AF"), TGp("W")), "kgCO2e/m2", "PM_C50"),
        ("P&M structure A1-C4 P75", '=SUMIFS(%s,%s,PM_GROUP)' % (TGp("AG"), TGp("W")), "kgCO2e/m2", "PM_C75"),
        # ---- B6 / B7 factor sets (RICS set 1; REF_B6B7 from the library)
        ("FES grid factor, average over the RSP from the first year of operation", FES_AVG_F, "kgCO2/kWh", "FES_AVG"),
        ("Energy decarbonisation factor (FES average / current generation factor)", "=IFERROR(FES_AVG/%s,0)" % rb("ELEC-GEN"), "", "DECARB"),
        ("Electricity, without grid decarbonisation (gen + T&D + WTT + infrastructure)",
         "=%s+%s+%s+%s+%s+%s" % tuple(rb(x) for x in ("ELEC-GEN", "ELEC-TD", "ELEC-WTT-GEN", "ELEC-WTT-TD", "ELEC-INFRA", "ELEC-INFRA-TD")), "kgCO2e/kWh", "F_E_ND"),
        ("Electricity, with grid decarbonisation (FES; upstream x decarbonisation; infrastructure x 0.5)",
         "=FES_AVG+(%s+%s+%s)*DECARB+(%s+%s)*%s" % tuple(rb(x) for x in ("ELEC-TD", "ELEC-WTT-GEN", "ELEC-WTT-TD", "ELEC-INFRA", "ELEC-INFRA-TD", "EMB-DECARB")),
         "kgCO2e/kWh", "F_E_D"),
        ("Natural gas, without grid decarbonisation (combustion + WTT)", "=%s+%s" % (rb("GAS"), rb("GAS-WTT")), "kgCO2e/kWh", "F_G_ND"),
        ("Natural gas, with grid decarbonisation (WTT decarbonised)", "=%s+%s*DECARB" % (rb("GAS"), rb("GAS-WTT")), "kgCO2e/kWh", "F_G_D"),
        ("Water, without grid decarbonisation (supply + wastewater share x treatment)",
         "=%s+%s*%s" % (rb("WATER-SUPPLY"), rb("WASTEWATER-SHARE"), rb("WATER-TREAT")), "kgCO2e/m3", "F_W_ND"),
        ("Water, with grid decarbonisation (x decarbonisation factor, PS G4)", "=F_W_ND*DECARB", "kgCO2e/m3", "F_W_D"),
        ("UKNZCBS energy use intensity limit (construction start year)", "=" + tsum("UKNZCBS", '"UKNZCBS "&UK_YEAR', "EUI"), "kWh/m2 GIA/yr", "UK_EUI"),
        ("RIBA 2030 Climate Challenge EUI, 2025 target", "=" + tsum("RIBA 2030", '"RIBA 2025"', "EUI"), "kWh/m2 GIA/yr", "RIBA_EUI25"),
        ("RIBA 2030 Climate Challenge EUI, 2030 target", "=" + tsum("RIBA 2030", '"RIBA 2030"', "EUI"), "kWh/m2 GIA/yr", "RIBA_EUI30"),
        ("RIBA 2025 potable water target", "=" + tsum("RIBA 2030", '"RIBA 2025"', "WATER_LPD"), "l/person/day", "RIBA_W_LPD"),
        ("RIBA 2025 potable water target (per pupil, schools)", "=" + tsum("RIBA 2030", '"RIBA 2025"', "WATER_M3PY"), "m3/person/yr", "RIBA_W_M3PY"),
        ("Water budget: occupancy x RIBA 2025 target", "=N(OCC)*(RIBA_W_LPD*OCC_DAYS/1000+RIBA_W_M3PY)", "m3/yr", "W_BUDGET"),
    ]
    for lbl, f, unit, nm in derived:
        P.cell(r, 2, lbl)
        P.cell(r, 3, f)
        P.cell(r, 4, unit)
        name(wb, nm, "Project!$C$%d" % r)
        r += 1

    ctx = dict(m_first=m_first, m_last=m_last, l_last=l_last, LA=LA, BA=BA, BD=BD, BM=BM, BN=BN, BG=BG,
               KC=KC, KO=KO, KR=KR, TGc=TGc, TGb=TGb, params=params, lib_src=lib_src, bm_src=bm_src,
               rng=rng, gb_col=gb_col, pm_col=pm_col)
    for o in opts:
        o.update(build_option(wb, ctx, o))

    # =================================================================================== Dashboard + Compare
    from wlca_dashboard import build_dashboard, build_compare    # noqa: E402
    build_dashboard(wb, ctx, opts, cfg.get("branding") or {})
    build_compare(wb, ctx, opts, cfg.get("branding") or {})

    # =================================================================================== README
    D = wb["README"]
    D.column_dimensions["A"].width = 26
    D.column_dimensions["B"].width = 130
    D["A1"] = "WLCA calculation workbook v%s - %s" % (VERSION, header.get("project_id", ""))
    D["A1"].font = TITLE
    D["A2"] = "Generated by carbon/calc/wlca_calc_template.py (AI-drafted - review before trusting). Regenerate when a take-off export changes."
    readme = [
        ("Inputs", "Take-offs: %s\nLibrary: %s\nBenchmarks: %s\nTargets: %s" % (
            "; ".join("%s = %s" % (o["name"], os.path.basename(o["csv"])) for o in opts), lib_src, bm_src,
            os.path.basename(cfg["targets_xlsx"]) if cfg.get("targets_xlsx") else "none")),
        ("Standard", "RICS Whole life carbon assessment for the built environment, 2nd ed. (PS). RSP 60 years. Modules A1-A5, B2, B3, B4, C1-C4 (embodied) and B6, B7 (operational, when assessed); D and biogenic carbon reported separately. B1 and B5 not in this version."),
        ("How to use", "1 Project: GIA per option, use class, stage, facade route, targets.\n2 O<k>_QA: every FLAG must be fixed or justified.\n"
                       "3 Dashboard: pick an option at the top. Compare: all options side by side. O<k>_Results: the element x module table.\n"
                       "Change factors or layer quantities on Materials / Layers only for project-specific overrides; library-wide changes go to the library owner."),
        ("Sheets", "Project - Dashboard - Compare - per option O<k>_Results / O<k>_QA / O<k>_Takeoff / O<k>_Lines / O<k>_Mat - Materials - Layers - Buildups - Benchmarks - "
                   "Targets - Rebar_Rates, Offsite_Params, REF_* (library reference tables)."),
        ("A1-A3", "Quantity x factor per material layer (GWP fossil + LULUC; biogenic separate). Offsite-manufactured build-ups (CWCT): + factory assembly and transport to fabricator per kg, "
                  "+ fabrication waste WR/(1-WR) x (A1-A3 + A2), reported in A3 (PS 5.1.4: offsite manufacturing waste belongs in A1-A3)."),
        ("A4", "Per material and option: lorry loads = ROUNDUP(max(mass/payload, volume/capacity)) x Table 17 road km x (c1 + empty running x c2), plus sea km x tonnes x sea factor. "
               "Spread over the material's project mass. DEFRA factors already include the return leg, so the empty running factor is 1.0 (see REF_T17_Transport)."),
        ("A5", "A5.1 = existing GIA demolished x rate; A5.2 = GIA x rate (PS defaults). A5.3 = Table 18 waste rate x (A1-A3 + A3 offsite + A4 + C2-C4 of the wasted material) (PS 5.1.4). "
               "A layer may carry its own site waste rate (e.g. unitised curtain wall). A5.4 not assessed."),
        ("B2 / B3", "Whole-building allowances (PS 5.2.2 - 5.2.3), one-off totals over the RSP: B2 = the greater of GIA x 10 kgCO2e/m2 and 1% of A1-A5 "
                    "(all elements incl. external works and site lumps, before contingency). B3 = 25% of the non-MEP part of B2 (B2 x non-services share of A1-A5) "
                    "+ 10% of services (group 5) A1-A3 (CIBSE TM65). Contingency added at the end, like the other modules. Not allocated to elements: "
                    "see the B2 / B3 block under the element table on O<k>_Results."),
        ("B4", "Replacements = ROUNDUP(RSP / life) - 1 (PS 5.2.4: whole replacements, rounded up). Life = layer override, else Table 20 by material. "
               "B4 = replacements x (A1-A3 + A3 offsite + A4 + A5.3). End of life of replaced items is counted in C2-C4 (x (1 + replacements))."),
        ("C1-C4", "C1 = GIA x lump rate. C2 = lorry loads to processing / EfW / landfill (Table 26 distances, 50% laden). "
                  "C3 / C4 follow each material's Table 23 route (a material with no route counts as 100% landfill): "
                  "reuse + recycling -> C3 = share x Table 28 (PS 5.6.4: landfill default when no EPD processing data); "
                  "incineration -> share x fossil CO2 released on combustion (Materials col W: EPD C3 where known, else carbon content x 44/12), "
                  "in C3 as energy recovery or in C4 as disposal (Project input); landfill -> C4 = share x Table 28. "
                  "Biogenic CO2 released at end of life is in the biogenic line; Table 28 wood adds the fugitive methane on the landfilled share (PS 5.6.5 example). "
                  "Library C3 / C4 values replace the defaults and must already be adjusted to the project route shares."),
        ("D", "Memo only, never in A-C. Library value per kg where given (metals); otherwise D1 energy recovery = "
              "-incinerated share x Table 27 electricity (kWh/t) x grid factor (PS Annex K example 2c; heat not counted)."),
        ("B6 / B7", "Operational energy and water per option (options table), both grid scenarios always reported (PS 4.11.4), no contingency. "
                    "Without decarbonisation: kWh x (DESNZ generation + T&D + WTT + infrastructure) x RSP; gas incl. WTT (net CV, as the RICS tables). "
                    "With decarbonisation: the FES 'Falling Behind' (successor to 'Falling Short') average over the RSP from the first year of operation, "
                    "upstream and water x (FES average / current factor), infrastructure x 0.5 (PS G4). Water: m3 x (supply + 90% x treatment). "
                    "Whole life carbon needs a predictive energy model (PS 5.3); the 'Indicative' basis uses each option's EUI estimate, else the UKNZCBS limit, x GIA. "
                    "Water without an engineer's figure: occupancy x RIBA 2025 potable water target x occupied days (a budget). PV exported is reported in D2 (memo). "
                    "Material decarbonisation (PS 4.11.5 Table 15): B2-B4 and C1-C2 x 0.5, A1-A5, C3-C4 and biogenic unchanged, D x 0.5; "
                    "the decarbonised whole life figure = decarbonised A-C + decarbonised B6 + B7."),
        ("Benchmarks", "Elements not taken off use kgCO2e/m2 GIA by use class (PS 4.5.2 method 2). Benchmark A5 is reported in A5.3. Modelled layers in a benchmarked group are dropped (no double counting)."),
        ("Contingency", "PS Table 10 percentage applied to A-C (not D, not biogenic) as a separate row."),
        ("Assumptions", "Take-off rows whose notes start with [ASSUMED] rest on office defaults (no SI, no detail yet). Takeoff column Z marks them; QA, Results and the Dashboard show how many and "
                        "what share of A1-A5 they carry. They are replaced by drawings, model quantities or supplied figures as the project develops - nothing in this workbook needs to change."),
        ("Targets", "Each scheme is compared on its own boundary, incl. contingency: UKNZCBS and LETI upfront = A1-A5 excl. external works (cat. 8), excl. sequestration; "
                    "LETI life cycle and RIBA 2030 = A1-A5 + B + C1-C4 excl. cat. 8, incl. sequestration; GLA = A1-A5 and A-C (excl. B6/B7) incl. cat. 8. "
                    "UKNZCBS Offices whole building includes Cat B and loose FF&E, which this shell & core + Cat A assessment does not - the comparison flatters the building. "
                    "B2 and B3 are PS default allowances; B1 and B5 are not assessed (B5 = 0 like-for-like)."),
        ("Scenario and budget", "The headline always uses typical benchmarks for elements not yet designed. The scenario re-runs those elements at the GLA aspirational (or benchmark) "
                                "tier, scaled per module group, and the budget chart shows what is left for them under the chosen target. Neither is a compliance result."),
    ]
    for i, (k, v) in enumerate(readme):
        D.cell(4 + i, 1, k).font = BOLD
        D.cell(4 + i, 2, v).alignment = WRAP
        D.cell(4 + i, 1).alignment = WRAP

    for ws in wb.worksheets:
        ws.sheet_view.zoomScale = 90
    wb["Dashboard"].sheet_view.zoomScale = 100
    tab = (cfg.get("branding") or {}).get("project") or "1F4E79"
    wb["Dashboard"].sheet_properties.tabColor = tab
    wb["Compare"].sheet_properties.tabColor = tab
    wb["Project"].sheet_properties.tabColor = "FFC000"
    for o in opts:
        wb["O%d_Results" % o["k"]].sheet_properties.tabColor = "70AD47"
        wb["O%d_QA" % o["k"]].sheet_properties.tabColor = "C00000"
    apply_font(wb, (cfg.get("branding") or {}).get("font"))
    wb.active = wb.sheetnames.index("Dashboard")
    wb.calculation.fullCalcOnLoad = True        # no cached values from openpyxl: Excel calculates on open
    return wb


def build_option(wb, ctx, o):
    """Takeoff, Lines, Mat, Results (+ Summary block) and QA for one design option."""
    k, header, rows = o["k"], o["header"], o["rows"]
    TK, LN, MT, RS, QS = ("O%d_%s" % (k, s) for s in ("Takeoff", "Lines", "Mat", "Results", "QA"))
    GIA, GIA_DEMO, FAC = "O%d_GIA" % k, "O%d_GIA_DEMO" % k, "O%d_FACADE" % k
    m_first, m_last, l_last, LA = ctx["m_first"], ctx["m_last"], ctx["l_last"], ctx["LA"]
    BA, BD, BM, BN, BG, KC, KO, KR = (ctx[x] for x in ("BA", "BD", "BM", "BN", "BG", "KC", "KO", "KR"))
    params = ctx["params"]

    # ---------------------------------------------------------------- Takeoff
    T = wb[TK]
    title(T, "%s - take-off rows from the interchange export (schema %s)" % (o["name"], header.get("schema_version", "0.1")),
          "Source: %s. A-I are inputs; J onwards are formulas. Spare rows at the bottom take manual additions." % os.path.basename(o["csv"]))
    t_hdr = ["element_id", "nrm_code", "build_up_code", "quantity", "unit", "method", "include", "source_layer", "notes",
             "Facade routed", "Build-up used", "RICS code used", "In library", "Declared unit", "Unit check", "Library check",
             "Active", "A1-A3 incl. offsite (kgCO2e)", "A4", "A5.3", "A1-A5 (kgCO2e)", "A1-A5 per unit", "B4", "C2-C4",
             "Status", "Basis"]
    head(T, 4, t_hdr, [9, 7, 18, 11, 6, 9, 7, 20, 50, 7, 18, 8, 7, 7, 12, 22, 6, 13, 11, 11, 13, 10, 11, 11, 26, 12])
    t_first = 5
    t_last = t_first + len(rows) + SPARE_ROWS - 1
    for i in range(len(rows) + SPARE_ROWS):
        r = t_first + i
        if i < len(rows):
            d = rows[i]
            vals = [d.get("element_id"), d.get("nrm_code"), d.get("build_up_code"), d["quantity"], d.get("unit"),
                    d.get("method"), d.get("include", "Y") or "Y", d.get("source_layer"), d.get("notes")]
            for c, v in enumerate(vals, start=1):
                T.cell(r, c, v)
        for c in range(1, 10):
            T.cell(r, c).fill = INPUT
        fx = {
            10: '=IF(AND(%s="Unitised curtain wall",' % FAC + 'B{r}&""="2.5.1",C{r}=OPQ_BU),1,0)',
            11: '=IF(C{r}="","",IF(J{r}=1,CW_BU,C{r}))',
            12: '=IF(C{r}="","",IF(J{r}=1,"2.5.2",B{r}&""))',
            13: '=IF(K{r}="","",IF(COUNTIF(%s,K{r})>0,"Y","N"))' % BA,
            14: '=IF(K{r}="","",IFERROR(INDEX(%s,MATCH(K{r},%s,0)),""))' % (BD, BA),
            15: '=IF(K{r}="","",IF(N{r}="","No build-up",IF(N{r}=E{r},"OK","Unit mismatch")))',
            16: '=IF(K{r}="","",IFERROR(INDEX(%s,MATCH(K{r},%s,0)),"Not in library"))' % (BM, BA),
            17: '=IF(AND(K{r}<>"",G{r}="Y",M{r}="Y",O{r}="OK"),1,0)',
            18: '=SUMIF(%s!$A:$A,ROW(),%s!$Q:$Q)+SUMIF(%s!$A:$A,ROW(),%s!$T:$T)' % (LN, LN, LN, LN),
            19: '=SUMIF(%s!$A:$A,ROW(),%s!$U:$U)' % (LN, LN),
            20: '=SUMIF(%s!$A:$A,ROW(),%s!$W:$W)' % (LN, LN),
            21: '=R{r}+S{r}+T{r}',
            22: '=IF(N(D{r})=0,"",U{r}/D{r})',
            23: '=SUMIF(%s!$A:$A,ROW(),%s!$Z:$Z)' % (LN, LN),
            24: ('=SUMIF(%s!$A:$A,ROW(),%s!$AA:$AA)+SUMIF(%s!$A:$A,ROW(),%s!$AB:$AB)'
                 '+SUMIF(%s!$A:$A,ROW(),%s!$AC:$AC)') % (LN, LN, LN, LN, LN, LN),
            25: ('=IF(C{r}="","",IF(G{r}<>"Y","Excluded (include = N)",IF(M{r}="N","NOT IN LIBRARY - not counted",'
                 'IF(O{r}<>"OK","UNIT MISMATCH - not counted",IF(P{r}<>"OK","Build-up "&P{r},"OK")))))'),
            26: ('=IF(C{r}="","",IF(F{r}="Supplied","Supplied",IF(LEFT(I{r},9)="[ASSUMED]","Assumed",'
                 'IF(F{r}="Modelled","Modelled","Derived"))))'),
        }
        for c, f in fx.items():
            T.cell(r, c, f.format(r=r))
    T.freeze_panes = "D5"
    dvt = DataValidation(type="list", formula1='"Y,N"')
    dvt.add("G%d:G%d" % (t_first, t_last))
    T.add_data_validation(dvt)
    TR = lambda c: "%s!$%s$%d:$%s$%d" % (TK, c, t_first, c, t_last)    # noqa: E731

    # ---------------------------------------------------------------- Lines
    N = wb[LN]
    title(N, "%s - each take-off row expanded into its build-up layers (do not edit)" % o["name"],
          "kgCO2e. One block of %d slots per Takeoff row; empty slots read 0. Layers whose element group is benchmarked are dropped (col I)." % SLOTS)
    n_hdr = ["Takeoff row", "Slot", "Build-up", "Layers row", "Valid", "material_id", "Reports to", "Group key",
             "Dropped (benchmarked)", "Active", "Qty/unit (decl.)", "Mass/unit (kg)", "Take-off qty", "Qty (decl. unit)",
             "Mass (kg)", "Materials row", "A1-A3", "Offsite flag", "Fab. waste rate", "A3 offsite (CWCT)", "A4",
             "Site waste rate", "A5.3", "Life (yrs)", "Replacements", "B4", "C2", "C3", "C4", "D", "Biogenic stored",
             "Biogenic released"]
    head(N, 4, n_hdr, [7, 5, 18, 8, 5, 20, 8, 7, 9, 6, 9, 9, 10, 11, 11, 8] + [10] * 16)
    n_first = 5
    MAcol = "Materials!$A$%d:$A$%d" % (m_first, m_last)
    mi = lambda c: "INDEX(Materials!$%s$%d:$%s$%d,P{r})" % (c, m_first, c, m_last)   # noqa: E731
    mo = lambda c: "INDEX(%s!$%s$%d:$%s$%d,P{r})" % (MT, c, m_first, c, m_last)       # noqa: E731
    li = lambda c: "INDEX(Layers!$%s$1:$%s$%d,D{r})" % (c, c, l_last + SLOTS)         # noqa: E731
    r = n_first
    for t in range(t_first, t_last + 1):
        for s in range(1, SLOTS + 1):
            fx = {
                1: t, 2: s,
                3: '=%s!K%d' % (TK, t),
                4: '=IF(C{r}="",0,IFERROR(MATCH(C{r},%s,0)+B{r}-1,0))' % LA,
                5: '=IF(D{r}=0,0,IF(%s=C{r},1,0))' % li("A"),
                6: '=IF(E{r}=1,%s,"")' % li("C"),
                7: '=IF(E{r}<>1,"",IF(%s&""="",%s!L%d,%s&""))' % (li("M"), TK, t, li("M")),
                8: '=IF(G{r}="","",IF(OR(LEFT(G{r},3)="2.7",LEFT(G{r},3)="2.8"),"2.7/2.8",LEFT(G{r},1)))',
                9: '=IF(H{r}="",0,IF(COUNTIFS(%s,H{r},%s,1)>0,1,0))' % (KC, KO),
                10: '=IF(AND(E{r}=1,%s!Q%d=1,I{r}=0),1,0)' % (TK, t),
                11: '=IF(E{r}=1,N(%s),0)' % li("H"),
                12: '=IF(E{r}=1,N(%s),0)' % li("I"),
                13: '=N(%s!D%d)' % (TK, t),
                14: '=J{r}*M{r}*K{r}',
                15: '=J{r}*M{r}*L{r}',
                16: '=IF(F{r}="",0,IFERROR(MATCH(F{r},%s,0),0))' % MAcol,
                17: '=IF(P{r}=0,0,N{r}*N(%s))' % mi("F"),
                18: '=IF(E{r}<>1,0,IF(IFERROR(INDEX(%s,MATCH(C{r},%s,0)),"N")="Y",1,0))' % (BN, BA),
                19: '=IF(E{r}=1,N(%s),0)' % li("Q"),
                20: '=IF(J{r}=0,0,R{r}*(O{r}*(CWCT_ASM+CWCT_A2F)+IF(S{r}>0,S{r}/(1-S{r})*(Q{r}+O{r}*CWCT_A2F),0)))',
                21: '=IF(P{r}=0,0,O{r}*%s)' % mo("AJ"),
                22: '=IF(E{r}<>1,0,IF(%s&""<>"",N(%s),IF(P{r}=0,0,%s)))' % (li("P"), li("P"), mi("Y")),
                23: '=IF(P{r}=0,0,V{r}*(Q{r}+T{r}+U{r}+O{r}*%s))' % mo("AP"),
                24: '=IF(E{r}<>1,RSP,IF(%s&""<>"",N(%s),IF(P{r}=0,RSP,%s)))' % (li("O"), li("O"), mi("X")),
                25: '=IF(OR(J{r}=0,X{r}<=0),0,MAX(0,ROUNDUP(RSP/X{r},0)-1))',
                26: '=Y{r}*(Q{r}+T{r}+U{r}+W{r})',
                27: '=IF(P{r}=0,0,O{r}*%s*(1+Y{r}))' % mo("AL"),
                28: '=IF(P{r}=0,0,O{r}*%s*(1+Y{r}))' % mi("AM"),
                29: '=IF(P{r}=0,0,O{r}*%s*(1+Y{r}))' % mi("AN"),
                30: '=IF(P{r}=0,0,O{r}*%s*(1+Y{r}))' % mi("AO"),
                31: '=IF(P{r}=0,0,N{r}*N(%s)*(1+Y{r}))' % mi("G"),
                32: '=-AE{r}',
            }
            for c, f in fx.items():
                N.cell(r, c, f.format(r=r) if isinstance(f, str) else f)
            r += 1
    n_last = r - 1
    N.freeze_panes = "C5"
    NR = lambda c: "%s!$%s$%d:$%s$%d" % (LN, c, n_first, c, n_last)     # noqa: E731

    # ---------------------------------------------------------------- Mat (project-mass-dependent per-kg factors)
    X = wb[MT]
    title(X, "%s - per-material project totals: A4 and C2 lorry loads depend on this option's masses" % o["name"],
          "Rows line up with Materials. Top 10 materials by A1-A3 in AT:AV (read by the Dashboard).")
    hdr = {1: "material_id", 2: "Name", 33: "Project mass (kg)", 34: "Project volume (m3)", 35: "A4 total (kgCO2e)",
           36: "A4 (kgCO2e/kg)", 37: "C2 total (kgCO2e)", 38: "C2 (kgCO2e/kg)", 42: "EoL C2-C4 (kgCO2e/kg)",
           43: "A1-A3 project (kgCO2e)", 44: "In use", 45: "Rank key", 46: "Rank", 47: "A1-A3 (tCO2e)", 48: "Material"}
    for c, h in hdr.items():
        x = X.cell(4, c, h)
        x.font, x.fill, x.alignment = BOLD, HEAD, WRAP
    X.column_dimensions["B"].width = 34
    for i in range(m_last - m_first + 1):
        r = m_first + i
        f = {
            1: '=Materials!A{r}', 2: '=Materials!B{r}',
            33: '=SUMIF(%s!$F:$F,A{r},%s!$O:$O)' % (LN, LN),
            34: '=IF(N(Materials!E{r})>0,AG{r}/Materials!E{r},0)',
            35: ('=IF(AG{r}=0,0,ROUNDUP(MAX(AG{r}/1000/V1_PAY,AH{r}/V1_VOL),0)*Materials!Z{r}*(V1_C1+V1_ERF*V1_C2)'
                 '+AG{r}/1000*Materials!AA{r}*SEA_F)'),
            36: '=IF(AG{r}=0,0,AI{r}/AG{r})',
            37: ('=IF(AG{r}=0,0,(ROUNDUP(MAX((Materials!AB{r}+Materials!AC{r})*AG{r}/1000/V2_PAY,(Materials!AB{r}+Materials!AC{r})*AH{r}/V2_VOL),0)*D_PROC'
                 '+ROUNDUP(MAX(Materials!AD{r}*AG{r}/1000/V2_PAY,Materials!AD{r}*AH{r}/V2_VOL),0)*D_EFW'
                 '+ROUNDUP(MAX(Materials!AE{r}*AG{r}/1000/V2_PAY,Materials!AE{r}*AH{r}/V2_VOL),0)*D_LAND)*2*V2_F50)'),
            38: '=IF(AG{r}=0,0,AK{r}/AG{r})',
            42: '=AL{r}+Materials!AM{r}+Materials!AN{r}',
            43: '=SUMIF(%s!$F:$F,A{r},%s!$Q:$Q)' % (LN, LN),
            44: '=IF(AG{r}>0,"Y","")',
            45: '=AQ{r}+ROW()*1E-9',
        }
        for c, ftxt in f.items():
            X.cell(r, c, ftxt.format(r=r))
    RK = "%s!$AS$%d:$AS$%d" % (MT, m_first, m_last)
    for j in range(10):
        r = m_first + j
        X.cell(r, 46, j + 1)
        X.cell(r, 47, '=IFERROR(LARGE(%s,%d)/1000,0)' % (RK, j + 1))
        X.cell(r, 48, '=IFERROR(INDEX(Materials!$B$%d:$B$%d,MATCH(LARGE(%s,%d),%s,0)),"")' % (m_first, m_last, RK, j + 1, RK))
    X.freeze_panes = "C5"
    MO = lambda c: "%s!$%s$%d:$%s$%d" % (MT, c, m_first, c, m_last)     # noqa: E731
    MR_ = lambda c: "Materials!$%s$%d:$%s$%d" % (c, m_first, c, m_last)  # noqa: E731

    # ---------------------------------------------------------------- Results
    R = wb[RS]
    title(R, "%s - whole life carbon by RICS element and module (tCO2e unless stated)" % o["name"],
          "RICS WLCA PS 2nd ed., RSP from Project. A-C = A1-A5 + B2 + B3 + B4 + C1-C4 (B2 / B3 whole-building allowances in the block below the table; "
          "B1, B5 not assessed; B6 / B7 in the Summary). D and biogenic are reported separately.")
    rh = ["Code", "Element", "Method", "A1-A3", "A4", "A5.1", "A5.2", "A5.3", "A1-A5", "B4", "C1", "C2", "C3", "C4",
          "A-C excl. biogenic", "D (memo)", "Biogenic stored, A1-A3 (reported separately)", "A1-A5 kgCO2e/m2 GIA", "A-C excl. biogenic kgCO2e/m2 GIA", "Group (GLA)", "Kind", "Structure (P&M)"]
    head(R, 13, rh, [8, 46, 12] + [10] * 14 + [11, 11, 16, 11, 9])
    line_cols = {"A1-A3": ["Q", "T"], "A4": ["U"], "A5.3": ["W"], "B4": ["Z"], "C2": ["AA"], "C3": ["AB"], "C4": ["AC"],
                 "D": ["AD"], "Bio": ["AE"]}
    mod_col = {m: 4 + i for i, m in enumerate(["A1-A3", "A4", "A5.1", "A5.2", "A5.3", "A1-A5", "B4", "C1", "C2", "C3",
                                                "C4", "A-C", "D", "Bio"])}

    def row_formulas(rr):
        R.cell(rr, mod_col["A1-A5"], "=SUM(D{r}:H{r})".format(r=rr))
        R.cell(rr, mod_col["A-C"], "=I{r}+SUM(J{r}:N{r})".format(r=rr))
        R.cell(rr, 18, "=IF({g}>0,I{r}*1000/{g},0)".format(r=rr, g=GIA))
        R.cell(rr, 19, "=IF({g}>0,O{r}*1000/{g},0)".format(r=rr, g=GIA))
        for c in range(4, 20):
            R.cell(rr, c).number_format = "#,##0.0"

    def tag(rr, code, kind):
        R.cell(rr, 20, GROUP_OF.get(code, "Superstructure"))
        R.cell(rr, 21, kind)
        R.cell(rr, 22, "Y" if (kind == "Modelled" and code in STRUCTURE) else "")

    rr = 14
    R.cell(rr, 1, "Modelled (take-off x build-up library)").font = BOLD
    for c in range(1, 23):
        R.cell(rr, c).fill = SECTION
    rr += 1
    mod_first = rr
    method_rows = []
    for code, label, kind in MODELLED_ROWS:
        crit = '"%s"' % code if kind == "x" else '"%s*"' % code
        R.cell(rr, 1, code)
        R.cell(rr, 2, label)
        method_rows.append((rr, crit))
        for m, cols in line_cols.items():
            parts = ["SUMIFS(%s,%s,%s)" % (NR(cc), NR("G"), crit) for cc in cols]
            R.cell(rr, mod_col[m], "=(" + "+".join(parts) + ")/1000")
        row_formulas(rr)
        tag(rr, code, "Modelled")
        rr += 1
    mod_last = rr - 1
    R.cell(rr, 1, "Benchmarked (kgCO2e/m2 GIA x GIA)").font = BOLD
    for c in range(1, 23):
        R.cell(rr, c).fill = SECTION
    rr += 1
    bm_first = rr
    for key, label in BENCH_KEYS:
        R.cell(rr, 1, key)
        R.cell(rr, 2, label)
        R.cell(rr, 3, '=IF(COUNTIFS(%s,"%s",%s,1)>0,"Benchmark","")' % (KR("C"), key, KR("O")))
        s = lambda c: "SUMIFS(%s,%s,\"%s\",%s,1)*%s/1000" % (KR(c), KR("C"), key, KR("O"), GIA)   # noqa: E731
        R.cell(rr, mod_col["A1-A3"], "=" + s("D"))
        R.cell(rr, mod_col["A4"], "=" + s("E"))
        R.cell(rr, mod_col["A5.3"], "=" + s("F"))
        R.cell(rr, mod_col["B4"], "=" + s("G"))
        R.cell(rr, mod_col["C1"], '=IF(BM_C1="Include",1,0)*' + s("H"))
        R.cell(rr, mod_col["C2"], "=" + s("I"))
        R.cell(rr, mod_col["C3"], "=" + s("J"))
        R.cell(rr, mod_col["C4"], "=" + s("K"))
        R.cell(rr, mod_col["D"], "=" + s("L"))
        row_formulas(rr)
        tag(rr, key, "Benchmark")
        rr += 1
    bm_last = rr - 1
    R.cell(rr, 1, "Site (whole building, PS defaults)").font = BOLD
    for c in range(1, 23):
        R.cell(rr, c).fill = SECTION
    rr += 1
    site_first = rr
    for code, label, mod, f in [("A5.1", "Pre-construction demolition (existing GIA x rate)", "A5.1", "=%s*R_A51/1000" % GIA_DEMO),
                                ("A5.2", "Construction site activities (GIA x rate)", "A5.2", "=%s*R_A52/1000" % GIA),
                                ("C1", "Deconstruction / demolition at end of life (GIA x rate)", "C1", "=%s*R_C1/1000" % GIA)]:
        R.cell(rr, 1, code)
        R.cell(rr, 2, label)
        R.cell(rr, 3, "Site lump")
        R.cell(rr, mod_col[mod], f)
        row_formulas(rr)
        tag(rr, code, "Site")
        rr += 1
    site_last = rr - 1
    rr += 1
    sub = rr
    R.cell(rr, 2, "Subtotal").font = BOLD
    for c in range(4, 18):
        R.cell(rr, c, "=SUM({l}{a}:{l}{b})".format(l=col(c), a=mod_first, b=site_last))
    row_formulas(rr)
    rr += 1
    cont = rr
    R.cell(rr, 2, "Contingency (PS Table 10)")
    R.cell(rr, 3, "=CONT")
    R.cell(rr, 3).number_format = "0%"
    for c in range(4, 16):
        if c in (mod_col["A1-A5"], mod_col["A-C"]):
            continue
        R.cell(rr, c, "={l}{s}*CONT".format(l=col(c), s=sub))
    row_formulas(rr)
    rr += 1
    tot = rr
    R.cell(rr, 2, "TOTAL incl. contingency").font = BOLD
    for c in range(4, 18):
        R.cell(rr, c, "={l}{a}+{l}{b}".format(l=col(c), a=sub, b=cont))
    row_formulas(rr)
    for c in range(1, 20):
        R.cell(rr, c).fill = TOTAL
        R.cell(rr, c).font = BOLD

    # ---- B2 / B3: whole-building allowances (PS 5.2.2 - 5.2.3), one-off totals over the RSP, tCO2e, before contingency
    rr += 2
    R.cell(rr, 1, "B2 / B3 - whole-building allowances (PS 5.2.2 - 5.2.3), tCO2e over the RSP; not allocated to elements").font = BOLD
    for c in range(1, 23):
        R.cell(rr, c).fill = SECTION
    rr += 1
    RGc = lambda c: "$%s$%d:$%s$%d" % (c, mod_first, c, site_last)        # noqa: E731
    mep_up = 'SUMIFS(%s,%s,"Services/MEP")' % (RGc("I"), RGc("T"))
    keys = ["g", "p", "b2", "sh", "b2n", "b3n", "mep", "b3m", "b3", "b2c", "b3c"]
    b = dict((k_, rr + i) for i, k_ in enumerate(keys))                   # row of each line
    fx = dict(GIA=GIA, sub=sub, **b)
    lines = {
        "g": ("B2 option 1: GIA x rate", "={GIA}*B2_RATE/1000", '="GIA "&TEXT({GIA},"#,##0")&" m2 x "&B2_RATE&" kgCO2e/m2"'),
        "p": ("B2 option 2: share of A1-A5 (before contingency)", "=B2_PCT*I{sub}", '=TEXT(B2_PCT,"0%")&" x A1-A5 "&TEXT(I{sub},"#,##0")&" t"'),
        "b2": ("B2 maintenance - the greater of the two", "=MAX(C{g},C{p})", '=IF(C{g}>=C{p},"GIA x rate","% of A1-A5")'),
        "sh": ("Services (group 5) share of A1-A5", "=IF(I{sub}=0,0,MEP_UP/I{sub})", '="element group 5, before contingency"'),
        "b2n": ("B2 non-MEP part", "=C{b2}*(1-C{sh})", '="B2 x (1 - services share)"'),
        "b3n": ("B3 repair - non-MEP: share of B2 non-MEP", "=B3_PCT*C{b2n}", '=TEXT(B3_PCT,"0%")&" of the non-MEP part of B2"'),
        "mep": ("Services A1-A3 (initial installation)", "=MEP_A13", '="element group 5"'),
        "b3m": ("B3 repair - MEP: share of services A1-A3", "=B3_MEP*C{mep}", '=TEXT(B3_MEP,"0%")&" of services A1-A3 (CIBSE TM65)"'),
        "b3": ("B3 repair", "=C{b3n}+C{b3m}", '="25% non-MEP + 10% MEP rule (PS 5.2.3)"'),
        "b2c": ("B2 incl. contingency", "=C{b2}*(1+CONT)", '="kgCO2e/m2 GIA: "&TEXT(IF({GIA}>0,C{b2c}*1000/{GIA},0),"#,##0.0")'),
        "b3c": ("B3 incl. contingency", "=C{b3}*(1+CONT)", '="kgCO2e/m2 GIA: "&TEXT(IF({GIA}>0,C{b3c}*1000/{GIA},0),"#,##0.0")'),
    }
    for key in keys:
        lbl, f, note = lines[key]
        r_ = b[key]
        R.cell(r_, 2, lbl)
        R.cell(r_, 3, f.format(**fx).replace("MEP_UP", mep_up).replace("MEP_A13", 'SUMIFS(%s,%s,"Services/MEP")' % (RGc("D"), RGc("T"))))
        R.cell(r_, 4, note.format(**fx))
        R.cell(r_, 3).number_format = "0%" if key == "sh" else "#,##0.0"
        if key in ("b2", "b3"):
            R.cell(r_, 2).font = BOLD
            R.cell(r_, 3).font = BOLD
    rr = b["b3c"] + 1

    # headline block
    R["A4"] = "Headline (incl. contingency; building-level figures to 2 s.f. at RIBA 0-2, 3 s.f. from RIBA 3 - PS 6.1.4)"
    R["A4"].font = BOLD
    hl = [("Upfront carbon A1-A5", "=" + sf("R%d" % tot), "kgCO2e/m2 GIA"),
          ("Embodied carbon A-C incl. B2 / B3 and biogenic (excl. B6/B7)", "=" + sf("Y%d" % SUM_ROW["AC_ALL"]), "kgCO2e/m2 GIA"),
          ("A1-A5 total", "=" + sf("I%d" % tot), "tCO2e"),
          ("Modelled share of A1-A5 (before contingency)", "=IF(I%d=0,0,SUM(I%d:I%d)/I%d)" % (sub, mod_first, mod_last, sub), ""),
          ("Benchmarked share of A1-A5", "=IF(I%d=0,0,SUM(I%d:I%d)/I%d)" % (sub, bm_first, bm_last, sub), ""),
          ("Share of A1-A5 from [ASSUMED] quantities", "=IF(I%d=0,0,SUMIF(%s,\"Assumed\",%s)/1000/I%d)"
           % (sub, TR("Z"), TR("U"), sub), "replace as project information arrives"),
          ("Module D (memo, not in totals)", "=IF(%s>0,P%d*1000/%s,0)" % (GIA, sub, GIA), "kgCO2e/m2 GIA"),
          ("QA flags open", "=%s!C4" % QS, "see %s" % QS)]
    for i, (lbl, f, unit) in enumerate(hl):
        R.cell(5 + i, 1, lbl)
        c = R.cell(5 + i, 3, f)
        c.font = BOLD
        c.number_format = "0%" if "share" in lbl.lower() else "#,##0"
        R.cell(5 + i, 4, unit)
    R.freeze_panes = "D14"

    # ---------------------------------------------------------------- Summary block (X:Z) for Dashboard / Compare
    RG = lambda c: "$%s$%d:$%s$%d" % (c, mod_first, c, site_last)        # noqa: E731
    G1 = "(1+CONT)*1000/%s" % GIA
    I_ = "I%d" % sub
    BB = "($C$%d+$C$%d)" % (b["b2"], b["b3"])                              # B2 + B3 before contingency (tCO2e)
    O_ = "(O%d+%s)" % (sub, BB)                                             # A-C incl. B2 / B3, before contingency
    BIO = "(Y%d+Y%d)" % (SUM_ROW["BIO_STORED"], SUM_ROW["BIO_RELEASED"])     # nets to zero unless products stay stored
    ext = lambda c, kind=None: ('SUMIFS(%s,%s,"External works"%s)' % (RG(c), RG("T"), ',%s,"%s"' % (RG("U"), kind) if kind else ""))  # noqa: E731
    kind = lambda c, k_: 'SUMIFS(%s,%s,"%s")' % (RG(c), RG("U"), k_)       # noqa: E731
    asm = 'SUMIF(%s,"Assumed",%s)/1000' % (TR("Z"), TR("U"))
    bm_up, bm_ac = kind("I", "Benchmark"), kind("O", "Benchmark")
    bm_up8, bm_ac8 = "(%s-%s)" % (bm_up, ext("I", "Benchmark")), "(%s-%s)" % (bm_ac, ext("O", "Benchmark"))
    # Cat A part of groups 3 / 4 / 5 for the UKNZCBS offices shell & core comparison (share inputs on Project)
    cat_a = lambda c, k_=None: "+".join('(1-%s)*SUMIFS(%s,%s,"%s"%s)' % (nm_, RG(c), RG("T"), g_, ',%s,"%s"' % (RG("U"), k_) if k_ else "")  # noqa: E731
                                        for nm_, g_ in (("SC_FIN", "Internal finishes"), ("SC_FFE", "FFE"), ("SC_MEP", "Services/MEP")))
    code = lambda c, cd: 'SUMIFS(%s,%s,"%s")' % (RG(c), RG("A"), cd)       # noqa: E731
    scn = lambda up, ac: "(%s*SCN_FUP+(%s-%s)*SCN_FBC)" % (up, ac, up)      # noqa: E731
    letis = '{"A++","A+","A","B","C","D","E","F","G","worse than G"}'
    tsub = "%s,\"%s\",%s,USE,%s,\"<\"&%s"      # band = first band whose limit the project does not exceed
    TGc = ctx["TGc"]
    F = {
        "GIA": "=%s" % GIA, "UP_T": "=I%d" % tot, "AC_T": "=O%d+C%d+C%d" % (tot, b["b2c"], b["b3c"]),
        "UP_ALL": "=IF(%s>0,%s*%s,0)" % (GIA, I_, G1),
        "UP_X8": "=IF(%s>0,(%s-%s)*%s,0)" % (GIA, I_, ext("I"), G1),
        "BC_ALL": "=IF(%s>0,(%s-%s)*%s,0)" % (GIA, O_, I_, G1),
        # A-C includes biogenic carbon (PS glossary): stored in A1-A3 (negative) and released at end of life; no contingency on it
        "BIO_STORED": "=IF(%s>0,Q%d*1000/%s,0)" % (GIA, sub, GIA),
        "BIO_RELEASED": "=IF(%s>0,SUM(%s)/%s,0)" % (GIA, NR("AF"), GIA),
        "AC_ALL": "=IF(%s>0,%s*%s,0)+%s" % (GIA, O_, G1, BIO),
        "AC_X8": "=IF(%s>0,(%s-%s)*%s,0)+%s" % (GIA, O_, ext("O"), G1, BIO),
        "M_A13": "=IF(%s>0,D%d*1000/%s,0)" % (GIA, tot, GIA), "M_A4": "=IF(%s>0,E%d*1000/%s,0)" % (GIA, tot, GIA),
        "M_A5": "=IF(%s>0,SUM(F%d:H%d)*1000/%s,0)" % (GIA, tot, tot, GIA), "M_B4": "=IF(%s>0,J%d*1000/%s,0)" % (GIA, tot, GIA),
        "M_C": "=IF(%s>0,SUM(K%d:N%d)*1000/%s,0)" % (GIA, tot, tot, GIA), "M_D": "=IF(%s>0,P%d*1000/%s,0)" % (GIA, sub, GIA),
        "M_B2": "=IF(%s>0,C%d*1000/%s,0)" % (GIA, b["b2c"], GIA), "M_B3": "=IF(%s>0,C%d*1000/%s,0)" % (GIA, b["b3c"], GIA),
        "B2_BASIS": "=D%d" % b["b2"],
        "E_CONT": "=IF(%s>0,I%d*1000/%s,0)" % (GIA, cont, GIA),
        "B_ASM": "=IF(%s>0,%s*1000/%s,0)" % (GIA, asm, GIA),
        "B_MOD": "=IF(%s>0,%s*1000/%s,0)-Y%d" % (GIA, kind("I", "Modelled"), GIA, SUM_ROW["B_ASM"]),
        "B_BM": "=IF(%s>0,%s*1000/%s,0)" % (GIA, bm_up, GIA), "B_SITE": "=IF(%s>0,%s*1000/%s,0)" % (GIA, kind("I", "Site"), GIA),
        "B_CONT": "=Y%d" % SUM_ROW["E_CONT"],
        "SH_ASM": "=IF(%s=0,0,%s/%s)" % (I_, asm, I_),
        "STRUCT": '=IF(%s>0,(SUMIFS(%s,%s,"Y")+%s)*1000/%s,0)' % (GIA, RG("I"), RG("V"), code("I", "A5.2"), GIA),
        "STRUCT_AC": '=IF(%s>0,(SUMIFS(%s,%s,"Y")+%s+%s)*1000/%s,0)' % (GIA, RG("O"), RG("V"), code("O", "A5.2"), code("O", "C1"), GIA),
        "UP_SC": "=IF(%s>0,(%s-%s-(%s))*%s,0)" % (GIA, I_, ext("I"), cat_a("I"), G1),
        "MOD_UP_SC": "=IF(%s>0,(%s-%s-%s-(%s))*%s,0)" % (GIA, I_, bm_up, ext("I", "Modelled"), cat_a("I", "Modelled"), G1),
        "BM_UP_SC": "=IF(%s>0,(%s-%s-(%s))*%s,0)" % (GIA, bm_up, ext("I", "Benchmark"), cat_a("I", "Benchmark"), G1),
        "BMS_UP_SC": "=Y%d*SCN_FUP" % SUM_ROW["BM_UP_SC"],
        "UP_WB": '=IF(USE<>"Office",Y{x},IF(ISNUMBER(CATB),Y{x}+CATB*(1+CONT),"not assessed"))'.format(x=SUM_ROW["UP_X8"]),
        "UK_LIM_SC": "=UK_LIM_SC",
        "MOD_UP_ALL": "=IF(%s>0,(%s-%s)*%s,0)" % (GIA, I_, bm_up, G1),
        "MOD_UP_X8": "=IF(%s>0,(%s-%s-%s)*%s,0)" % (GIA, I_, bm_up, ext("I", "Modelled"), G1),
        "MOD_AC_ALL": "=IF(%s>0,(%s-%s)*%s,0)+%s" % (GIA, O_, bm_ac, G1, BIO),        # B2 / B3 sit with modelled + site
        "MOD_AC_X8": "=IF(%s>0,(%s-%s-%s)*%s,0)+%s" % (GIA, O_, bm_ac, ext("O", "Modelled"), G1, BIO),
        "BM_UP_ALL": "=IF(%s>0,%s*%s,0)" % (GIA, bm_up, G1), "BM_UP_X8": "=IF(%s>0,%s*%s,0)" % (GIA, bm_up8, G1),
        "BM_AC_ALL": "=IF(%s>0,%s*%s,0)" % (GIA, bm_ac, G1), "BM_AC_X8": "=IF(%s>0,%s*%s,0)" % (GIA, bm_ac8, G1),
        "BMS_UP_ALL": "=IF(%s>0,%s*SCN_FUP*%s,0)" % (GIA, bm_up, G1), "BMS_UP_X8": "=IF(%s>0,%s*SCN_FUP*%s,0)" % (GIA, bm_up8, G1),
        "BMS_AC_ALL": "=IF(%s>0,%s*%s,0)" % (GIA, scn(bm_up, bm_ac), G1), "BMS_AC_X8": "=IF(%s>0,%s*%s,0)" % (GIA, scn(bm_up8, bm_ac8), G1),
        "SCN_UP_ALL": "=Y%d+Y%d" % (SUM_ROW["MOD_UP_ALL"], SUM_ROW["BMS_UP_ALL"]),
        "SCN_AC_ALL": "=Y%d+Y%d" % (SUM_ROW["MOD_AC_ALL"], SUM_ROW["BMS_AC_ALL"]),
        "UK_LIM": "=UK_LIM",
        "LETI_UP": '=IF(COUNTIFS(%s,"LETI upfront",%s,USE)=0,"n/a",INDEX(%s,COUNTIFS(%s)+1))'
                   % (TGc("A"), TGc("C"), letis, tsub % (TGc("A"), "LETI upfront", TGc("C"), TGc("F"), "Y%d" % SUM_ROW["UP_X8"])),
        "LETI_LC": '=IF(COUNTIFS(%s,"LETI life cycle",%s,USE)=0,"n/a",INDEX(%s,COUNTIFS(%s)+1))'
                   % (TGc("A"), TGc("C"), letis, tsub % (TGc("A"), "LETI life cycle", TGc("C"), TGc("F"), "Y%d" % SUM_ROW["AC_X8"])),
        "QA_FLAGS": "=%s!C4" % QS, "NAME": "=O%d_NAME" % k,
    }
    for i, g in enumerate(GROUPS):
        F["E_%d" % i] = '=IF(%s>0,SUMIFS(%s,%s,"%s")*1000/%s,0)' % (GIA, RG("I"), RG("T"), g, GIA)
    # ---- B6 / B7 (PS 5.3 / 5.4, App. G): per option; no contingency; indicative route = UKNZCBS EUI limit x GIA, all electric
    eng, ind = 'LEFT(OP_ROUTE,8)="Engineer"', 'LEFT(OP_ROUTE,10)="Indicative"'
    E_ = "IF(%s,N(O%d_E_ELEC),IF(%s,IF(N(O%d_EUI_IN)>0,O%d_EUI_IN,UK_EUI)*%s,0))" % (eng, k, ind, k, k, GIA)
    G_, DH_ = "IF(%s,N(O%d_E_GAS),0)" % (eng, k), "IF(%s,N(O%d_E_DH),0)" % (eng, k)
    PVS_, PVX_ = "IF(%s,N(O%d_E_PVS),0)" % (eng, k), "IF(%s,N(O%d_E_PVX),0)" % (eng, k)
    W_ = "Y%d" % SUM_ROW["W_M3"]                                            # engineer's water, else the occupancy budget
    y = lambda key: "Y%d" % SUM_ROW[key]                                     # noqa: E731
    F.update({
        "EUI": '=IF(OP_ROUTE="%s","not assessed",IF(%s>0,(%s+%s+%s+%s)/%s,0))' % (OP_ROUTES[2], GIA, E_, G_, DH_, PVS_, GIA),
        "EUI_LIM": "=UK_EUI",
        "B6_ND": "=IF(%s>0,RSP*(%s*F_E_ND+%s*F_G_ND+%s*F_DH)/%s,0)" % (GIA, E_, G_, DH_, GIA),
        "B6_D": "=IF(%s>0,RSP*(%s*F_E_D+%s*F_G_D+%s*F_DH)/%s,0)" % (GIA, E_, G_, DH_, GIA),
        "B7_ND": "=IF(%s>0,RSP*%s*F_W_ND/%s,0)" % (GIA, W_, GIA),
        "B7_D": "=IF(%s>0,RSP*%s*F_W_D/%s,0)" % (GIA, W_, GIA),
        "D2_PV": "=IF(%s>0,-RSP*%s*FES_AVG/%s,0)" % (GIA, PVX_, GIA),
        "WLC_ND": '=IF(OP_ROUTE="%s","not assessed",%s+%s+%s)' % (OP_ROUTES[2], y("AC_ALL"), y("B6_ND"), y("B7_ND")),
        "WLC_D": '=IF(OP_ROUTE="%s","not assessed",%s+%s+%s)' % (OP_ROUTES[2], y("AC_DEC"), y("B6_D"), y("B7_D")),
        "W_M3": "=IF(N(O%d_W)>0,O%d_W,W_BUDGET)" % (k, k),
        "AC_DEC": "=IF(%s>0,(I%d+0.5*(J%d+%s)+0.5*(K%d+L%d)+M%d+N%d)*%s,0)+%s" % (GIA, sub, sub, BB, sub, sub, sub, sub, G1, BIO),
        "D_DEC": "=IF(%s>0,0.5*P%d*1000/%s,0)" % (GIA, sub, GIA),
        "OP_STATUS": ('=IF(OP_ROUTE="%s","Energy not assessed - embodied carbon only",IF(%s,"Energy: engineer prediction",'
                      'IF(N(O%d_EUI_IN)>0,"Energy: indicative (EUI estimate)","Energy: indicative (UKNZCBS limit)")))'
                      '&IF(N(O%d_W)>0," | water: engineer",IF(W_BUDGET>0," | water: indicative budget (occupancy x RIBA 2025)"," | water: not assessed"))'
                      % (OP_ROUTES[2], eng, k, k)),
    })
    R.cell(4, 24, "Summary (read by Dashboard and Compare - do not move)").font = BOLD
    R.column_dimensions["X"].width = 12
    R.column_dimensions["Y"].width = 12
    R.column_dimensions["Z"].width = 48
    for key, lbl, unit in SUMMARY:
        rr_ = SUM_ROW[key]
        R.cell(rr_, 24, key)
        c = R.cell(rr_, 25, F[key])
        c.number_format = "0%" if key.startswith("SH_") else "#,##0.0"
        R.cell(rr_, 26, "%s%s" % (lbl, " (%s)" % unit if unit else ""))
    # NB: the Summary block sits in rows 5..~65 of columns X:Z, right of the element table (A:V)

    # ---------------------------------------------------------------- QA
    Q = wb[QS]
    title(Q, "%s - QA checks" % o["name"], "FLAG = must be resolved or justified before issue. INFO = read and confirm.")
    Q["A4"] = "Flags open:"
    Q["A4"].font = BOLD
    head(Q, 6, ["#", "Check", "Result", "Status", "What to do"], [5, 60, 14, 9, 90])
    cw_ref = params.get("cw_reference_a1a5")
    cw_sum = ('SUMPRODUCT((%s=CW_BU)*(%s+%s+%s+%s))' % (NR("C"), NR("Q"), NR("T"), NR("U"), NR("W")))
    cw_qty = 'SUMIFS(%s,%s,CW_BU,%s,1)' % (TR("D"), TR("K"), TR("Q"))
    checks = [
        ("GIA entered and > 0", "=%s" % GIA, '=IF(N(C{r})>0,"OK","FLAG")', "Enter GIA for this option on Project (from the area schedule or the take-off)."),
        ("Take-off rows whose build-up is not in the library (not counted)", '=COUNTIF(%s,"N")' % TR("M"),
         '=IF(C{r}=0,"OK","FLAG")', "Add the build-up to the library (owner) or change the code on the row. These quantities are currently missing from the result."),
        ("Take-off rows with a unit that differs from the build-up's declared unit (not counted)",
         '=COUNTIF(%s,"Unit mismatch")' % TR("O"), '=IF(C{r}=0,"OK","FLAG")', "Fix the export unit or the build-up."),
        ("Take-off rows using an incomplete build-up (layer missing a factor)",
         '=SUMPRODUCT((%s<>"")*(%s<>"OK")*(%s="Y"))' % (TR("K"), TR("P"), TR("M")),
         '=IF(C{r}=0,"OK","FLAG")', "Layers without a factor count as zero. Add the factor in the library."),
        ("Materials in use with no A1-A3 factor", '=SUMPRODUCT((%s>0)*(%s=""))' % (MO("AG"), MR_("F")),
         '=IF(C{r}=0,"OK","FLAG")', "See Materials column F for rows with project mass on %s." % MT),
        ("Materials in use with status CHECK", '=SUMPRODUCT((%s="Y")*(%s="CHECK"))' % (MO("AR"), MR_("P")),
         '=IF(C{r}=0,"OK","INFO")', "Usable, but provenance or value still to confirm (library owner)."),
        ("DRAFT build-ups in use", '=SUMPRODUCT((%s=1)*(COUNTIFS(%s,%s,%s,"DRAFT")>0))' % (TR("Q"), BA, TR("K"), BG),
         '=IF(C{r}=0,"OK","INFO")', "Expected before library v1.0. Only RELEASED build-ups should be used for issue."),
        ("Layers dropped because their element group is benchmarked", '=SUMPRODUCT((%s=1)*(%s=1))' % (NR("E"), NR("I")),
         '=IF(C{r}=0,"OK","INFO")', "Avoids double counting (BR-14 one method per element). Check the benchmark covers them."),
        ("Take-off rows excluded by include = N", '=COUNTIF(%s,"N")' % TR("G"), '=IF(C{r}=0,"OK","INFO")', "Confirm the exclusions are intended."),
        ("Opaque facade rows routed to the curtain wall build-up", '=SUM(%s)' % TR("J"), '=IF(C{r}=0,"OK","INFO")',
         "Facade route on Project. Whole envelope is reported in 2.5.2."),
        ("Curtain wall A1-A5 per m2 facade vs reference (kgCO2e/m2)",
         ('=IF(%s=0,"",%s/%s)' % (cw_qty, cw_sum, cw_qty)),
         ('=IF(C{r}="","INFO",IF(ABS(C{r}/%s-1)<=0.05,"OK","FLAG"))' % cw_ref) if cw_ref else '="INFO"',
         ("Reference %s kgCO2e/m2 FSA (ETPP-28 EWS-04A, CWCT). Within 5%% = OK. Differences come from A4 vehicles and Table 18 site waste." % cw_ref)
         if cw_ref else "No reference value configured."),
        ("Take-off rows resting on office assumptions ([ASSUMED] in notes; Takeoff column Z)", '=COUNTIF(%s,"Assumed")' % TR("Z"),
         '=IF(C{r}=0,"OK","INFO")',
         "Expected at RIBA 0-2. Replace each with drawings, model quantities or supplied figures (engineer / user) as they arrive: "
         "set the take-off CONFIG source to 'supplied' or re-run with the new drawing, then regenerate. Filter Takeoff column Z = Assumed for the list."),
        ("Share of A1-A5 resting on assumed quantities (before contingency)",
         '=IFERROR(%s/%s!I%d,0)' % (asm, RS, sub), '="INFO"',
         "The part of the result that moves when assumptions are replaced."),
        ("Contingency at RIBA 0-2 at least 15%", "=CONT", '=IF(AND(STAGE<=2,CONT<0.15),"FLAG","OK")', "PS Table 10."),
        ("Scope elements with neither take-off nor benchmark (GAPS)", None, '=IF(C{r}=0,"OK","FLAG")',
         "RICS PS 4.4.1: all in-scope elements are included at every stage. Add a quantity, a benchmark, or mark the element not applicable in the coverage table below."),
    ]
    r = 7
    gap_row = None
    for i, (lbl, res, st, todo) in enumerate(checks, start=1):
        Q.cell(r, 1, i)
        Q.cell(r, 2, lbl)
        if res is not None:
            Q.cell(r, 3, res)
        else:
            gap_row = r
        Q.cell(r, 4, st.format(r=r))
        Q.cell(r, 5, todo).alignment = WRAP
        Q.cell(r, 3).number_format = "0%" if lbl.startswith("Share") else "#,##0.0"
        r += 1
    q_last = r - 1
    Q["C4"] = '=COUNTIF(D7:D%d,"FLAG")' % q_last
    Q["C4"].font = Font(bold=True, color="C00000")
    r += 1
    Q.cell(r, 1, "Scope coverage").font = BOLD
    r += 1
    head(Q, r, ["Code", "Element", "Expected? (Y/N)", "Modelled (kgCO2e)", "Benchmarked", "Status", "Why (from the take-off)"], None, 30)
    r += 1
    cov_first = r
    for code, label, exp in COVERAGE:
        Q.cell(r, 1, code)
        Q.cell(r, 2, label)
        exp, why = coverage_expected(code, exp, header)
        Q.cell(r, 3, '=IF(%s>0,"Y","N")' % GIA_DEMO if exp == "DEMO" else exp).fill = INPUT
        Q.cell(r, 7, why)
        crit = '"%s"' % code if code.count(".") >= 1 and code not in ("2.4",) else '"%s*"' % code
        if code in ("2.4",):
            crit = '"2.4*"'
        Q.cell(r, 4, "=SUMIFS(%s,%s,%s)+SUMIFS(%s,%s,%s)" % (NR("Q"), NR("G"), crit, NR("T"), NR("G"), crit))
        Q.cell(r, 4).number_format = "#,##0"
        key = "2.7/2.8" if code in ("2.7", "2.8") else (code if code in ("3", "4", "5", "8") else "")
        Q.cell(r, 5, '=IF(COUNTIFS(%s,"%s",%s,1)>0,"Y","")' % (KR("C"), key, KR("O")) if key else "")
        Q.cell(r, 6, '=IF(C{r}<>"Y","n/a",IF(D{r}>0,"Modelled",IF(E{r}="Y","Benchmark","GAP")))'.format(r=r))
        r += 1
    cov_last = r - 1
    Q.column_dimensions["G"].width = 80
    for rr_, crit_ in method_rows:
        cs = 'INDEX(%s!$F${a}:$F${b},MATCH(A{r}&"",%s!$A${a}:$A${b},0))' % (QS, QS)
        R.cell(rr_, 3, ('=IF(COUNTIFS({g},{c},{j},1)>0,"Modelled",IFERROR(IF(' + cs + '="Benchmark","see benchmark",'
                        + cs + '),""))').format(r=rr_, g=NR("G"), c=crit_, j=NR("J"), a=cov_first, b=cov_last))
    R.conditional_formatting.add("C%d:C%d" % (mod_first, mod_last), CellIsRule(operator="equal", formula=['"GAP"'], fill=RED))
    Q.cell(gap_row, 3, '=COUNTIF(F%d:F%d,"GAP")' % (cov_first, cov_last))
    for rng_, rules in (("D7:D%d" % q_last, [("OK", GREEN), ("INFO", AMBER), ("FLAG", RED)]),
                        ("F%d:F%d" % (cov_first, cov_last), [("Modelled", GREEN), ("Benchmark", GREEN), ("GAP", RED)])):
        for txt, fill in rules:
            Q.conditional_formatting.add(rng_, CellIsRule(operator="equal", formula=['"%s"' % txt], fill=fill))
    T.conditional_formatting.add("Z%d:Z%d" % (t_first, t_last), CellIsRule(operator="equal", formula=['"Assumed"'], fill=AMBER))
    T.conditional_formatting.add("Y%d:Y%d" % (t_first, t_last),
                                 FormulaRule(formula=['AND(Y%d<>"",Y%d<>"OK")' % (t_first, t_first)], fill=AMBER))
    return dict(TK=TK, LN=LN, MT=MT, RS=RS, QS=QS, sub=sub, cont=cont, tot=tot, t_first=t_first, t_last=t_last)


def apply_font(wb, font):
    """Use one font family across the workbook (sizes, weights and colours kept). The workbook stores only the name:
    anyone opening it without that font installed sees Excel's substitute."""
    if not font:
        return
    from copy import copy
    from openpyxl.utils.indexed_list import IndexedList
    swapped = {}
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if not c.has_style:
                    continue
                fid = c._style.fontId
                if fid not in swapped:
                    f = copy(wb._fonts[fid])
                    f.name, f.scheme = font, None       # a 'minor' / 'major' scheme makes Excel use the theme font instead
                    swapped[fid] = f
                c.font = swapped[fid]
    try:        # default font (unstyled cells, the Normal style)
        f0 = copy(wb._fonts[0])
        f0.name, f0.scheme = font, None
        wb._fonts = IndexedList([f0] + list(wb._fonts)[1:])
    except Exception:
        pass


def main(argv):
    if len(argv) < 2:
        print(__doc__)
        return 2
    cfg = json.load(open(argv[1], encoding="utf-8"))
    wb = build(cfg)
    out = cfg["out"]
    wb.save(out)
    print("Wrote", out)
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
