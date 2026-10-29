"""Reconciliation (wlca_reconcile.py) on an INVENTED building, library and model.
Drafted by an AI agent (Claude) - review before trusting."""
import json

import pytest

import wlca_reconcile as wr
import wlca_state as ws

B = {"pile": "PILE", "cap_raft": "CAP", "core_wall": "CORE", "column": "COLD", "edge_beam": "EB", "soffit_finish": "SOF",
     "louvre": "LOUV", "door_ext": "DOOR", "facade_glazed": "CW", "facade_opaque": "OPQ", "upper_floor": "FLAT",
     "lowest_slab": "GF", "liner_wall": "LINER"}
LIB = {"FLAT": {"unit": "m2", "conc": 0.30}, "GF": {"unit": "m2", "conc": 0.35}, "COLD": {"unit": "m", "conc": 0.5},
       "EB": {"unit": "m", "conc": 0.14}, "SLAB-M3": {"unit": "m3", "conc": 1.0}, "COL-M3": {"unit": "m3", "conc": 1.0},
       "RET": {"unit": "m2", "conc": 0.3}}


def r(eid, code, bu, q, unit="m2", method="Derived", notes=""):
    return {"project_id": "T", "element_id": eid, "nrm_code": code, "build_up_code": bu, "quantity": str(q), "unit": unit,
            "method": method, "stage": "2", "model_version": "A", "author": "XX", "source_layer": "M", "include": "Y",
            "notes": notes}


def derived():
    return [r("MS-1", "2.2.1", "FLAT", 2000), r("MS-2", "2.5.1", "OPQ", 600), r("MS-3", "2.5.2", "CW", 400),
            r("MS-5", "2.1.1", "CORE", 800), r("MS-6", "2.1.1", "COLD", 100, "m"), r("MS-7", "2.1.2", "EB", 500, "m"),
            r("MS-8", "1.2.1", "GF", 1000), r("MS-9", "1.2.3", "RET", 520), r("MS-10", "1.2.3", "RET", 99, method="Supplied")]


def geom():
    # B1 (4 m), L0..L2 (4 m each); schedule GEA differs from the massing at L2
    return {"schedule_reconciliation": [
        {"level": "B1", "aod": -4.0, "ftf": 4.0, "schedule_gea": 1000, "massing_gea": 1000},
        {"level": "0", "aod": 0.0, "ftf": 4.0, "schedule_gea": 1000, "massing_gea": 1010},
        {"level": "1", "aod": 4.0, "ftf": 4.0, "schedule_gea": 1000, "massing_gea": 990},
        {"level": "2", "aod": 8.0, "ftf": 4.0, "schedule_gea": 1000, "massing_gea": 900}],
        "basement": [{"name": "B1", "z": -4.0, "ftf": 4.0}],
        "storeys": [{"name": "0", "z": 0.0, "height": 4.0}, {"name": "1", "z": 4.0, "height": 4.0},
                    {"name": "2", "z": 8.0, "height": 4.0}],
        "structure": {"typical_level": "1", "basement_depth_m": 4.0},
        "basement_wall_perimeter": {"m": 130.0}}


def slab(top, area, t=0.3):
    return [top - t, top, area, area * t, 0.0, 0]


def column(z0, z1, a=0.25):
    return [z0, z1, a, a * (z1 - z0), 2.0 * (z1 - z0), 1]


def measured():
    slabs = [slab(0.0, 950), slab(4.0, 940), slab(7.9, 930), slab(12.0, 300)]          # 12.0: the roof, not a level
    cols = [column(0.3, 4.0) for _ in range(10)] + [column(4.3, 8.0) for _ in range(12)]
    core = [[-4.0, 14.0, 50.0, 300.0, 400.0, 1]]                                          # core runs 2 m above the roof
    ret = [[-4.5, 0.0, 10.0, 120.0, 135.0 * 4.5, 1]]
    st = lambda **k: dict({"n": 1, "open": 0, "vol": 0.0, "vol_v": 0.0, "vol_h": 0.0, "plan": 0.0, "face": 0.0, "outside": 0}, **k)  # noqa: E731
    return {"layers": {
        "S::Slabs": st(vol=sum(o[3] for o in slabs), plan=sum(o[2] for o in slabs), objects=slabs, open=2),
        "S::Frame": st(vol=sum(o[3] for o in cols), vol_v=sum(o[3] for o in cols), objects=cols),
        "S::Core": st(vol=300.0, objects=core),
        "S::Ret": st(vol=120.0, face=607.5, objects=ret),
        "F::Glass": st(face=300.0, plan=10.0),
        "F::Brick": st(face=900.0, plan=400.0, outside=3, objects=[[0, 0.1, 400.0, 40.0, 20.0, 0]] + [[0, 0.065, 0.02, 0.001, 0.02, 0]] * 50),
        "F::Mortar": st(face=10.0, plan=900.0)},                                       # no records: not checked
        "duplicates": [["S::Slabs", "S::Slabs", 1], ["X::Other", "X::Other2", 5]], "units": "m"}


RULES = [
    {"layer": "S::Slabs", "rics": "2.2.1", "measure": "volume", "build_up": "SLAB-M3", "factor": 1.0, "replaces": "floors"},
    {"layer": "S::Frame", "rics": "2.1.1", "measure": "volume_v", "build_up": "COL-M3", "factor": 1.0, "replaces": "columns"},
    {"layer": "S::Core", "rics": "2.1.1", "measure": "core_ref", "build_up": "", "factor": 1.0, "replaces": "none"},
    {"layer": "S::Ret", "rics": "1.2.3", "measure": "volume", "build_up": "RET", "factor": 1.0, "replaces": "basement_walls"},
    {"layer": "F::Glass", "rics": "2.6", "measure": "share", "build_up": "WIN", "factor": 1.25, "replaces": "facade"},
    {"layer": "F::Brick", "rics": "2.5.1", "measure": "share", "build_up": "WALL", "factor": 1.0, "replaces": "facade"},
    {"layer": "F::Mortar", "rics": "2.5.1", "measure": "share", "build_up": "WALL", "factor": 1.0, "replaces": "facade"},
]


def run(option=None, rules=RULES, mixed=0.4):
    return dict((x["key"], x) for x in wr.reconcile(derived(), geom(), {"columns": {"count_per_floor": 12}}, rules, measured(),
                                                     LIB, B, option or {}, mixed))


def test_levels_against_the_schedule():
    R = run()
    assert R["gea:2"]["values"] == {"massing": 900.0, "schedule": 1000.0, "model": 930.0}
    assert R["gea:2"]["flag"] and R["gea:2"]["ref"] == "schedule"                    # massing -10%
    assert R["gea:1"]["flag"] and R["gea:1"]["diff"] == {"massing": -0.01, "model": -0.06}   # model 6% under
    assert not R["gea:0"]["flag"] and R["gea:B1"]["values"]["model"] is None          # no slab at B1
    assert R["gea:total"]["values"]["model"] == 950 + 940 + 930                        # the roof slab is not a level
    assert "300 m2" in R["gea:total"]["note"]


def test_totals():
    R = run()
    assert R["slabs_m3"]["values"] == {"massing": 600.0, "model": pytest.approx(0.3 * 3120)}
    assert R["slabs_t"]["values"]["massing"] == pytest.approx(0.30) and R["slabs_t"]["values"]["model"] == pytest.approx(0.30)
    assert R["columns_m3"]["values"]["massing"] == 50.0
    assert R["columns_n"]["values"] == {"plan": 12.0, "model": 12.0} and not R["columns_n"]["flag"]
    assert R["facade"]["values"] == {"massing": 1000.0, "model": 300 * 1.25 + 900 + 10}
    assert R["glazing"]["values"]["massing"] == 0.4
    assert R["glazing"]["values"]["model"] == pytest.approx(375 / 1285, abs=1e-3)
    assert R["basement_len"]["values"] == {"massing": 130.0, "model": 135.0}
    assert "beams_m3" not in R                                                          # no modelled beams


def test_choices_and_scale_factors():
    R = run()
    assert R["core_height"]["values"] == {"plan": 16.0, "model": 18.0} and R["core_height"]["flag"]
    assert R["core_height"]["choice"]["value"] == "storeys"
    assert wr.scale_factors(list(R.values())) == {}
    R = run({"reconcile": {"core_height": "model", "basement_depth": "model"}})
    f = wr.scale_factors(list(R.values()))
    assert f["core_wall"] == pytest.approx(18 / 16) and f["basement_wall"] == pytest.approx(4.5 / 4.0)
    rows, notes = ws.apply_reconcile(derived(), list(R.values()), {}, B)
    q = dict((x["element_id"], float(x["quantity"])) for x in rows)
    assert q["MS-5"] == pytest.approx(900.0) and q["MS-9"] == pytest.approx(585.0)
    assert q["MS-10"] == 99.0 and q["MS-6"] == 100.0                                  # engineer's figure, columns untouched
    assert len(notes) == 2


def test_model_checks():
    R = run()
    assert "dup:S::Slabs|S::Slabs" in R and "dup:X::Other|X::Other2" not in R           # only layers in the table
    assert "open:S::Slabs" in R and "out:F::Brick" in R
    assert "flat:F::Brick" in R and "400 m2" in R["flat:F::Brick"]["note"]             # the plate, not the bricks
    assert "flat:F::Glass" not in R and "flat:F::Mortar" not in R


def test_without_the_model():
    R = run(rules=[])
    assert set(k for k in R) == {"gea:B1", "gea:0", "gea:1", "gea:2", "gea:total"}
    assert "model" not in R["gea:1"]["values"]
    R = run({"use_modelled": False})
    assert "slabs_m3" not in R


def test_state_choice_and_stale(tmp_path):
    st = ws.default_state()
    st["options"] = [ws.default_option("A")]
    rows = list(run().values())
    p = tmp_path / "A_reconciliation.json"
    ws.store_reconciliation(st, st["options"][0], rows, str(p))
    assert json.loads(p.read_text())["rows"][0]["key"] == "gea:B1"
    assert ws.reconcile_stale(st) == []
    ws.set_reconcile_choice(st, 0, "core_height", "model")
    assert ws.reconcile_stale(st) == ["A"]
    ws.set_reconcile_choice(st, 0, "core_height", "storeys")
    assert ws.reconcile_stale(st) == []
    with pytest.raises(ValueError):
        ws.set_reconcile_choice(st, 0, "core_height", "guess")
    assert ws.recon_path("C:/x/A_takeoff.csv") == "C:/x/A_takeoff_reconciliation.json"


def test_glazing_mix():
    o = {"facade": [{"build_up": "CW", "kind": ws.FACADE_UNITISED, "share": 0.5},
                    {"build_up": "PUN", "kind": ws.FACADE_PUNCHED, "share": 0.5, "wwr": 0.4}]}
    assert ws.glazing_mix(o, {}) == pytest.approx(0.7)


def test_workbook_sheet(tmp_path):
    from openpyxl import Workbook
    import wlca_calc_template as ct
    rows = list(run({"reconcile": {"core_height": "model"}}).values())
    p = tmp_path / "A_reconciliation.json"
    ws.store_reconciliation({"options": []}, {"name": "A", "reconcile": {"core_height": "model"}}, rows, str(p))
    wb = Workbook()
    wb.create_sheet("O1_Recon")
    ct.build_recon(wb, {"k": 1, "name": "A", "recon": ct.read_recon(str(p))})
    S = wb["O1_Recon"]
    cells = [[c.value for c in r] for r in S.iter_rows(min_row=5)]
    assert cells[0][0] == "GEA per level"
    core = next(r for r in cells if r[0] == "Core walls: height")
    assert core[3] == 16.0 and core[5] == 18.0 and core[7] == "CHECK" and core[8] == "Used: Core in the model"
    assert any(r[0] == "Duplicate pieces" and r[7] == "CHECK" for r in cells)
    assert ct.read_recon(None) is None
