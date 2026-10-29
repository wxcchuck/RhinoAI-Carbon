"""Dialog state: checks, facade mix and config builders (no Rhino). Drafted by an AI agent (Claude) - review before trusting."""
import json
import os

import pytest

import wlca_state as ws

OFFICE = {"facade_types": [
    {"code": "CW", "label": "Unitised CW", "kind": ws.FACADE_UNITISED},
    {"code": "SFS", "label": "SFS + brick", "kind": ws.FACADE_PUNCHED, "glazing_build_up": "WIN"},
    {"id": "SFS-T", "code": "SFS", "label": "SFS + brick, timber windows", "kind": ws.FACADE_PUNCHED, "glazing_build_up": "WIN-T"},
    {"code": "OLD", "label": "Punched, no window build-up", "kind": ws.FACADE_PUNCHED},
]}


def _rows():
    base = dict(project_id="T", method="Derived", stage="2", model_version="v", author="XX", source_layer="L", include="Y", unit="m2")
    return [dict(base, element_id="MS-0008", nrm_code="2.5.2", build_up_code="CW", quantity="600.0", notes="glazed"),
            dict(base, element_id="MS-0009", nrm_code="2.5.1", build_up_code="SFS", quantity="400.0", notes="opaque"),
            dict(base, element_id="MS-0007", nrm_code="2.5.1", build_up_code="SOFFIT", quantity="83.0", notes="soffit")]


def _option(mix):
    o = ws.default_option("A", layer="L")
    o["facade"] = mix
    o["substructure"]["basement"]["answer"] = "none"
    return o


def test_facade_mix_splits_net_facade():
    o = _option([{"build_up": "CW", "kind": ws.FACADE_UNITISED, "share": 0.7},
                 {"build_up": "SFS", "kind": ws.FACADE_PUNCHED, "share": 0.3, "wwr": 0.4}])
    out = ws.apply_facade_mix(_rows(), o, OFFICE, glazed_code="CW", opaque_code="SFS")
    q = dict(((r["build_up_code"], r["nrm_code"]), r["quantity"]) for r in out)
    assert q[("CW", "2.5.2")] == pytest.approx(700.0)          # 0.7 x 1000
    assert q[("SFS", "2.5.1")] == pytest.approx(180.0)         # 0.3 x 0.6 x 1000
    assert q[("WIN", "2.6")] == pytest.approx(120.0)           # 0.3 x 0.4 x 1000, windows report to 2.6
    assert q[("SOFFIT", "2.5.1")] == "83.0"                    # untouched
    assert sum(float(r["quantity"]) for r in out if r["build_up_code"] != "SOFFIT") == pytest.approx(1000.0)


def test_facade_types_sharing_a_wall_keep_their_own_windows():
    """Two catalogue types on the same wall build-up (ids differ) route to different window build-ups; a punched
    type without a window build-up keeps the curtain wall stand-in at 2.5.2."""
    o = _option([{"build_up": "SFS", "kind": ws.FACADE_PUNCHED, "share": 0.5, "wwr": 0.4},
                 {"build_up": "SFS-T", "kind": ws.FACADE_PUNCHED, "share": 0.3, "wwr": 0.5},
                 {"build_up": "OLD", "kind": ws.FACADE_PUNCHED, "share": 0.2, "wwr": 0.5}])
    out = ws.apply_facade_mix(_rows(), o, OFFICE, glazed_code="CW", opaque_code="SFS")
    q = {}
    for r in out:
        q[(r["build_up_code"], r["nrm_code"])] = q.get((r["build_up_code"], r["nrm_code"]), 0) + float(r["quantity"])
    assert q[("SFS", "2.5.1")] == pytest.approx(0.5 * 0.6 * 1000 + 0.3 * 0.5 * 1000)
    assert q[("WIN", "2.6")] == pytest.approx(200.0) and q[("WIN-T", "2.6")] == pytest.approx(150.0)
    assert q[("OLD", "2.5.1")] == pytest.approx(100.0) and q[("CW", "2.5.2")] == pytest.approx(100.0)
    assert [m for lv, m in ws.validate(dict(ws.default_state(), options=[o]), OFFICE) if "not in the office" in m] == []


def test_validate_flags_bad_mix_and_missing_inputs():
    st = ws.default_state()
    st["project"]["name"] = "P"
    st["output"]["folder"] = "x"
    st["options"] = [_option([{"build_up": "CW", "kind": ws.FACADE_UNITISED, "share": 0.6},
                              {"build_up": "SFS", "kind": ws.FACADE_PUNCHED, "share": 0.3, "wwr": 0.4}])]
    errs = [m for lv, m in ws.validate(st, OFFICE) if lv == "error"]
    assert any("90%" in m for m in errs)
    st["options"][0]["facade"][0]["share"] = 0.7
    assert [m for lv, m in ws.validate(st, OFFICE) if lv == "error"] == []
    st["options"].append(dict(ws.default_option("A"), layer=""))
    errs = [m for lv, m in ws.validate(st, OFFICE) if lv == "error"]
    assert any("different" in m for m in errs) and any("massing layer" in m for m in errs)


def test_state_roundtrip_fills_new_keys():
    st = ws.loads(json.dumps({"project": {"name": "P"}}))
    assert st["project"]["name"] == "P" and st["project"]["use_class"] == "Office"
    assert st["targets"]["sc_share"]["services"] == 1.0
    assert ws.loads(ws.dumps(st)) == st


def test_office_paths_resolve_relative_to_settings(tmp_path):
    p = tmp_path / "07_Settings" / "office.json"
    p.parent.mkdir()
    p.write_text(json.dumps({"library_xlsx": "../03/lib.xlsx", "benchmarks": {"path": "b.xlsx", "sheet": "S"}}))
    off = ws.load_office(str(p))
    assert off["library_xlsx"] == os.path.normpath(str(tmp_path / "03" / "lib.xlsx"))
    assert off["benchmarks"]["path"] == str(p.parent / "b.xlsx")


def test_configs(defaults, tmp_path):
    st = ws.default_state()
    st["project"].update(name="2109 Test", start_year=2028)
    st["output"]["folder"] = str(tmp_path)
    st["schedule"].update(path="s.xlsx", sheet="Office")
    st["targets"]["show"]["RIBA"] = False
    st["targets"]["catb_allowance"] = "50"
    o = ws.default_option("A - CW", layer="01::Opt")
    tk = ws.takeoff_config(st, o, {}, defaults)
    assert tk["massing_layers"] == ["01::Opt"] and tk["area_schedule"]["sheet"] == "Office"
    assert tk["output_dir"] == os.path.join(str(tmp_path), "takeoff") and tk["file_stem"] == "A_-_CW_takeoff"
    assert defaults["massing_layers"] is None                   # base CONFIG not mutated
    off = {"library_xlsx": "lib", "benchmarks": {"path": "b"}, "params": {"rsp": 60}}
    cc = ws.calc_config(st, off, [(o, "a.csv")])
    assert cc["params"]["construction_start_year"] == 2028 and cc["params"]["show_targets"]["RIBA"] == "N"
    assert cc["params"]["catb_allowance"] == 50.0 and cc["params"]["rsp"] == 60
    assert cc["options"][0]["facade_route"] == "As exported"
    assert cc["out"] == os.path.join(str(tmp_path), "2109_Test_WLCA.xlsx")


def test_mixed_export_builds_workbook(tmp_path):
    """Take-off export -> facade mix -> CSV -> workbook generator (invented library from the calc tests)."""
    openpyxl = pytest.importorskip("openpyxl")
    import wlca_calc_template as ct
    from test_calc_template import _bench, _export, _library
    lib, exp, bm = str(tmp_path / "lib.xlsx"), str(tmp_path / "e.csv"), str(tmp_path / "bm.xlsx")
    _library(lib), _export(exp), _bench(bm)
    header, rows = ws.read_export(exp)
    office = {"facade_types": [{"code": "PANEL", "label": "Panel", "kind": ws.FACADE_UNITISED},
                               {"code": "OPQ", "label": "Opaque", "kind": ws.FACADE_PUNCHED, "glazing_build_up": "PANEL"}],
              "library_xlsx": lib, "benchmarks": {"path": bm, "sheet": "Bench", "source": "invented"},
              "site": {"A51": 35, "A52": 40, "C1": 12},
              "params": {"road_vehicle": "HGV", "eol_vehicle": "HGV", "sea_factor": 0.016, "eol_distances_km": [50, 60, 70],
                         "facade_opaque_build_up": "OPQ", "cw_build_up": "PANEL"}}
    o = _option([{"build_up": "PANEL", "kind": ws.FACADE_UNITISED, "share": 0.5},
                 {"build_up": "OPQ", "kind": ws.FACADE_PUNCHED, "share": 0.5, "wwr": 0.2}])
    mixed = ws.apply_facade_mix(rows, o, office, glazed_code="PANEL", opaque_code="OPQ")
    out = str(tmp_path / "mixed.csv")
    ws.write_export(out, header, mixed)
    st = ws.default_state()
    st["project"]["name"] = "T"
    st["output"]["folder"] = str(tmp_path)
    cfg = ws.calc_config(st, office, [(o, out)])
    wb = ct.build(cfg)
    T = wb["O1_Takeoff"]
    got = sorted((T.cell(r, 3).value, T.cell(r, 2).value, T.cell(r, 4).value) for r in range(5, 12) if T.cell(r, 3).value)
    # net facade = 500 (PANEL) + 200 (OPQ) = 700: PANEL 350 + 0.5 x 0.2 x 700 = 70 (windows -> 2.6); OPQ 0.5 x 0.8 x 700 = 280
    assert ("OPQ", "2.5.1", 280.0) in got and ("PANEL", "2.5.2", 350.0) in got and ("PANEL", "2.6", 70.0) in got


def test_plan_route_rhino_feeds_the_takeoff():
    st = ws.default_state()
    assert st["plan"]["route"] == "rhino"                                  # D5: Rhino layers are the default route
    st["project"]["name"], st["output"]["folder"] = "P", "out"
    st["options"] = [ws.default_option("A", layer="L")]
    assert any("No plan" in m for lv, m in ws.validate(st, {}))           # nothing picked yet: office defaults, warned
    st["plan"].update(walls_layer="EF_25 walls", json="out/takeoff/plan_rhino.json")
    assert not any("No plan" in m for lv, m in ws.validate(st, {}))
    cfg = ws.takeoff_config(st, st["options"][0], {}, {"structure": {}, "build_ups": {}})
    assert cfg["structure"]["plan_json"] == "out/takeoff/plan_rhino.json"
    plan = {"core": {"wall_length_m_per_floor": 137.1, "mean_thickness_m": 0.291, "pieces": 30},
            "columns": {"count_per_floor": 20, "size_m": 0.75, "shapes": ["circle"]},
            "stairs": {"count": 2, "flights_per_stair": [2, 2], "treads_per_flight": [11] * 4, "flight_width_m": 1.25, "going_m": 0.25}}
    lines = ws.plan_summary(plan, ["Floor plate: ..."])
    assert lines[0] == "Core walls: 137 m per floor, 291 mm average (30 pieces)"
    assert lines[1] == "Columns: 20 per floor, 0.75 m circle" and lines[2].startswith("Stairs: 2 - 2/2 flights, 11 treads")


def test_operational_inputs_reach_the_workbook_config():
    st = ws.default_state()
    st["project"]["name"], st["output"]["folder"] = "P", "out"
    o = ws.default_option("A", layer="L")
    st["options"] = [o]
    assert st["operational"]["route"] == "Not assessed"
    st["operational"].update(route="Engineer prediction (TM54)", op_year=2031, dh_factor=0.05)
    assert any("predicted annual energy" in m for lv, m in ws.validate(st, {}) if lv == "error")
    o["operational"].update(elec=120000, water=900)
    assert not [m for lv, m in ws.validate(st, {}) if lv == "error" and "energy" in m]
    cc = ws.calc_config(st, {"library_xlsx": "lib", "benchmarks": {"path": "b"}}, [(o, "a.csv")])
    assert cc["params"]["op_route"] == "Engineer prediction (TM54)" and cc["params"]["op_year"] == 2031
    assert cc["params"]["dh_factor"] == 0.05 and cc["options"][0]["operational"] == {"elec": 120000, "water": 900}
    st["project"].update(occupancy=850, occ_days=250)                   # asked on step 1: sets the water budget
    cc = ws.calc_config(st, {"library_xlsx": "lib", "benchmarks": {"path": "b"}}, [(o, "a.csv")])
    assert cc["params"]["occupancy"] == 850 and cc["params"]["occ_days"] == 250


def test_water_warning_follows_occupancy():
    st = ws.default_state()
    st["project"]["name"], st["output"]["folder"] = "P", "out"
    o = ws.default_option("A", layer="L")
    o["operational"]["elec"] = 1000
    st["options"] = [o]
    st["operational"]["route"] = "Engineer prediction (TM54)"
    water = lambda: [m for lv, m in ws.validate(st, {}) if "water" in m.lower()]  # noqa: E731
    assert "not assessed" in water()[0]
    st["project"]["occupancy"] = 850
    assert "occupancy water budget" in water()[0]
    o["operational"]["water"] = 900
    assert water() == []


def test_workbook_font_and_logo(tmp_path):
    st = ws.default_state()
    st["project"]["name"], st["output"]["folder"] = "P", str(tmp_path)
    st["options"] = [ws.default_option("A", layer="L")]
    office = {"library_xlsx": "lib", "benchmarks": {"path": "b"}, "branding": {"font": "Suisse Int'l", "logo": "office.png"}}
    cfg = lambda: ws.calc_config(st, office, [(st["options"][0], "a.csv")])["branding"]  # noqa: E731
    assert cfg()["font"] == "Arial" and cfg()["logo"] == "office.png"       # dialog default font; office logo until changed
    st["output"]["font"] = "Georgia"
    st["output"]["logo"] = ""                                                # "No logo"
    assert cfg()["font"] == "Georgia" and cfg()["logo"] is None
    logo = tmp_path / "ours.png"
    st["output"]["logo"] = str(logo)
    assert cfg()["logo"] == str(logo)
    assert any("Logo file not found" in m for _, m in ws.validate(st, office))
    logo.write_bytes(b"x")
    assert not any("Logo" in m for _, m in ws.validate(st, office))
    assert ws.loads('{"output": {"folder": "f"}}')["output"]["font"] == "Arial"   # older saves get the default


def test_apply_font_keeps_size_and_weight():
    import openpyxl
    from openpyxl.styles import Font
    import wlca_calc_template as ct
    wb = openpyxl.Workbook()
    a = wb.active
    a["A1"], a["A2"] = "title", "plain"
    a["A1"].font = Font(name="Calibri", size=18, bold=True, color="FF0000")
    a["A2"].number_format = "0.0"
    ct.apply_font(wb, "Georgia")
    assert (a["A1"].font.name, a["A1"].font.sz, a["A1"].font.b, a["A1"].font.color.rgb) == ("Georgia", 18, True, "00FF0000")
    assert a["A2"].font.name == "Georgia" and wb._fonts[0].name == "Georgia"
    assert wb._fonts[0].scheme is None and a["A2"].font.scheme is None      # no theme font overriding the name


def test_logo_without_pillow(tmp_path, monkeypatch):
    """The dashboard logo is embedded even where openpyxl cannot import Pillow (Rhino 8)."""
    import struct
    import zlib
    import openpyxl
    import openpyxl.drawing.image as oi
    import wlca_dashboard as db
    raw = b"\x00" + b"\x00\x00\x00" * 3               # 3 x 1 RGB, one filter byte
    chunk = lambda t, d: struct.pack(">I", len(d)) + t + d + struct.pack(">I", zlib.crc32(t + d))  # noqa: E731
    png = b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", 3, 1, 8, 2, 0, 0, 0)) + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b"")
    p = tmp_path / "logo.png"
    p.write_bytes(png)

    def no_pillow(img):
        raise ImportError("You must install Pillow to fetch image objects")
    monkeypatch.setattr(oi, "_import_image", no_pillow)
    img = db.logo_image(str(p))
    assert (img.format, img.width, img.height) == ("png", 3, 1)
    wb = openpyxl.Workbook()
    wb.active.add_image(img, "A1")
    out = tmp_path / "x.xlsx"
    wb.save(out)
    import zipfile
    assert zipfile.ZipFile(out).read("xl/media/image1.png") == png


def test_office_settings_lookup_order(tmp_path, monkeypatch):
    pkg = tmp_path / "pkg"
    pkg.mkdir()
    old = tmp_path / "07_Settings" / "wlca_office_settings.json"
    old.parent.mkdir()
    old.write_text("{}")
    assert ws.find_office_settings(str(pkg), env={}, user_cfg={}) == (str(old), "next to the package folder")
    shared = tmp_path / "shared.json"
    shared.write_text("{}")
    (pkg / "wlca_package.json").write_text(json.dumps({"office_settings": str(shared), "name": "Office WLCA", "version": "0.1.0"}))
    assert ws.find_office_settings(str(pkg), env={}, user_cfg={}) == (str(shared), "package default")
    mine = tmp_path / "mine.json"
    mine.write_text("{}")
    assert ws.find_office_settings(str(pkg), env={}, user_cfg={"office_settings": str(mine)})[1] == "chosen on this PC"
    assert ws.find_office_settings(str(pkg), env={"WLCA_OFFICE_SETTINGS": str(old)}, user_cfg={"office_settings": str(mine)})[0] == str(old)
    gone = {"office_settings": str(tmp_path / "missing.json")}           # a remembered file that has gone: fall through
    assert ws.find_office_settings(str(pkg), env={}, user_cfg=gone)[1] == "package default"
    assert ws.package_info(str(pkg)) == ("Office WLCA", "0.1.0")
    monkeypatch.setenv("APPDATA", str(tmp_path / "appdata"))
    ws.remember_office_settings(str(mine))
    assert ws.read_json(ws.user_config_path())["office_settings"] == str(mine)
    ws.remember_office_settings("")
    assert "office_settings" not in ws.read_json(ws.user_config_path())



def test_edits_from_the_web_window():
    st = ws.default_state()
    st["options"] = [ws.default_option("A", layer="L")]
    assert ws.apply_edit(st, ["project", "name"], "Opt 7") == "Opt 7"
    assert ws.apply_edit(st, ["project", "riba_stage"], "3", "int") == 3
    assert ws.apply_edit(st, ["project", "riba_stage"], "x", "int") == 3                # bad text keeps the old value
    assert ws.apply_edit(st, ["options", 0, "facade", 0, "share"], "70%", "pct") == pytest.approx(0.7)
    assert ws.apply_edit(st, ["options", 0, "facade", 0, "share"], "0.5", "pct") == pytest.approx(0.5)
    assert ws.apply_edit(st, ["options", 0, "operational", "gas"], "12,000", "num") == 12000
    assert ws.apply_edit(st, ["options", 0, "operational", "gas"], "", "num") is None
    assert ws.apply_edit(st, ["project", "demolished_gia_m2"], "", "num0") == 0.0
    assert ws.apply_edit(st, ["targets", "show", "LETI"], False, "bool") is False
    assert ws.apply_edit(st, ["output", "font"], "  ") == "Arial"
    with pytest.raises((KeyError, IndexError)):
        ws.apply_edit(st, ["options", 5, "name"], "x")


# ------------------------------------------------------------------------------------- substructure step (step 3)
def _steps(path, name="STEPS"):
    """A list constant from a Rhino-only script, read without importing Rhino."""
    import ast
    tree = ast.parse(open(os.path.join(os.path.dirname(__file__), "..", "rhino_ui", path), encoding="utf-8").read())
    for node in tree.body:
        if isinstance(node, ast.Assign) and getattr(node.targets[0], "id", None) == name:
            return ast.literal_eval(node.value)
    raise AssertionError(name + " not found")


def test_step_order():
    web = _steps("wlca_web.py")
    assert web == ["Project & options", "Drawings, schedule & model", "Substructure", "Superstructure - Structure",
                   "Superstructure - Façade", "Targets & reporting", "Run", "Results"]
    js = open(os.path.join(os.path.dirname(__file__), "..", "rhino_ui", "web", "app.js"), encoding="utf-8").read()
    assert "const PAGES = [pageProject, pageDrawings, pageSubstructure, pageStructure, pageFacade, pageTargets, pageRun, pageResults];" in js


def test_check_messages_point_at_the_right_step(tmp_path):
    st = ws.default_state()
    st["project"]["name"] = ""
    st["options"] = [_option([{"build_up": "CW", "kind": ws.FACADE_UNITISED, "share": 0.5}])]
    st["options"][0]["substructure"]["piles"].update(source="engineer", count=None)
    st["targets"]["sc_share"]["services"] = 1.5
    st["output"]["logo"] = str(tmp_path / "missing.png")
    msgs = [m for _, m in ws.validate(st, OFFICE)]
    where = {"Project name": 1, "facade shares": 5, "piles from the engineer": 3, "Shell & core share": 6,
             "where to save": 6, "Logo file": 7}
    for key, n in where.items():
        hit = [m for m in msgs if key.lower() in m.lower()]
        assert hit and all("(step %d)" % n in m for m in hit), (key, hit)


def test_substructure_defaults_match_the_takeoff(defaults):
    assert ws.default_substructure() == defaults["substructure"]
    assert ws.default_option("A")["substructure"] == defaults["substructure"]


def test_older_saves_get_the_substructure_block():
    old = ws.default_option("A")
    del old["substructure"]
    st = ws.loads(json.dumps({"project": {"name": "P"}, "options": [old]}))
    assert st["options"][0]["substructure"] == ws.default_substructure()
    part = {"piles": {"length_m": 30.0}}
    o = ws.fill_substructure(dict(ws.default_option("B"), substructure=part))
    assert o["substructure"]["piles"]["length_m"] == 30.0 and o["substructure"]["piles"]["diameter_m"] == 0.75


def test_option_substructure_reaches_the_takeoff(defaults):
    st = ws.default_state()
    st["project"]["name"] = "P"
    st["output"]["folder"] = "x"
    o = ws.default_option("A", layer="L")
    o["substructure"]["retaining_walls"].update(type="rc_wall", perimeter_m=180.0)
    o["substructure"]["piles"].update(source="engineer", count=120, diameter_m=0.9, length_m=30.0)
    o["substructure"]["basement_depth_m"] = 7.5
    office = {"takeoff_defaults": {"substructure": {"piles": {"length_m": 20.0}, "caps_raft": {"depth_m": 1.5}}}}
    tk = ws.takeoff_config(st, o, office, defaults)
    sub = tk["substructure"]
    assert sub["retaining_walls"]["type"] == "rc_wall" and sub["retaining_walls"]["perimeter_m"] == 180.0
    assert sub["piles"]["source"] == "engineer" and sub["piles"]["length_m"] == 30.0     # option wins over the office
    assert sub["basement_depth_m"] == 7.5
    assert defaults["substructure"]["piles"]["source"] == "assumed"                     # base CONFIG not mutated


def test_substructure_checks():
    def errs(**upd):
        o = ws.default_option("A", layer="L")
        o["substructure"]["basement"]["answer"] = "none"
        for k, v in upd.items():
            o["substructure"][k].update(v)
        return [(lv, m) for lv, m in ws.check_substructure(o, "A")]
    assert errs() == []
    assert any("spacing" in m for _, m in errs(retaining_walls={"pile_spacing_m": 2.1}))       # firm pile cannot reach
    assert any("spacing" in m for _, m in errs(retaining_walls={"pile_spacing_m": 0.9}))       # hard piles overlap
    assert errs(retaining_walls={"type": "rc_wall", "pile_spacing_m": None}) == []
    assert any("volume, an area or a perimeter" in m for _, m in errs(retaining_walls={"source": "engineer"}))
    assert errs(retaining_walls={"source": "engineer", "area_m2": 900}) == []
    assert errs(piles={"source": "engineer"}, caps_raft={"source": "engineer", "volume_m3": 300}) == \
        [("error", "A: piles from the engineer need the number, diameter and length (step 3).")]
    assert errs(piles={"source": "engineer", "count": 40}, caps_raft={"source": "engineer", "volume_m3": 300}) == []
    w = errs(piles={"source": "none"})
    assert [lv for lv, _ in w] == ["warning"]                                                  # assumed caps need assumed piles
    assert any("need a volume" in m for _, m in errs(caps_raft={"source": "engineer"}))
    assert any("diameter, piles per column" in m for _, m in errs(piles={"length_m": None}))


def test_basement_from_schedule(schedule_xlsx):
    b = ws.basement_from_schedule(schedule_xlsx, "Office")
    assert b["levels"] == ["B2", "B1"] and b["depth_m"] == 9.0
    assert b["lowest_gea"] == 1200 and b["ground_gea"] == 1000
    assert ws.basement_from_schedule("", None) is None and ws.basement_from_schedule("nope.xlsx") is None


def test_basement_must_be_confirmed_unless_the_schedule_has_one(schedule_xlsx):
    o = ws.default_option("A", layer="L")
    bs = o["substructure"]["basement"]
    msg = lambda sb=None: [(lv, m) for lv, m in ws.check_substructure(o, "A", sb) if "basement" in m.lower()]  # noqa: E731
    assert msg()[0][0] == "error" and "no area schedule" in msg()[0][1]
    sched = ws.basement_from_schedule(schedule_xlsx, "Office")
    assert msg(sched) == []                                             # B2, B1 in the schedule: nothing to confirm
    no_bsm = dict(sched, levels=[])
    assert "no levels below ground" in msg(no_bsm)[0][1]
    bs["answer"] = "none"
    assert msg(no_bsm) == []
    bs["answer"] = "levels"
    assert "at least one basement level" in msg()[0][1]
    bs["levels"] = [ws.new_basement_level(1)]
    assert [lv for lv, _ in msg()] == ["warning"]                       # GEA / GIA from the massing
    bs["levels"][0].update(gea=900.0, gia=850.0)
    assert msg() == []
    bs["levels"][0]["ftf"] = None
    assert "floor-to-floor" in msg()[0][1]


def test_raft_check():
    o = ws.default_option("A", layer="L")
    o["substructure"]["basement"]["answer"] = "none"
    o["substructure"]["caps_raft"].update(source="raft")
    assert ws.check_substructure(o, "A") == []
    o["substructure"]["caps_raft"]["raft_thickness_m"] = None
    assert "raft thickness" in ws.check_substructure(o, "A")[0][1]


def test_paths_follow_a_moved_project(tmp_path):
    old, new = tmp_path / "P" / "proj", tmp_path / "C" / "samples" / "proj"
    for base in (old, new):
        (base / "WLCA").mkdir(parents=True)
        (base / "schedule.xlsx").write_text("x")
    other = tmp_path / "V" / "logo.png"
    other.parent.mkdir()
    other.write_text("x")
    st = ws.default_state()
    st["schedule"]["path"] = str(old / "schedule.xlsx")
    st["output"].update(folder=str(old / "WLCA"), logo=str(other), last_workbook=str(old / "WLCA" / "P_WLCA.xlsx"))
    st["options"] = [dict(ws.default_option("A"), source="file", file=str(tmp_path / "elsewhere.3dm"))]
    saved = ws.dumps(st, str(old))
    rel = json.loads(saved)["_rel"]
    assert rel == {"schedule.path": "schedule.xlsx", "output.folder": "WLCA",
                   "output.last_workbook": os.path.join("WLCA", "P_WLCA.xlsx")}          # logo / option file are elsewhere
    assert "_rel" not in st                                                           # in-memory state untouched
    # the whole folder moved: inputs found there, output follows the model, outside paths unchanged
    moved = ws.loads(saved, str(new))
    assert moved["schedule"]["path"] == str(new / "schedule.xlsx")
    assert moved["output"]["folder"] == str(new / "WLCA")
    assert moved["output"]["logo"] == str(other) and moved["options"][0]["file"] == str(tmp_path / "elsewhere.3dm")
    assert moved["output"]["last_workbook"] == str(new / "WLCA" / "P_WLCA.xlsx")      # neither exists: follows the model
    # only the .3dm copied somewhere else: the schedule is still read from where it is
    lone = tmp_path / "D" / "copy"
    lone.mkdir(parents=True)
    copied = ws.loads(saved, str(lone))
    assert copied["schedule"]["path"] == str(old / "schedule.xlsx")
    assert copied["output"]["folder"] == str(lone / "WLCA")
    assert ws.loads(saved)["schedule"]["path"] == str(old / "schedule.xlsx")         # unsaved model: as saved
    assert ws.loads(ws.dumps(moved, str(new)), str(new)) == moved


def test_facade_header():
    o = _option([{"build_up": "CW", "kind": ws.FACADE_UNITISED, "share": 0.25},
                 {"build_up": "SFS", "kind": ws.FACADE_PUNCHED, "share": 0.75, "wwr": 0.2}])
    assert ws.facade_header(o, OFFICE) == {"facade_curtain_wall_share": 0.25, "facade_punched_window_share": 0.75}
    o = _option([{"build_up": "SFS", "kind": ws.FACADE_PUNCHED, "share": 1.0, "wwr": 0.0}])
    assert ws.facade_header(o, OFFICE) == {"facade_curtain_wall_share": 0.0, "facade_punched_window_share": 0.0}


def test_interface_choice_is_remembered_per_pc(tmp_path):
    """XW 27/10/2026: the new panel (web_next) is the default; the classic one is kept for anyone who switches back,
    and is the fallback when the new page is not in the package."""
    cfg = str(tmp_path / "cfg" / "wlca_user.json")
    assert ws.ui_choice(cfg) == "next"
    ws.remember_ui("classic", cfg)
    assert ws.ui_choice(cfg) == "classic"
    ws.remember_ui("next", cfg)
    assert ws.ui_choice(cfg) == "next" and "ui" not in json.load(open(cfg))
    ws.remember_ui("bogus", cfg)
    assert ws.ui_choice(cfg) == "next"
    pkg = tmp_path / "pkg"
    (pkg / "web").mkdir(parents=True)
    (pkg / "web" / "index.html").write_text("x")
    assert ws.ui_page(str(pkg), "next").endswith(os.path.join("web", "index.html"))      # not shipped: current panel
    (pkg / "web_next").mkdir()
    (pkg / "web_next" / "index.html").write_text("x")
    assert ws.ui_page(str(pkg), "next").endswith(os.path.join("web_next", "index.html"))
    assert ws.ui_page(str(pkg), "classic").endswith(os.path.join("web", "index.html"))
